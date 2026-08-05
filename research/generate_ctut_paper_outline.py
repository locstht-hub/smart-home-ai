from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from generate_vietnamese_paper_final_revision import (
    REFERENCES,
    add_figure,
    captioned_table,
    clear_after_front_matter,
    clone_paragraph,
    set_runs,
    set_text,
)


TITLE_VI = "THIẾT KẾ MÔ HÌNH GIÁM SÁT NĂNG LƯỢNG VÀ DỰ BÁO PHỤ TẢI CHO PHÒNG THÍ NGHIỆM ĐIỆN CÔNG NGHIỆP"
TITLE_EN = "Design of an Energy Monitoring and Load Forecasting Model for an Industrial Electrical Engineering Laboratory"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate a reduced CTUT-formatted paper outline.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def shade_and_border(paragraph, fill: str = "F3F6F8", border: str = "7A8793") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "5")
        node.set(qn("w:color"), border)
        borders.append(node)


def add_planning_note(document: Document, sample, text: str, *, keep_next: bool = False):
    paragraph = clone_paragraph(document, sample, text, keep_next=keep_next)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        run.italic = True
        run.font.color.rgb = None
    shade_and_border(paragraph, fill="F4F6F7", border="A7B0B8")
    return paragraph


def add_figure_slot(document: Document, samples: dict, number: int, title: str, required: str, caption: str) -> None:
    slot = clone_paragraph(
        document,
        samples["body"],
        f"VỊ TRÍ HÌNH {number}\n{title}\nBằng chứng cần có: {required}",
        keep_next=True,
    )
    slot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    slot.paragraph_format.space_before = Pt(6)
    slot.paragraph_format.space_after = Pt(6)
    slot.paragraph_format.left_indent = Inches(0.55)
    slot.paragraph_format.right_indent = Inches(0.55)
    for index, run in enumerate(slot.runs):
        run.bold = index == 0
        run.italic = True
    shade_and_border(slot, fill="EEF4F7", border="6D8797")
    caption_p = clone_paragraph(document, samples["figure_caption"], caption)
    caption_p.paragraph_format.keep_together = True


def populate_front_matter(document: Document) -> None:
    set_text(document.paragraphs[0], TITLE_VI)
    set_text(document.paragraphs[1], "BẢN DÀN Ý RÚT GỌN — CHƯA PHẢI BẢN THẢO HOÀN CHỈNH")
    document.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in document.paragraphs[1].runs:
        run.italic = True
        run.font.size = Pt(9)
    set_text(document.paragraphs[3], "[HỌ VÀ TÊN TÁC GIẢ]1, [HỌ VÀ TÊN ĐỒNG TÁC GIẢ]2")
    set_text(document.paragraphs[4], "1 [KHOA/ĐƠN VỊ, TRƯỜNG ĐẠI HỌC KỸ THUẬT - CÔNG NGHỆ CẦN THƠ]")
    set_text(document.paragraphs[5], "2 [ĐƠN VỊ CỦA ĐỒNG TÁC GIẢ, NẾU CÓ]")
    set_text(document.paragraphs[6], "Tác giả liên hệ: [EMAIL TÁC GIẢ CHỊU TRÁCH NHIỆM]")
    set_runs(
        document.tables[0].cell(0, 0).paragraphs[0],
        [
            "Thông tin chung\n",
            "Ngày nhận bài: dd/mm/yyyy\nNgày nhận bài sửa: dd/mm/yyyy\nNgày duyệt đăng: dd/mm/yyyy\n\n",
            "Từ khóa:\n",
            "dự báo phụ tải, giám sát năng lượng, MFM384, PLC S7-1200, quản lý phụ tải",
        ],
    )
    vi_outline = (
        "Khung tóm tắt sẽ được viết thành một đoạn 150–200 từ sau khi hoàn tất thí nghiệm. "
        "Đoạn mở đầu nêu nhu cầu giám sát và quản lý phụ tải trong phòng thí nghiệm Điện công nghiệp. "
        "Phần phương pháp giới thiệu chuỗi MFM384–Modbus RTU–PLC S7-1200–relay/contactor, Web/App và mô-đun dự báo 24 giờ. "
        "Phần đánh giá phải nêu dữ liệu, cách chia theo thời gian, kịch bản đo, điều khiển, timeout, Quota và phục hồi. "
        "Kết quả chỉ sử dụng số liệu đã kiểm chứng; kết quả dự báo công khai được tách khỏi kết quả phần cứng địa phương. "
        "Câu cuối nêu đóng góp, giới hạn và không tuyên bố tiết kiệm năng lượng hoặc sa thải tải tự động nếu chưa có log tải thật."
    )
    set_text(document.tables[0].cell(0, 1).paragraphs[1], vi_outline)
    set_runs(
        document.tables[1].cell(0, 0).paragraphs[0],
        [
            "Title:\n",
            TITLE_EN + "\n\n",
            "Keywords:\n",
            "energy monitoring, load forecasting, load management, MFM384, Siemens S7-1200 PLC",
        ],
    )
    en_outline = (
        "The English abstract will be drafted only after the Vietnamese version and the hardware evidence are approved. "
        "It will state the laboratory context, the MFM384–Modbus RTU–S7-1200 measurement and control chain, the Web/mobile monitoring functions, and the 24-hour forecasting task. "
        "The evaluation statement will distinguish the public-data benchmark from local hardware trials. "
        "Only verified measurement, command-feedback, quota, and forecast results will be reported. "
        "The final sentence will delimit claims about local accuracy, control latency, energy savings, and automatic load shedding."
    )
    set_text(document.tables[1].cell(0, 1).paragraphs[1], en_outline)


def build_outline(document: Document, samples: dict, figures_dir: Path) -> None:
    def p(kind: str, text: str):
        return clone_paragraph(document, samples[kind], text)

    p("h1", "1. ĐẶT VẤN ĐỀ")
    add_planning_note(
        document,
        samples["body"],
        "Khung viết: 4 đoạn. Đoạn 1 nêu nhu cầu đo–giám sát–điều khiển trong phòng thí nghiệm. "
        "Đoạn 2 nhóm nghiên cứu liên quan theo HEMS, dự báo và demand response. "
        "Đoạn 3 chỉ ra khoảng trống chuỗi MFM384–PLC–tải–Web/App. "
        "Đoạn 4 nêu mục tiêu, câu hỏi nghiên cứu và ba đóng góp.",
    )
    p("h2", "1.1. Khoảng trống và câu hỏi nghiên cứu")
    p("body", "RQ1: Chuỗi thiết bị nào liên kết phép đo V/I/P/E, PLC S7-1200, cơ cấu đóng cắt và giao diện giám sát mà vẫn phản hồi đúng trạng thái vận hành?")
    p("body", "RQ2: Random Forest và XGBoost dự báo h+1 đến h+24 như thế nào so với các đường cơ sở khi đánh giá theo thứ tự thời gian?")
    p("body", "RQ3: Các kịch bản Quota, cảnh báo, điều khiển và sa thải phụ tải nào đã có bằng chứng thực nghiệm, và kịch bản nào mới ở mức thiết kế an toàn?")
    p("h2", "1.2. Đóng góp và phạm vi")
    add_planning_note(
        document,
        samples["body"],
        "Chỉ giữ ba đóng góp: (1) chuỗi đo–điều khiển điện; (2) giám sát từ xa, phản hồi PLC và Quota; "
        "(3) đánh giá dự báo 24 giờ và quy trình kiểm chứng. Không nhận thuật toán AI mới, không nhận hiệu quả tiết kiệm khi chưa có thí nghiệm.",
    )

    p("h1", "2. CƠ SỞ LÝ THUYẾT VÀ MÔ HÌNH ĐỀ XUẤT")
    p("h2", "2.1. Công trình liên quan")
    p("body", "Phần này nhóm tài liệu theo đối tượng điện, chuỗi đo–điều khiển, dự báo và quản lý tải; không liệt kê tuần tự từng bài.")
    related_rows = [
        ["Nhóm công trình", "Đối tượng", "Đo/điều khiển", "Dự báo", "Quota/DR", "Bằng chứng"],
        ["Tổng quan HEMS", "Nhà ở", "Thiết bị thông minh", "Tùy công trình", "Có", "Tổng quan"],
        ["Dự báo phụ tải", "Hộ/cộng đồng", "Smart meter", "Có", "Ít", "Dữ liệu công khai"],
        ["HEMS và demand response", "Thiết bị gia dụng", "Điều phối tải", "Có/không", "Có", "Mô phỏng/thực nghiệm"],
        ["Nghiên cứu này", "Phòng thí nghiệm", "MFM384–PLC–tải", "24 giờ", "Quota", "Benchmark + thử tải thật"],
    ]
    captioned_table(
        document,
        samples,
        "Bảng 1. Khung đối chiếu các nghiên cứu gần đề tài",
        related_rows,
        "Ghi chú: Chỉ ghi “thử tải thật” sau khi có log và trạng thái tiếp điểm được chấp nhận.",
        [1.20, 1.10, 1.50, 1.05, 0.95, 1.10],
        size=7.8,
    )

    p("h2", "2.2. Kiến trúc tổng thể")
    p("body", "Mục này giải thích hai nhánh chính: nhánh dữ liệu đo từ MFM384 đến lưu trữ/dự báo và nhánh điều khiển từ Web/App qua PLC đến relay/contactor rồi trả phản hồi.")
    add_figure(
        document,
        samples["image"],
        samples["figure_caption"],
        figures_dir / "final_architecture_vi.png",
        "Hình 1. Kiến trúc dự kiến của chuỗi đo–điều khiển và quản lý phụ tải; các kết luận phần cứng phải được xác nhận bằng log thử nghiệm",
        "Sơ đồ MFM384, PLC S7-1200, relay hoặc contactor, kho dữ liệu, mô-đun dự báo và Web/App.",
        width=6.40,
    )

    p("h2", "2.3. Phần cứng, đo lường và mạch điều khiển")
    p("body", "Trình bày MFM384, RS-485/Modbus RTU, S7-1200, nguồn điều khiển, relay/contactor, bảo vệ, phụ tải, tiếp điểm phản hồi, nút dừng khẩn và cách ly mạch lực–mạch điều khiển.")
    device_rows = [
        ["Thiết bị", "Model/thông số", "Tín hiệu/đại lượng", "Giao tiếp", "Vai trò", "Bằng chứng cần có"],
        ["Công-tơ đa năng", "MFM384 — xác minh mã", "V/I/P/Q/PF/E", "RS-485", "Đo năng lượng", "Manual + raw log"],
        ["PLC", "S7-1200", "I/O và liên động", "S7/Modbus", "Điều khiển", "Chương trình + log"],
        ["Relay/contactor", "Điền theo thiết bị thật", "Command/feedback", "Digital I/O", "Đóng cắt tải", "Ảnh + thử tải"],
        ["Phụ tải", "Tên, công suất định mức", "P và trạng thái", "Mạch lực", "Đối tượng thử", "Ảnh + timestamp"],
    ]
    captioned_table(
        document,
        samples,
        "Bảng 2. Thiết bị và bằng chứng cần thu thập",
        device_rows,
        "Ghi chú: Không tự điền địa chỉ thanh ghi, hệ số tỉ lệ hoặc cấp chính xác khi chưa có manual đúng model.",
        [1.16, 1.26, 1.34, 0.90, 1.06, 1.34],
        size=7.6,
    )
    add_figure_slot(
        document,
        samples,
        2,
        "Ảnh thật gồm: (a) toàn cảnh; (b) MFM384 và RS-485; (c) S7-1200; (d) relay/contactor, bảo vệ và tải.",
        "Ảnh gốc rõ nét, nhãn thiết bị, sơ đồ đấu dây đã xác nhận và điều kiện thử nghiệm.",
        "Hình 2. Mô hình phần cứng thực nghiệm dùng để đo điện năng và kiểm tra điều khiển phụ tải",
    )

    p("h2", "2.4. Web/App, phản hồi PLC và Quota")
    p("body", "Chỉ mô tả chức năng phục vụ vận hành điện: hiển thị đại lượng, dự báo, Quota, gửi lệnh và nhận trạng thái. Không mở rộng sang kiến trúc phần mềm, endpoint hoặc giao diện đăng nhập.")
    add_figure_slot(
        document,
        samples,
        3,
        "Ba ảnh chụp màn hình: (a) Dashboard V/I/P/E; (b) dự báo 24 giờ và Quota; (c) điều khiển kèm thành công/từ chối/timeout.",
        "Screenshot đúng phiên bản App, dữ liệu hợp lệ, che tài khoản và địa chỉ mạng.",
        "Hình 3. Giao diện giám sát, dự báo–Quota và điều khiển phụ tải kèm trạng thái phản hồi",
    )
    p("body", "Logic điều khiển được diễn giải bằng văn bản thay cho một hình tuần tự riêng: kiểm tra điều kiện, đọc trạng thái hiện tại, phát xung khi cần, polling statusTag và trả kết quả có nguyên nhân.")

    p("h2", "2.5. Dữ liệu và mô hình dự báo phụ tải")
    add_planning_note(
        document,
        samples["body"],
        "Nêu nguồn dữ liệu, tổng hợp theo giờ, xử lý thiếu theo chiều thời gian, đặc trưng lag/rolling, Random Forest, XGBoost, "
        "ba baseline và dự báo trực tiếp h+1…h+24. Không tạo bảng cấu hình riêng; đưa các tham số quan trọng vào đoạn phương pháp.",
    )

    p("h1", "3. THIẾT KẾ THỰC NGHIỆM")
    p("h2", "3.1. Quy trình thu dữ liệu và kịch bản thử")
    p("body", "Mô tả sơ đồ mạch, danh sách tải, chu kỳ lấy mẫu, đồng bộ timestamp, thiết bị tham chiếu, thời lượng thu dữ liệu và quy trình an toàn.")
    scenario_rows = [
        ["Mã", "Kịch bản", "Tác động", "Dữ liệu phải ghi", "Chỉ tiêu"],
        ["S1", "Giám sát có tải", "Đóng tải theo kế hoạch", "V/I/P/E, timestamp", "Sai số, mất mẫu"],
        ["S2", "Điều khiển từ App", "Bật/tắt tải", "Command, statusTag, tiếp điểm", "Thành công, RTT"],
        ["S3", "Mất kết nối/timeout", "Ngắt gateway hoặc PLC", "Mã lỗi, trạng thái tải", "Fail-closed"],
        ["S4", "Quota/sa thải tải", "Tiến gần Q hoặc vượt ngưỡng", "E, Q, P trước–sau, tải bị cắt", "Cảnh báo, thời gian cắt"],
        ["S5", "Manual override/khôi phục", "Thao tác tại chỗ", "Trạng thái và log phục hồi", "An toàn, phục hồi"],
    ]
    captioned_table(
        document,
        samples,
        "Bảng 3. Kịch bản kiểm chứng mô hình",
        scenario_rows,
        "Ghi chú: Sa thải tải chỉ được báo cáo như kết quả sau khi thử trên tải thật với interlock và dừng khẩn.",
        [0.52, 1.36, 1.42, 2.28, 1.20],
        size=7.8,
    )

    p("h2", "3.2. Chỉ tiêu đánh giá")
    result_rows = [
        ["Nhóm đánh giá", "Chỉ tiêu", "Cách báo cáo", "Bằng chứng tối thiểu", "Trạng thái hiện tại"],
        ["Đo lường", "Sai số V/I/P/E; mất mẫu", "Mean, max, tỷ lệ (%)", "Log MFM384 + thiết bị tham chiếu", "Chưa có"],
        ["Điều khiển", "Thành công; RTT; timeout", "n/N, mean ± SD", "Command + statusTag + tiếp điểm", "Phần mềm một phần"],
        ["Quota/sa thải", "Cảnh báo đúng; thời gian cắt", "n/N và P trước–sau", "Log sự kiện + tải thật", "Thiết kế"],
        ["Dự báo", "MAE, RMSE, MAPE, R²", "Mean ± sample SD", "Rolling-origin + seed + log", "Có benchmark UCI"],
    ]
    captioned_table(
        document,
        samples,
        "Bảng 4. Khung báo cáo kết quả và trạng thái bằng chứng",
        result_rows,
        "Ghi chú: Bảng này thay cho hai bảng kết quả riêng trong dàn ý dài; chỉ thay trạng thái bằng số liệu sau khi kiểm chứng.",
        [1.12, 1.32, 1.40, 2.04, 1.20],
        size=7.6,
    )

    p("h1", "4. KẾT QUẢ VÀ THẢO LUẬN")
    p("h2", "4.1. Kết quả đo lường và điều khiển")
    add_planning_note(
        document,
        samples["body"],
        "Viết sau khi có log phần cứng. Báo cáo sai số đo, mất mẫu, phản hồi tiếp điểm, timeout, Quota và phục hồi. "
        "Không thêm đồ thị riêng nếu dữ liệu này có thể trình bày gọn trong Bảng 4 và phần mô tả.",
    )
    p("h2", "4.2. Kết quả dự báo phụ tải")
    p("body", "So sánh Persistence, Seasonal Naive 24 h, Seasonal Naive 168 h, Random Forest và XGBoost. Báo cáo mean ± sample SD và giới hạn kết quả trong phạm vi bộ dữ liệu UCI.")
    add_figure(
        document,
        samples["image"],
        samples["figure_caption"],
        figures_dir / "final_forecast_horizon_vi.png",
        "Hình 4. Sai số MAE và RMSE tại các chân trời h+1, h+6, h+12 và h+24; error bar biểu diễn sample SD",
        "Hai biểu đồ đường thể hiện MAE và RMSE của các mô hình tại bốn chân trời dự báo.",
        width=5.75,
    )
    p("h2", "4.3. Thảo luận và nguy cơ đối với tính hợp lệ")
    add_planning_note(
        document,
        samples["body"],
        "Thảo luận bốn điểm: giá trị đối với phòng thí nghiệm; vai trò phản hồi PLC; mức hỗ trợ của dự báo đối với Quota; "
        "và lý do dữ liệu UCI chưa đại diện cho mô hình tại Cần Thơ. Tách nguy cơ nội tại, ngoại tại, đo lường và mô hình.",
    )

    p("h1", "5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    add_planning_note(
        document,
        samples["body"],
        "Kết luận trả lời trực tiếp ba RQ và chỉ nêu số đã kiểm chứng. Hướng phát triển gồm dữ liệu MFM384 địa phương, "
        "sai số đo, độ trễ App–PLC, sa thải tải an toàn và tích hợp thời tiết có đánh giá ablation.",
    )
    p("h2", "5.1. Điều kiện trước khi viết bản hoàn chỉnh")
    p("body", "Chốt manual MFM384; chụp Hình 2; chụp Hình 3; thực hiện các kịch bản Bảng 3; điền Bảng 4; sau đó mới viết tóm tắt, kết quả, kết luận và xuất bản nộp.")

    clone_paragraph(document, samples["ack"], "Lời cảm ơn (nếu có): [BỔ SUNG SAU KHI XÁC NHẬN CƠ QUAN HỖ TRỢ/TÀI TRỢ].")
    clone_paragraph(document, samples["ref_heading"], "Tài liệu tham khảo dự kiến", keep_next=True)
    selected_references = [REFERENCES[index] for index in (0, 2, 5, 6, 7, 8, 9, 10, 11, 13, 17)]
    for number, reference in enumerate(selected_references, start=1):
        renumbered = re.sub(r"^\[\d+\]", f"[{number}]", reference)
        paragraph = clone_paragraph(document, samples["ref"], renumbered)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def validate_output(path: Path) -> None:
    document = Document(path)
    if len(document.sections) != 1:
        raise ValueError("Outline must retain one section")
    if len(document.tables) != 6:
        raise ValueError(f"Expected 2 front-matter + 4 scientific tables, got {len(document.tables)}")
    captions = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for number in range(1, 5):
        if f"Hình {number}." not in captions:
            raise ValueError(f"Missing Figure {number} caption")
        if f"Bảng {number}." not in captions:
            raise ValueError(f"Missing Table {number} caption")
    if "Hình 5." in captions or "Bảng 5." in captions:
        raise ValueError("Outline exceeds reduced 4-figure/4-table scope")


def main() -> int:
    args = parse_args()
    if not args.template.exists():
        raise FileNotFoundError(args.template)
    for name in ("final_architecture_vi.png", "final_forecast_horizon_vi.png"):
        if not (args.figures_dir / name).exists():
            raise FileNotFoundError(args.figures_dir / name)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.template, args.output)
    document = Document(args.output)
    samples = {
        "h1": document.paragraphs[10],
        "body": document.paragraphs[11],
        "h2": document.paragraphs[16],
        "h3": document.paragraphs[21],
        "table_caption": document.paragraphs[25],
        "table_note": document.paragraphs[26],
        "image": document.paragraphs[29],
        "figure_caption": document.paragraphs[30],
        "ack": document.paragraphs[36],
        "ref_heading": document.paragraphs[37],
        "ref": document.paragraphs[39],
    }
    populate_front_matter(document)
    clear_after_front_matter(document)
    build_outline(document, samples, args.figures_dir)
    document.core_properties.title = "Dàn ý bài báo HEMS theo mẫu CTUT"
    document.core_properties.subject = "Bản dàn ý rút gọn với 4 hình và 4 bảng"
    document.core_properties.comments = "Outline only; hardware and App figures remain evidence-gated placeholders."
    document.save(args.output)
    validate_output(args.output)
    print(args.output)
    print("scientific_tables=4 proposed_figures=4 embedded_images=2")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
