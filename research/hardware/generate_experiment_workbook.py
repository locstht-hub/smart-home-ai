from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


DEFAULT_OUTPUT = Path(
    "outputs/20260712-smart-home-research/HEMS_Phieu_Thu_Thap_Du_Lieu_Thuc_Nghiem.xlsx"
)

NAVY = "1F4E78"
BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "F4CCCC"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
INPUT_BLUE = "0000FF"
FORMULA_BLACK = "000000"
THIN = Side(style="thin", color="B7C9D6")


LISTS = {
    "yes_no": ["YES", "NO"],
    "yes_no_na": ["YES", "NO", "N/A"],
    "device_group": ["MFM384", "PLC", "RELAY_CONTACTOR", "LOAD", "REFERENCE_METER", "NETWORK", "SERVER", "OTHER"],
    "device_status": ["CHỜ XÁC MINH", "ĐÃ XÁC MINH", "ĐÃ HIỆU CHUẨN", "KHÔNG SỬ DỤNG"],
    "phase": ["L1", "L2", "L3", "TOTAL", "N/A"],
    "load_state": ["ON", "OFF", "UNKNOWN", "TRANSITION"],
    "quality": ["OK", "MISSING", "STALE", "OUT_OF_RANGE", "COMM_ERROR", "MANUAL_ENTRY"],
    "source": ["plc-s7-1200", "mfm384-direct", "reference-meter", "manual", "mock", "mock-fallback"],
    "session_status": ["PLANNED", "RUNNING", "COMPLETED", "REJECTED", "ABORTED"],
    "interface": ["MOBILE_APP", "WEB_DASHBOARD", "LOCAL_PANEL", "SCRIPT", "OTHER"],
    "requested_state": ["ON", "OFF"],
    "permission": ["ALLOWED", "DENIED", "NOT_CHECKED"],
    "quota_state": ["NORMAL", "NEAR", "EXCEEDED", "NOT_APPLICABLE"],
    "interlock": ["READY", "BLOCKED", "TRIPPED", "NOT_CHECKED"],
    "command_result": ["VERIFIED", "DENIED", "ERROR", "TIMEOUT", "CANCELLED"],
    "network_state": ["ONLINE", "OFFLINE", "DEGRADED", "UNKNOWN"],
    "alert_channel": ["APP", "WEB", "TELEGRAM", "MULTI_CHANNEL", "NONE"],
    "incident_type": ["NETWORK_LOSS", "TIMEOUT", "STALE_DATA", "MODBUS_ERROR", "INTERLOCK", "EMERGENCY_STOP", "MANUAL_OVERRIDE", "OVERCURRENT", "OTHER"],
    "safe_state": ["LOAD_OFF", "HOLD_LAST_SAFE", "COMMAND_REJECTED", "ALARM_ONLY", "OTHER"],
    "evidence_type": ["RAW_CSV", "RAW_JSON", "PHOTO", "SCREENSHOT", "VIDEO", "TIA_EXPORT", "MANUAL_PDF", "CALIBRATION_CERT", "OTHER"],
    "scenario_status": ["CHƯA THỰC HIỆN", "ĐANG THỰC HIỆN", "ĐỦ DỮ LIỆU", "KHÔNG ĐẠT", "KHÔNG ĐƯỢC PHÉP"],
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the HEMS laboratory data-collection workbook.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def style_header(ws, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.row_dimensions[1].height = 34
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def setup_input_sheet(ws, headers: list[str], widths: dict[str, float], max_rows: int) -> None:
    ws.append(headers)
    style_header(ws, headers)
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = widths.get(header, max(12, min(26, len(header) + 3)))
        ws.cell(1, index).comment = Comment(f"Trường dữ liệu: {header}. Xem TU_DIEN_DU_LIEU để biết cách nhập.", "Codex")
    for row in range(2, max_rows + 1):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Arial", size=9, color=INPUT_BLUE)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws.sheet_format.defaultRowHeight = 18


def add_validation(ws, column: str, values_key: str, max_rows: int, list_ranges: dict[str, str]) -> None:
    validation = DataValidation(type="list", formula1=f"={list_ranges[values_key]}", allow_blank=True)
    validation.error = "Chọn một giá trị trong danh sách."
    validation.errorTitle = "Giá trị không hợp lệ"
    validation.prompt = "Chọn giá trị chuẩn để dữ liệu có thể phân tích tự động."
    validation.promptTitle = "Danh sách chuẩn"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(f"{column}2:{column}{max_rows}")


def add_formula_column(ws, column: str, max_rows: int, formula_builder, number_format: str = "0.00") -> None:
    for row in range(2, max_rows + 1):
        cell = ws[f"{column}{row}"]
        cell.value = formula_builder(row)
        cell.font = Font(name="Arial", size=9, color=FORMULA_BLACK)
        cell.number_format = number_format


def create_lists(wb: Workbook) -> dict[str, str]:
    ws = wb.create_sheet("DANH_MUC")
    ranges: dict[str, str] = {}
    for col, (key, values) in enumerate(LISTS.items(), 1):
        ws.cell(1, col, key)
        for row, value in enumerate(values, 2):
            ws.cell(row, col, value)
        letter = get_column_letter(col)
        wb.defined_names.add(
            DefinedName(
                key,
                attr_text=f"'DANH_MUC'!${letter}$2:${letter}${len(values) + 1}",
            )
        )
        ranges[key] = key
    ws.sheet_state = "hidden"
    return ranges


def create_guide(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "HUONG_DAN"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 105
    ws.merge_cells("A1:B1")
    ws["A1"] = "PHIẾU THU THẬP DỮ LIỆU THỰC NGHIỆM HEMS"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    rows = [
        ("Phạm vi", "Mô hình phụ tải điện quy mô nhỏ đặt tại Phòng thí nghiệm Điện công nghiệp; tích hợp MFM384, PLC Siemens S7-1200, cơ cấu đóng cắt, Web/App, Quota và dự báo 24 giờ."),
        ("Mục đích", "Ghi bằng chứng thô để cập nhật bài báo theo cấu trúc 4 hình + 4 bảng. Workbook không tự biến kế hoạch thành kết quả khoa học."),
        ("Bước 1", "Điền THIET_BI trước khi thử: đúng model, serial, manual, firmware, cấu hình Modbus, định mức và tình trạng hiệu chuẩn."),
        ("Bước 2", "Mỗi lần thử tạo một session_id duy nhất trong PHIEN_THU, ví dụ LAB-20260810-01. Đồng bộ thời gian cho PLC, máy chủ, MFM384 và thiết bị tham chiếu."),
        ("Bước 3", "Nhập hoặc dán dữ liệu đo vào DU_LIEU_MFM384. Giữ timestamp, timezone, đơn vị và đường dẫn file raw. Không làm tròn hoặc sửa số thô."),
        ("Bước 4", "Ghi từng lệnh ON/OFF trong LENH_PLC, gồm trạng thái trước, phản hồi độc lập, tiếp điểm, công suất trước/sau, kết quả và RTT."),
        ("Bước 5", "Ghi sự kiện Quota/cảnh báo trong QUOTA_CANH_BAO. Cảnh báo và khuyến nghị không đồng nghĩa tối ưu hóa. Không đánh dấu sa thải tải đã thực hiện nếu chưa thử an toàn."),
        ("Bước 6", "Ghi timeout, mất mạng, interlock, manual override và dừng khẩn trong SU_CO_AN_TOAN. Trạng thái thực tế phải dựa trên PLC/tiếp điểm, không dựa vào thông báo giao diện."),
        ("Bước 7", "Đăng ký ảnh, video, CSV, JSON, TIA export và manual trong NHAT_KY_TEP. Che token, mật khẩu, IP công khai và dữ liệu cá nhân."),
        ("Dữ liệu tối thiểu gửi lại", "Workbook đã điền; CSV/JSON raw; manual đúng model MFM384; ảnh testbed thật; ảnh chụp App đã ẩn danh; TIA/PLC tag export nếu được phép; thông tin thiết bị tham chiếu."),
        ("Quy tắc nguồn", "Chỉ dữ liệu source=plc-s7-1200 hoặc mfm384-direct và có effective mode thật mới được xem xét làm bằng chứng phần cứng. Mock/fallback phải giữ để truy lỗi nhưng không đưa vào thống kê luận văn."),
        ("An toàn", "Không tự động đóng/cắt tải công suất lớn bằng workbook hoặc script. Thử nghiệm phải có người chuyên môn, bảo vệ mạch, interlock, manual override và dừng khẩn."),
        ("Khi gửi lại cho Codex", "Đính kèm file này cùng các raw log. Tôi sẽ kiểm tra thiếu dữ liệu, làm sạch có truy vết, tính thống kê và chỉ cập nhật các claim được bằng chứng cho phép."),
    ]
    for row, (label, text) in enumerate(rows, 3):
        ws.cell(row, 1, label)
        ws.cell(row, 2, text)
        ws.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=NAVY)
        ws.cell(row, 2).font = Font(name="Arial", size=10)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE if row % 2 else GRAY)
        ws.cell(row, 2).fill = PatternFill("solid", fgColor="F8FBFD" if row % 2 else "F2F2F2")
        for col in range(1, 3):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row, col).border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.row_dimensions[row].height = 48
    ws["A18"] = "Lưu ý bắt buộc"
    ws["B18"] = "Không điền số ước lượng, không tạo ảnh minh họa giả và không công bố sa thải tải tự động khi chưa có log tải thật."
    for cell in ws[18]:
        cell.fill = PatternFill("solid", fgColor=RED)
        cell.font = Font(name="Arial", size=10, bold=True, color="9C0006")
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def create_paper_map(wb: Workbook) -> None:
    ws = wb.create_sheet("ANH_XA_BAI_BAO")
    headers = ["Mã", "Loại", "Nội dung", "Nguồn dữ liệu", "Điều kiện sử dụng", "Trạng thái hiện tại"]
    setup_input_sheet(ws, headers, {"Mã": 10, "Loại": 12, "Nội dung": 42, "Nguồn dữ liệu": 38, "Điều kiện sử dụng": 55, "Trạng thái hiện tại": 24}, 20)
    rows = [
        ["H1", "Hình", "Kiến trúc MFM384–PLC–tải–Web/App–dự báo", "Sơ đồ thiết kế đã xác minh với cấu hình thật", "Ghi rõ ranh giới thiết kế và phần cứng đã kiểm chứng", "CÓ THỂ CẬP NHẬT"],
        ["H2", "Hình", "Ảnh testbed phần cứng thật", "NHAT_KY_TEP + ảnh gốc", "Có toàn cảnh, MFM384/RS-485, PLC/I/O, relay/contactor/tải và chú thích", "CHỜ ẢNH THẬT"],
        ["H3", "Hình", "V/I/P/E, trạng thái tải và sự kiện điều khiển", "DU_LIEU_MFM384 + LENH_PLC", "Timestamp đồng bộ, source thật, phản hồi độc lập và đủ điều kiện tải", "CHỜ LOG THẬT"],
        ["H4", "Hình", "MAE/RMSE theo chân trời h+1, h+6, h+12, h+24", "canonical_results.json", "Chỉ diễn giải trong phạm vi UCI", "ĐÃ CÓ"],
        ["B1", "Bảng", "So sánh nghiên cứu liên quan", "Tài liệu đã xác minh", "Không dùng nguồn giả hoặc nguồn không liên quan", "CẦN CHỐT TRÍCH DẪN"],
        ["B2", "Bảng", "Thiết bị và cấu hình thử nghiệm", "THIET_BI + PHIEN_THU", "Manual, serial, firmware, định mức và thiết bị tham chiếu phải rõ", "CHỜ ĐIỀN"],
        ["B3", "Bảng", "Kịch bản và kết quả đo–điều khiển", "BANG_KICH_BAN + các log thực nghiệm", "Chỉ ghi kết quả sau khi đủ số lần và bằng chứng", "CHỜ THỬ"],
        ["B4", "Bảng", "Kết quả benchmark dự báo", "canonical_results.json", "RF/XGBoost gần tương đương; không suy ra độ chính xác địa phương", "ĐÃ CÓ"],
    ]
    for row_index, values in enumerate(rows, 2):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9, color=FORMULA_BLACK)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = "A1:F9"


def create_device_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("THIET_BI")
    headers = ["device_id", "nhom_thiet_bi", "hang_san_xuat", "model", "serial", "firmware", "manual_ten_phien_ban", "manual_da_xac_minh", "dien_ap_dinh_muc_v", "dong_dinh_muc_a", "cong_suat_dinh_muc_w", "cap_chinh_xac", "giao_tiep", "modbus_slave_id", "baud_rate", "parity", "stop_bits", "byte_order", "he_so_ty_le", "plc_address_tag", "ngay_lap_dat", "ngay_hieu_chuan", "thiet_bi_tham_chieu_chung_nhan", "trang_thai", "ghi_chu"]
    setup_input_sheet(ws, headers, {"device_id": 14, "nhom_thiet_bi": 20, "manual_ten_phien_ban": 28, "he_so_ty_le": 24, "plc_address_tag": 22, "thiet_bi_tham_chieu_chung_nhan": 32, "ghi_chu": 36}, 300)
    add_validation(ws, "B", "device_group", 300, ranges)
    add_validation(ws, "H", "yes_no", 300, ranges)
    add_validation(ws, "X", "device_status", 300, ranges)
    defaults = [
        ["MFM-01", "MFM384", "", "MFM384", "", "", "", "NO", "", "", "", "", "RS-485/Modbus RTU", "", "", "", "", "", "", "", "", "", "", "CHỜ XÁC MINH", "Điền đúng manual và model thực tế trước khi đo."],
        ["PLC-01", "PLC", "Siemens", "S7-1200", "", "", "", "NO", "", "", "", "", "S7/Modbus", "", "", "", "", "", "", "", "", "", "", "CHỜ XÁC MINH", "Điền CPU/module, firmware và phiên bản chương trình TIA."],
        ["SW-01", "RELAY_CONTACTOR", "", "", "", "", "", "NO", "", "", "", "", "Digital I/O", "", "", "", "", "", "", "", "", "", "", "CHỜ XÁC MINH", "Ghi định mức tiếp điểm, cuộn hút và tiếp điểm phản hồi."],
        ["LOAD-01", "LOAD", "", "", "", "", "", "NO", "", "", "", "", "Mạch lực", "", "", "", "", "", "", "", "", "", "", "CHỜ XÁC MINH", "Ghi tên tải, công suất định mức và mức ưu tiên cắt."],
        ["REF-01", "REFERENCE_METER", "", "", "", "", "", "NO", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "CHỜ XÁC MINH", "Dùng để đối chiếu sai số MFM384 nếu có."],
    ]
    for row_index, values in enumerate(defaults, 2):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)


def create_session_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("PHIEN_THU")
    headers = ["session_id", "ngay_thu", "bat_dau", "ket_thuc", "thoi_luong_phut", "nguoi_thuc_hien", "nguoi_giam_sat", "dia_diem", "ma_kich_ban", "muc_tieu", "cau_hinh_testbed", "plc_cpu_module", "plc_program_version", "app_version_commit", "backend_version_commit", "mfm_config_version", "chu_ky_lay_mau_s", "phuong_phap_dong_bo_gio", "thiet_bi_tham_chieu", "nhiet_do_moi_truong_c", "kiem_tra_an_toan", "da_thu_dung_khan", "da_thu_manual_override", "che_do_mang", "thu_muc_du_lieu", "trang_thai_phien", "ghi_chu"]
    setup_input_sheet(ws, headers, {"session_id": 20, "muc_tieu": 34, "cau_hinh_testbed": 40, "phuong_phap_dong_bo_gio": 28, "thu_muc_du_lieu": 35, "ghi_chu": 38}, 300)
    add_validation(ws, "U", "yes_no", 300, ranges)
    add_validation(ws, "V", "yes_no_na", 300, ranges)
    add_validation(ws, "W", "yes_no_na", 300, ranges)
    add_validation(ws, "Z", "session_status", 300, ranges)
    add_formula_column(ws, "E", 300, lambda row: f'=IF(OR(C{row}="",D{row}=""),"",ROUND((D{row}-C{row})*1440,1))', "0.0")
    for column in ("B", "C", "D"):
        for row in range(2, 301):
            ws[f"{column}{row}"].number_format = "dd/mm/yyyy hh:mm:ss.000" if column != "B" else "dd/mm/yyyy"


def create_mfm_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("DU_LIEU_MFM384")
    headers = ["record_id", "session_id", "timestamp_local", "timezone", "mfm_device_id", "phase", "voltage_v", "current_a", "active_power_kw", "reactive_power_kvar", "apparent_power_kva", "power_factor", "frequency_hz", "energy_kwh", "load_id", "load_state", "contact_feedback", "modbus_ok", "data_quality", "source_system", "effective_mode", "sampling_interval_s", "raw_log_file", "raw_register_snapshot", "ghi_chu"]
    setup_input_sheet(ws, headers, {"record_id": 20, "session_id": 20, "timestamp_local": 24, "timezone": 18, "raw_log_file": 38, "raw_register_snapshot": 38, "ghi_chu": 36}, 10000)
    add_validation(ws, "F", "phase", 10000, ranges)
    add_validation(ws, "P", "load_state", 10000, ranges)
    add_validation(ws, "Q", "load_state", 10000, ranges)
    add_validation(ws, "R", "yes_no", 10000, ranges)
    add_validation(ws, "S", "quality", 10000, ranges)
    add_validation(ws, "T", "source", 10000, ranges)
    for row in range(2, 10001):
        ws[f"C{row}"].number_format = "dd/mm/yyyy hh:mm:ss.000"
        for column in ("G", "H", "I", "J", "K", "L", "M", "N", "V"):
            ws[f"{column}{row}"].number_format = "0.000"
    ws.conditional_formatting.add("S2:S10000", FormulaRule(formula=['AND($S2<>"",$S2<>"OK")'], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add("T2:T10000", FormulaRule(formula=['OR($T2="mock",$T2="mock-fallback")'], fill=PatternFill("solid", fgColor=YELLOW)))


def create_api_latency_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("DO_TRE_API")
    headers = ["run_id", "trial_id", "network_label", "endpoint", "started_at_utc", "finished_at_utc", "latency_ms", "http_status", "accepted", "source", "effective_mode", "voltage_v", "current_a", "power_kw", "energy_kwh", "reading_timestamp", "error"]
    setup_input_sheet(ws, headers, {"run_id": 30, "network_label": 18, "endpoint": 24, "started_at_utc": 28, "finished_at_utc": 28, "reading_timestamp": 28, "error": 40}, 3000)
    add_validation(ws, "I", "yes_no", 3000, ranges)
    add_validation(ws, "J", "source", 3000, ranges)
    ws["A3002"] = "Ghi chú"
    ws["B3002"] = "Có thể dán trực tiếp CSV do research/hardware/collect_hardware_trials.py tạo ra; chỉ accepted=YES, source=plc-s7-1200 và effective_mode=plc-real mới được xét là dữ liệu phần cứng."


def create_command_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("LENH_PLC")
    headers = ["command_id", "session_id", "timestamp_send", "timezone", "source_interface", "account_id_anon", "home_id_anon", "room_id", "device_id", "requested_state", "permission_result", "quota_state", "interlock_state", "plc_command_tag", "status_before", "ack_timestamp", "feedback_timestamp", "actual_state", "contact_feedback", "power_before_kw", "power_after_kw", "result", "error_code", "rtt_ms", "network_state", "manual_override", "operator_observation", "raw_log_file", "ghi_chu"]
    setup_input_sheet(ws, headers, {"command_id": 22, "session_id": 20, "timestamp_send": 24, "account_id_anon": 20, "plc_command_tag": 22, "operator_observation": 34, "raw_log_file": 36, "ghi_chu": 36}, 2000)
    add_validation(ws, "E", "interface", 2000, ranges)
    add_validation(ws, "J", "requested_state", 2000, ranges)
    add_validation(ws, "K", "permission", 2000, ranges)
    add_validation(ws, "L", "quota_state", 2000, ranges)
    add_validation(ws, "M", "interlock", 2000, ranges)
    add_validation(ws, "O", "load_state", 2000, ranges)
    add_validation(ws, "R", "load_state", 2000, ranges)
    add_validation(ws, "S", "load_state", 2000, ranges)
    add_validation(ws, "V", "command_result", 2000, ranges)
    add_validation(ws, "Y", "network_state", 2000, ranges)
    add_validation(ws, "Z", "yes_no", 2000, ranges)
    add_formula_column(ws, "X", 2000, lambda row: f'=IF(OR(C{row}="",Q{row}=""),"",ROUND((Q{row}-C{row})*86400000,0))', "0")
    for column in ("C", "P", "Q"):
        for row in range(2, 2001):
            ws[f"{column}{row}"].number_format = "dd/mm/yyyy hh:mm:ss.000"
    ws.conditional_formatting.add("V2:V2000", FormulaRule(formula=['$V2="VERIFIED"'], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add("V2:V2000", FormulaRule(formula=['OR($V2="ERROR",$V2="TIMEOUT",$V2="DENIED")'], fill=PatternFill("solid", fgColor=RED)))


def create_quota_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("QUOTA_CANH_BAO")
    headers = ["event_id", "session_id", "timestamp_event", "timezone", "account_id_anon", "period_month", "energy_month_kwh", "quota_kwh", "usage_percent", "forecast_24h_kwh", "projected_energy_kwh", "projected_usage_percent", "threshold_percent", "alert_triggered", "alert_channel", "alert_sent_timestamp", "alert_received_timestamp", "alert_latency_ms", "recommendation_type", "recommendation_text", "user_acknowledged", "load_shedding_requested", "load_shedding_executed", "operator_approval", "related_command_id", "evidence_file", "ghi_chu"]
    setup_input_sheet(ws, headers, {"event_id": 20, "session_id": 20, "timestamp_event": 24, "recommendation_text": 48, "related_command_id": 22, "evidence_file": 36, "ghi_chu": 36}, 1000)
    for column in ("N", "U", "V", "W", "X"):
        add_validation(ws, column, "yes_no", 1000, ranges)
    add_validation(ws, "O", "alert_channel", 1000, ranges)
    add_formula_column(ws, "I", 1000, lambda row: f'=IF(OR(G{row}="",H{row}="",H{row}=0),"",G{row}/H{row})', "0.0%")
    add_formula_column(ws, "K", 1000, lambda row: f'=IF(OR(G{row}="",J{row}=""),"",G{row}+J{row})', "0.000")
    add_formula_column(ws, "L", 1000, lambda row: f'=IF(OR(K{row}="",H{row}="",H{row}=0),"",K{row}/H{row})', "0.0%")
    add_formula_column(ws, "R", 1000, lambda row: f'=IF(OR(P{row}="",Q{row}=""),"",ROUND((Q{row}-P{row})*86400000,0))', "0")
    for column in ("C", "P", "Q"):
        for row in range(2, 1001):
            ws[f"{column}{row}"].number_format = "dd/mm/yyyy hh:mm:ss.000"
    ws.conditional_formatting.add("W2:W1000", FormulaRule(formula=['$W2="YES"'], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add("I2:I1000", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor=RED)))


def create_incident_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("SU_CO_AN_TOAN")
    headers = ["event_id", "session_id", "start_timestamp", "end_timestamp", "recovery_time_ms", "event_type", "trigger_method", "network_state", "expected_safe_state", "actual_safe_state", "fail_closed", "interlock_activated", "emergency_stop_used", "manual_override_used", "recovered", "operator_action", "root_cause", "affected_device", "evidence_file", "ghi_chu"]
    setup_input_sheet(ws, headers, {"event_id": 20, "session_id": 20, "start_timestamp": 24, "end_timestamp": 24, "trigger_method": 30, "operator_action": 38, "root_cause": 38, "evidence_file": 36, "ghi_chu": 36}, 500)
    add_validation(ws, "F", "incident_type", 500, ranges)
    add_validation(ws, "H", "network_state", 500, ranges)
    add_validation(ws, "I", "safe_state", 500, ranges)
    add_validation(ws, "J", "safe_state", 500, ranges)
    for column in ("K", "L", "M", "N", "O"):
        add_validation(ws, column, "yes_no", 500, ranges)
    add_formula_column(ws, "E", 500, lambda row: f'=IF(OR(C{row}="",D{row}=""),"",ROUND((D{row}-C{row})*86400000,0))', "0")
    for column in ("C", "D"):
        for row in range(2, 501):
            ws[f"{column}{row}"].number_format = "dd/mm/yyyy hh:mm:ss.000"
    ws.conditional_formatting.add("K2:K500", FormulaRule(formula=['$K2="NO"'], fill=PatternFill("solid", fgColor=RED)))


def create_scenario_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("BANG_KICH_BAN")
    headers = ["ma_kich_ban", "nhom", "muc_tieu", "thiet_lap", "tac_dong", "du_lieu_bat_buoc", "so_lan_khuyen_nghi", "tieu_chi_dat", "trang_thai", "vi_tri_ket_qua", "ket_luan_duoc_phep"]
    setup_input_sheet(ws, headers, {"ma_kich_ban": 14, "nhom": 20, "muc_tieu": 40, "thiet_lap": 45, "tac_dong": 42, "du_lieu_bat_buoc": 60, "so_lan_khuyen_nghi": 26, "tieu_chi_dat": 55, "trang_thai": 22, "vi_tri_ket_qua": 32, "ket_luan_duoc_phep": 55}, 100)
    add_validation(ws, "I", "scenario_status", 100, ranges)
    rows = [
        ["S1", "Đo lường", "Kiểm tra V/I/P/Q/PF/E theo các mức tải", "Không tải, tải thấp, trung bình, cao; có thiết bị tham chiếu nếu có", "Thay đổi tải trong giới hạn an toàn", "Timestamp; V/I/P/Q/PF/E; tải hoạt động; chất lượng dữ liệu; số tham chiếu", "Theo lịch 30 ngày; mỗi mức tải tối thiểu 3 phiên", "Không mất mẫu nghiêm trọng; đơn vị/scale đúng; sai số chỉ tính khi có thiết bị tham chiếu", "CHƯA THỰC HIỆN", "DU_LIEU_MFM384", "Chỉ mô tả đo thật khi source và log được chấp nhận"],
        ["S2", "Điều khiển", "Xác nhận lệnh App/Web bằng phản hồi độc lập", "PLC, relay/contactor, tiếp điểm và tải thật", "ON/OFF từ App và Web", "Command; status trước/sau; tiếp điểm; P trước/sau; RTT; lỗi", "Tối thiểu 20 lần ON và 20 lần OFF cho mỗi tải", "Chỉ VERIFIED khi actual_state và contact_feedback khớp; ghi mọi lỗi/timeout", "CHƯA THỰC HIỆN", "LENH_PLC", "Được báo tỷ lệ thành công và RTT khi đủ trial thật"],
        ["S3", "Fail-safe", "Kiểm tra mất mạng, timeout và dữ liệu stale", "Tình huống lỗi được giám sát; tải ở trạng thái an toàn", "Ngắt kết nối hoặc tạo timeout có kiểm soát", "Mốc thời gian; lỗi; trạng thái yêu cầu/thực; interlock; phục hồi", "Tối thiểu 10 lần cho mỗi loại lỗi", "Fail-closed; không báo thành công giả; phục hồi có kiểm soát", "CHƯA THỰC HIỆN", "SU_CO_AN_TOAN", "Được mô tả fail-safe khi có log và quan sát tải thật"],
        ["S4", "Quota/cảnh báo", "Kiểm tra cảnh báo gần/chạm Quota", "Quota thử nghiệm có ghi rõ; không nhầm kW và kWh", "Tạo các mức dưới, gần và vượt ngưỡng", "E tháng; Quota; dự báo 24 h; thời điểm gửi/nhận; nội dung cảnh báo", "Tối thiểu 5 lần tại mỗi ngưỡng đã chốt", "Cảnh báo đúng ngưỡng, đúng tài khoản, không gửi lặp ngoài quy tắc cooldown", "CHƯA THỰC HIỆN", "QUOTA_CANH_BAO", "Chỉ gọi là cảnh báo/khuyến nghị; không gọi là tối ưu hóa"],
        ["S5", "An toàn", "Kiểm tra manual override, dừng khẩn và khôi phục", "Có người giám sát và quy trình phục hồi", "Override hoặc dừng khẩn có kiểm soát", "Trạng thái trước/sau; tiếp điểm; thao tác; thời gian phục hồi; ảnh/video", "Tối thiểu 10 lần cho mỗi chức năng được phép", "Ưu tiên thao tác tại chỗ; hệ thống không tự ghi đè trạng thái an toàn", "CHƯA THỰC HIỆN", "SU_CO_AN_TOAN", "Chỉ kết luận trong phạm vi testbed"],
        ["S6", "Sa thải tải", "Kiểm chứng cắt tải phi thiết yếu có điều kiện", "Chỉ khi interlock, manual override, dừng khẩn, thứ tự ưu tiên và phê duyệt an toàn đầy đủ", "Vượt ngưỡng đã chốt và cho phép sa thải", "P/E trước/sau; tải bị cắt; nguyên nhân; tiếp điểm; interlock; phục hồi", "0 cho đến khi được phê duyệt; sau đó chốt số lần với người hướng dẫn", "Đúng tải ưu tiên, có hysteresis/thời gian giữ, không cắt tải thiết yếu, phục hồi an toàn", "KHÔNG ĐƯỢC PHÉP", "LENH_PLC + SU_CO_AN_TOAN", "Chỉ là thiết kế cho đến khi đủ điều kiện và log thật"],
    ]
    for row_index, values in enumerate(rows, 2):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9, color=FORMULA_BLACK)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row_index].height = 72


def create_evidence_sheet(wb: Workbook, ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("NHAT_KY_TEP")
    headers = ["evidence_id", "session_id", "timestamp", "evidence_type", "file_name", "relative_path", "sha256", "mo_ta", "record_ids_lien_quan", "da_an_danh", "duoc_chap_nhan", "ly_do_loai", "nguoi_xac_minh", "ghi_chu"]
    setup_input_sheet(ws, headers, {"evidence_id": 20, "session_id": 20, "timestamp": 24, "file_name": 35, "relative_path": 55, "sha256": 66, "mo_ta": 45, "record_ids_lien_quan": 35, "ly_do_loai": 38, "ghi_chu": 36}, 1000)
    add_validation(ws, "D", "evidence_type", 1000, ranges)
    add_validation(ws, "J", "yes_no", 1000, ranges)
    add_validation(ws, "K", "yes_no", 1000, ranges)
    for row in range(2, 1001):
        ws[f"C{row}"].number_format = "dd/mm/yyyy hh:mm:ss.000"


def create_dictionary(wb: Workbook) -> None:
    ws = wb.create_sheet("TU_DIEN_DU_LIEU")
    headers = ["Sheet", "Nhóm trường", "Trường tiêu biểu", "Bắt buộc", "Đơn vị/định dạng", "Mô tả và mục đích khoa học"]
    setup_input_sheet(ws, headers, {"Sheet": 24, "Nhóm trường": 24, "Trường tiêu biểu": 45, "Bắt buộc": 12, "Đơn vị/định dạng": 28, "Mô tả và mục đích khoa học": 85}, 80)
    rows = [
        ["THIET_BI", "Định danh", "device_id, model, serial, firmware", "YES", "Chuỗi", "Truy vết đúng thiết bị; không được thay bằng tên chung chung."],
        ["THIET_BI", "Manual/cấu hình", "manual_da_xac_minh, Modbus, byte_order, he_so_ty_le", "YES", "Theo manual đúng model", "Chứng minh số đo được giải mã đúng thanh ghi và hệ số tỉ lệ."],
        ["PHIEN_THU", "Phiên", "session_id, bat_dau, ket_thuc, ma_kich_ban", "YES", "Timestamp local + timezone", "Liên kết mọi log, ảnh và điều kiện thử."],
        ["PHIEN_THU", "An toàn", "kiem_tra_an_toan, dung_khan, manual_override", "YES", "YES/NO/N/A", "Ngăn mô tả thử tải khi điều kiện an toàn chưa được xác nhận."],
        ["DU_LIEU_MFM384", "Đại lượng điện", "voltage_v, current_a, active_power_kw, energy_kwh", "YES", "V, A, kW, kWh", "Dữ liệu chính cho Hình 3, kiểm tra tải và Quota."],
        ["DU_LIEU_MFM384", "Chất lượng", "modbus_ok, data_quality, source_system, effective_mode", "YES", "Danh sách chuẩn", "Loại mock/fallback và dữ liệu lỗi khỏi kết quả khoa học."],
        ["LENH_PLC", "Lệnh/phản hồi", "requested_state, actual_state, contact_feedback, result", "YES", "Trạng thái", "Phân biệt lệnh gửi với trạng thái tải thật."],
        ["LENH_PLC", "Thời gian", "timestamp_send, feedback_timestamp, rtt_ms", "YES", "ms", "Đo độ trễ App/Web–PLC; không dùng thời gian mô phỏng."],
        ["LENH_PLC", "Tác động điện", "power_before_kw, power_after_kw", "Khuyến nghị", "kW", "Hỗ trợ xác nhận tải có thay đổi sau lệnh."],
        ["QUOTA_CANH_BAO", "Quota", "energy_month_kwh, quota_kwh, usage_percent", "YES", "kWh, %", "Kiểm tra đúng ngưỡng năng lượng theo tháng."],
        ["QUOTA_CANH_BAO", "Khuyến nghị", "recommendation_type, recommendation_text", "YES", "Chuỗi", "Chỉ là hỗ trợ quyết định nếu chưa có thuật toán tối ưu được đánh giá."],
        ["QUOTA_CANH_BAO", "Sa thải tải", "requested, executed, approval", "Nếu có", "YES/NO", "Không được đánh dấu thực hiện nếu chưa có phê duyệt và log tải thật."],
        ["SU_CO_AN_TOAN", "Fail-safe", "expected_safe_state, actual_safe_state, fail_closed", "YES", "Danh sách chuẩn", "Chứng minh hệ thống từ chối an toàn khi lỗi."],
        ["NHAT_KY_TEP", "Provenance", "file_name, path, sha256, accepted", "YES", "Đường dẫn + SHA-256", "Liên kết số liệu trong bài với bằng chứng gốc."],
    ]
    for row_index, values in enumerate(rows, 2):
        for col, value in enumerate(values, 1):
            ws.cell(row_index, col, value)
            ws.cell(row_index, col).font = Font(name="Arial", size=9, color=FORMULA_BLACK)
            ws.cell(row_index, col).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row_index].height = 44


def create_summary(wb: Workbook) -> None:
    ws = wb.create_sheet("TONG_HOP")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 70
    ws.merge_cells("A1:C1")
    ws["A1"] = "TỔNG HỢP ĐỘ ĐẦY ĐỦ DỮ LIỆU"
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].font = Font(name="Arial", size=15, bold=True, color=WHITE)
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = ["Chỉ tiêu", "Giá trị", "Diễn giải"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(3, col, value)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")
    metrics = [
        ("Số thiết bị đã khai báo", '=COUNTIF(THIET_BI!A2:A300,"<>")', "Cần có MFM384, PLC, cơ cấu đóng cắt, tải và thiết bị tham chiếu nếu đánh giá sai số."),
        ("Thiết bị đã xác minh manual", '=COUNTIF(THIET_BI!H2:H300,"YES")', "Manual phải đúng model/phiên bản đang lắp."),
        ("Số phiên thử hoàn thành", '=COUNTIF(PHIEN_THU!Z2:Z300,"COMPLETED")', "Phiên bị REJECTED/ABORTED không dùng làm kết quả chính."),
        ("Số bản ghi MFM384", '=COUNTIF(DU_LIEU_MFM384!A2:A10000,"<>")', "Chỉ dữ liệu có timestamp, source và quality hợp lệ mới được phân tích."),
        ("Bản ghi MFM384 chất lượng OK", '=COUNTIF(DU_LIEU_MFM384!S2:S10000,"OK")', "So sánh với tổng số bản ghi để phát hiện dữ liệu thiếu/lỗi."),
        ("Bản ghi nguồn mock/fallback", '=COUNTIF(DU_LIEU_MFM384!T2:T10000,"mock")+COUNTIF(DU_LIEU_MFM384!T2:T10000,"mock-fallback")', "Không đưa vào kết quả phần cứng."),
        ("Tổng lệnh PLC", '=COUNTIF(LENH_PLC!A2:A2000,"<>")', "Bao gồm VERIFIED, DENIED, ERROR và TIMEOUT."),
        ("Lệnh VERIFIED", '=COUNTIF(LENH_PLC!V2:V2000,"VERIFIED")', "Chỉ khi phản hồi độc lập khớp yêu cầu."),
        ("Tỷ lệ VERIFIED", '=IF(B10=0,"",B11/B10)', "Không gọi là tỷ lệ đóng cắt tải thật nếu thiếu tiếp điểm/công suất."),
        ("RTT trung bình (ms)", '=IF(COUNT(LENH_PLC!X2:X2000)=0,"",AVERAGE(LENH_PLC!X2:X2000))', "Chỉ dùng RTT từ các phiên phần cứng thật được chấp nhận."),
        ("RTT P95 (ms)", '=IF(COUNT(LENH_PLC!X2:X2000)=0,"",PERCENTILE.INC(LENH_PLC!X2:X2000,0.95))', "Báo kèm n, mean, median, SD, min và max khi viết bài."),
        ("Số sự kiện Quota", '=COUNTIF(QUOTA_CANH_BAO!A2:A1000,"<>")', "Cần có tình huống dưới, gần và vượt ngưỡng."),
        ("Số cảnh báo đã kích hoạt", '=COUNTIF(QUOTA_CANH_BAO!N2:N1000,"YES")', "Phải kiểm tra đúng tài khoản và đúng ngưỡng kWh."),
        ("Số lần sa thải tải được ghi", '=COUNTIF(QUOTA_CANH_BAO!W2:W1000,"YES")', "Giá trị lớn hơn 0 cần phê duyệt, log tải thật và hồ sơ an toàn."),
        ("Sự cố/fail-safe đã ghi", '=COUNTIF(SU_CO_AN_TOAN!A2:A500,"<>")', "Bao gồm mất mạng, timeout, stale data, interlock và dừng khẩn."),
        ("Sự cố đạt fail-closed", '=COUNTIF(SU_CO_AN_TOAN!K2:K500,"YES")', "Cần đối chiếu expected_safe_state với actual_safe_state."),
        ("Tệp bằng chứng được chấp nhận", '=COUNTIF(NHAT_KY_TEP!K2:K1000,"YES")', "Ưu tiên raw log, ảnh gốc, manual và export cấu hình."),
    ]
    for row, (label, formula, note) in enumerate(metrics, 4):
        ws.cell(row, 1, label)
        ws.cell(row, 2, formula)
        ws.cell(row, 3, note)
        ws.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color=NAVY)
        ws.cell(row, 2).font = Font(name="Arial", size=10, color=FORMULA_BLACK)
        ws.cell(row, 3).font = Font(name="Arial", size=9)
        ws.cell(row, 2).number_format = "0.0%" if label == "Tỷ lệ VERIFIED" else "0.00"
        for col in range(1, 4):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row, col).border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.row_dimensions[row].height = 34
    warning_row = 23
    ws.merge_cells(start_row=warning_row, start_column=1, end_row=warning_row, end_column=3)
    ws.cell(warning_row, 1, "Bảng tổng hợp chỉ hỗ trợ kiểm tra độ đầy đủ. Quyền tuyên bố khoa học vẫn phải lấy từ canonical_results.json và raw log đã được chấp nhận.")
    ws.cell(warning_row, 1).fill = PatternFill("solid", fgColor=YELLOW)
    ws.cell(warning_row, 1).font = Font(name="Arial", bold=True, color="7F6000")
    ws.cell(warning_row, 1).alignment = Alignment(wrap_text=True)


def build_workbook(output: Path) -> None:
    wb = Workbook()
    wb.properties.title = "HEMS laboratory experiment data collection"
    wb.properties.subject = "MFM384, PLC S7-1200, load control, quota, alerts and evidence logging"
    wb.properties.creator = "Codex"
    create_guide(wb)
    ranges = create_lists(wb)
    create_paper_map(wb)
    create_device_sheet(wb, ranges)
    create_session_sheet(wb, ranges)
    create_mfm_sheet(wb, ranges)
    create_api_latency_sheet(wb, ranges)
    create_command_sheet(wb, ranges)
    create_quota_sheet(wb, ranges)
    create_incident_sheet(wb, ranges)
    create_scenario_sheet(wb, ranges)
    create_evidence_sheet(wb, ranges)
    create_dictionary(wb)
    create_summary(wb)
    wb.active = wb.sheetnames.index("HUONG_DAN")
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def validate_workbook(output: Path) -> None:
    wb = load_workbook(output, data_only=False)
    required = {
        "HUONG_DAN", "ANH_XA_BAI_BAO", "THIET_BI", "PHIEN_THU", "DU_LIEU_MFM384",
        "DO_TRE_API", "LENH_PLC", "QUOTA_CANH_BAO", "SU_CO_AN_TOAN", "BANG_KICH_BAN",
        "NHAT_KY_TEP", "TU_DIEN_DU_LIEU", "TONG_HOP", "DANH_MUC",
    }
    if set(wb.sheetnames) != required:
        raise ValueError(f"Unexpected sheets: {wb.sheetnames}")
    if wb["DANH_MUC"].sheet_state != "hidden":
        raise ValueError("DANH_MUC must remain hidden")
    if wb["BANG_KICH_BAN"].max_row < 7:
        raise ValueError("Scenario matrix is incomplete")
    if wb["ANH_XA_BAI_BAO"].max_row < 9:
        raise ValueError("The 4-figure/4-table map is incomplete")


def main() -> int:
    args = parse_args()
    build_workbook(args.output)
    validate_workbook(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
