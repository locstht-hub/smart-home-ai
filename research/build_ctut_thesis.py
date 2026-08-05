from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


TITLE = "XÂY DỰNG HỆ THỐNG QUẢN LÝ NĂNG LƯỢNG NHÀ THÔNG MINH TÍCH HỢP PLC, GIÁM SÁT TỪ XA VÀ DỰ BÁO PHỤ TẢI"
PLACEHOLDER = "[CHƯA CUNG CẤP]"


def args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the canonical thesis in CTUT Appendix-II format.")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--architecture", type=Path)
    parser.add_argument("--pipeline", type=Path)
    parser.add_argument("--forecast-figure", type=Path)
    parser.add_argument("--references", type=Path, default=Path("Danh_Mục_Tài_Liệu_Tham_Khảo.md"))
    return parser.parse_args()


def set_cell_text(cell, text: str, bold: bool = False, size: int = 11) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "insideH"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "8")
        tag.set(qn("w:color"), "000000")
    for edge in ("left", "right", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "Cập nhật trường trong Word (Ctrl+A, F9)"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, display, end])


def set_page_number_header(section, number_format: str, start: int = 1) -> None:
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE")
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), number_format)
    pg_num.set(qn("w:start"), str(start))


def configure_section(section) -> None:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1)
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for name, size, bold, italic, alignment, uppercase in (
        ("Heading 1", 14, True, False, WD_ALIGN_PARAGRAPH.CENTER, True),
        ("Heading 2", 13, True, False, WD_ALIGN_PARAGRAPH.JUSTIFY, True),
        ("Heading 3", 13, True, False, WD_ALIGN_PARAGRAPH.JUSTIFY, False),
        ("Heading 4", 13, True, True, WD_ALIGN_PARAGRAPH.JUSTIFY, False),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.alignment = alignment
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Table Caption CTUT" not in styles:
        style = styles.add_style("Table Caption CTUT", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Table Caption CTUT"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_after = Pt(3)

    if "Figure Caption CTUT" not in styles:
        style = styles.add_style("Figure Caption CTUT", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Figure Caption CTUT"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_before = Pt(3)


def centered(document: Document, text: str = "", size: int = 13, bold: bool = False, space_after: int = 6):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    return paragraph


def page_break(document: Document) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_cover(document: Document, secondary: bool = False) -> None:
    centered(document, "TRƯỜNG ĐẠI HỌC KỸ THUẬT - CÔNG NGHỆ CẦN THƠ", 13, True)
    centered(document, f"KHOA {PLACEHOLDER}", 13, True)
    for _ in range(3):
        centered(document)
    centered(document, PLACEHOLDER, 14, True)
    centered(document, "MSSV: " + PLACEHOLDER, 13, True)
    for _ in range(2):
        centered(document)
    centered(document, "ĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC", 14, True)
    centered(document, TITLE, 15, True)
    for _ in range(3):
        centered(document)
    centered(document, "Ngành: " + PLACEHOLDER, 13, True)
    centered(document, "Mã ngành: " + PLACEHOLDER, 13, True)
    centered(document, "CÁN BỘ HƯỚNG DẪN", 13, True)
    centered(document, PLACEHOLDER, 13, True)
    for _ in range(2):
        centered(document)
    centered(document, "Cần Thơ, 2026", 13, True)
    if secondary:
        centered(document, "(Trang bìa phụ - in trên giấy trắng)", 11, False)
    page_break(document)


def add_confirmation(document: Document) -> None:
    centered(document, "TRANG XÁC NHẬN CỦA CÁN BỘ HƯỚNG DẪN VÀ BAN CHẤM", 14, True)
    document.add_paragraph(f"Tên đề tài: {TITLE}")
    document.add_paragraph(f"Sinh viên thực hiện: {PLACEHOLDER}")
    document.add_paragraph(f"MSSV: {PLACEHOLDER}    Lớp: {PLACEHOLDER}    Khóa: {PLACEHOLDER}")
    document.add_paragraph(f"Cán bộ hướng dẫn: {PLACEHOLDER}")
    document.add_paragraph("Nhận xét của cán bộ hướng dẫn:")
    document.add_paragraph("............................................................................................................................")
    document.add_paragraph("............................................................................................................................")
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, ("TRƯỞNG BAN\n(Ký, ghi rõ họ tên)", "CÁN BỘ PHẢN BIỆN\n(Ký, ghi rõ họ tên)", "ỦY VIÊN, THƯ KÝ\n(Ký, ghi rõ họ tên)")):
        set_cell_text(cell, value, True, 11)
    page_break(document)


def add_front_section(document: Document, title: str, body: list[str]) -> None:
    document.add_heading(title, level=1)
    for item in body:
        document.add_paragraph(item)
    page_break(document)


def extract_sections(source: Document) -> dict[str, list[tuple[str, str]]]:
    sections: dict[str, list[tuple[str, str]]] = {}
    current = "front"
    for paragraph in source.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if paragraph.style.name.startswith("Heading 1"):
            current = text
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append((paragraph.style.name, text))
    return sections


def add_mapped_content(document: Document, items: list[tuple[str, str]], chapter_number: int, start_section: int = 1) -> int:
    section_number = start_section
    subsection_number = 1
    for style_name, text in items:
        if style_name.startswith("Heading 2"):
            clean = text.split(" ", 1)[1] if " " in text and text[0].isdigit() else text
            prefix = f"{section_number}." if chapter_number == 0 else f"{chapter_number}.{section_number}."
            document.add_heading(f"{prefix} {clean.upper()}", level=2)
            section_number += 1
            subsection_number = 1
        elif style_name.startswith("Heading 3"):
            clean = text.split(" ", 1)[1] if " " in text and text[0].isdigit() else text
            document.add_heading(f"{chapter_number}.{section_number - 1}.{subsection_number}. {clean}", level=3)
            subsection_number += 1
        else:
            document.add_paragraph(apply_citations(text))
    return section_number


def add_picture(document: Document, path: Path | None, caption: str, width_cm: float = 14.5) -> None:
    if not path or not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    document.add_paragraph(caption, style="Figure Caption CTUT")


def add_table(document: Document, title: str, headers: list[str], rows: list[list[object]]) -> None:
    document.add_paragraph(title, style="Table Caption CTUT")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    set_table_borders(table)


def canonical_metric_rows(canonical: dict) -> list[list[object]]:
    rows = []
    for item in canonical["forecast"]["models"]:
        if item["split"] != "test":
            continue
        def value(key, std_key):
            val = item.get(key)
            std = item.get(std_key)
            return "-" if val is None else (f"{val:.3f} ± {std:.3f}" if std is not None else f"{val:.3f}")
        rows.append([
            item["model"], value("mae_kw", "mae_kw_std"), value("rmse_kw", "rmse_kw_std"),
            value("mape_percent", "mape_percent_std"), value("r2", "r2_std"),
            value("inference_ms_per_sample", "inference_ms_per_sample_std"),
        ])
    return rows


def apply_citations(text: str) -> str:
    text = re.sub(r"\s*\[(?:\d+(?:\s*,\s*\d+)*)\]", "", text).strip().rstrip(" .")
    if text.startswith("HEMS thu thập dữ liệu"):
        return text + ". Các kiến trúc và mô hình HEMS được tổng hợp trong [1],[2], ISO 17800 mô tả information model cho forecast, monitoring và demand response [3], còn demand-side management được tổng quan trong [4],[5]."
    if text.startswith("S7-1200 được dùng làm bộ điều khiển"):
        return text + ". Các đặc điểm PLC và hai tầng Modbus được đối chiếu theo [6],[7],[8]."
    if text.startswith("Random Forest tạo tập hợp cây"):
        return (
            text + ". Tổng quan forecasting và mô hình LSTM được trình bày trong [9],[10],[11]; "
            "cơ sở của XGBoost và Random Forest lần lượt là [12] và [13]. Dataset benchmark được trích dẫn tại [14]."
        )
    if text.startswith("MAE thể hiện sai số"):
        return text + ". Việc lựa chọn và diễn giải metric tham chiếu [15]."
    if text.startswith("Đối với chuỗi thời gian, random split"):
        return text + ". Rolling-origin được dùng để bảo toàn thứ tự thời gian và đánh giá khả năng dự báo tương lai [16]. ISO 50001 cung cấp khung quản lý và cải tiến hiệu quả năng lượng [17]; backend gần PLC có thể đảm nhiệm vai trò edge gateway [18]."
    if text.startswith("Authentication xác định người dùng"):
        return text + ". Threat model tham chiếu NISTIR 7628, OWASP ASVS, NISTIR 8259A và NIST SP 800-82 Rev. 3 [19],[20],[21],[22]."
    if text.startswith("Ứng dụng được xây bằng Expo/React Native"):
        return text + ". Backend cục bộ có thể được xem như edge gateway giữa ứng dụng và PLC [18]."
    if text.startswith("Quota trong phiên bản hiện tại"):
        return text + ". Demand-side management và demand response tạo cơ sở lý thuyết cho policy tương lai [4],[5], trong khi ISO 50001 cung cấp khung quản lý và cải tiến hiệu quả năng lượng [17]."
    return text + ("." if text and text[-1] not in ".!?" else "")


def load_references(path: Path) -> list[str]:
    references = [
        line.strip() for line in path.read_text(encoding="utf-8").splitlines()
        if re.match(r"^\[\d+\]\s", line.strip())
    ]
    if len(references) != 22:
        raise ValueError(f"Expected 22 verified references, found {len(references)} in {path}")
    return references


def add_references(document: Document, references: list[str]) -> None:
    document.add_heading("TÀI LIỆU THAM KHẢO", level=1)
    for text in references:
        p = document.add_paragraph(text)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.first_line_indent = Cm(-1)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> int:
    options = args()
    source = Document(options.source)
    canonical = json.loads(options.canonical.read_text(encoding="utf-8"))
    references = load_references(options.references)
    content = extract_sections(source)

    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])
    document.sections[0].different_first_page_header_footer = True
    add_cover(document)
    add_cover(document, secondary=True)
    add_confirmation(document)

    front = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front)
    set_page_number_header(front, "lowerRoman", 1)
    add_front_section(document, "LỜI CAM ĐOAN", [
        "Tôi cam đoan nội dung và kết quả trong đồ án này được thực hiện trung thực. Các nguồn tham khảo được trích dẫn theo quy định; dữ liệu mô phỏng, dữ liệu mẫu và dữ liệu phần cứng thật được phân biệt rõ ràng.",
        f"Sinh viên thực hiện: {PLACEHOLDER}",
    ])
    add_front_section(document, "LỜI CẢM ƠN", [
        "Tác giả trân trọng cảm ơn cán bộ hướng dẫn, quý thầy cô và các cá nhân, đơn vị đã hỗ trợ quá trình thực hiện đề tài. Thông tin cụ thể sẽ được bổ sung sau khi người thực hiện xác nhận.",
    ])
    abstract = next((items for name, items in content.items() if name == "Tóm tắt"), [])
    add_front_section(document, "TÓM TẮT", [text for _, text in abstract] or ["[CHỜ BỔ SUNG TÓM TẮT]"])
    document.add_heading("MỤC LỤC", level=1)
    toc = document.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    page_break(document)
    add_front_section(document, "DANH MỤC CÁC TỪ VIẾT TẮT VÀ THUẬT NGỮ", [
        "API: Application Programming Interface", "HEMS: Home Energy Management System",
        "MFM384: Đồng hồ đo điện đa chức năng", "PLC: Programmable Logic Controller",
        "RBAC: Role-Based Access Control", "RTT: Round-Trip Time",
    ])
    add_front_section(document, "DANH MỤC HÌNH VẼ", ["Cập nhật tự động trong Microsoft Word sau khi hoàn thiện caption."])
    add_front_section(document, "DANH MỤC BẢNG BIỂU", ["Cập nhật tự động trong Microsoft Word sau khi hoàn thiện caption."])

    main_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main_section)
    set_page_number_header(main_section, "decimal", 1)

    document.add_heading("LỜI MỞ ĐẦU", level=1)
    intro_items = content.get("CHƯƠNG 1. TỔNG QUAN ĐỀ TÀI", [])
    add_mapped_content(document, intro_items, 0)

    document.add_heading("CHƯƠNG 1. TỔNG QUAN LÝ THUYẾT VÀ CƠ SỞ NGHIÊN CỨU", level=1)
    document.add_paragraph(
        "Các nguồn được trình bày theo thứ tự xuất hiện: HEMS và demand response [1],[2],[3],[4],[5]; "
        "Modbus và PLC [6],[7],[8]; forecasting, mô hình, dataset và đánh giá [9],[10],[11],[12],[13],[14],[15],[16]; "
        "quản lý năng lượng và edge computing [17],[18]; bảo mật smart-grid, ứng dụng, IoT và OT [19],[20],[21],[22]."
    )
    add_mapped_content(document, content.get("CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG TRÌNH LIÊN QUAN", []), 1)

    document.add_heading("CHƯƠNG 2. PHƯƠNG TIỆN VÀ PHƯƠNG PHÁP NGHIÊN CỨU", level=1)
    add_mapped_content(document, content.get("CHƯƠNG 3. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", []), 2)
    add_picture(document, options.architecture, "Hình 2.1. Kiến trúc tổng thể và ranh giới điều khiển của hệ thống")
    add_picture(document, options.pipeline, "Hình 2.2. Pipeline dữ liệu và evidence gate")
    add_mapped_content(document, content.get("CHƯƠNG 4. HIỆN THỰC HỆ THỐNG", []), 2, start_section=8)
    add_table(document, "Bảng 2.1. Ma trận phân quyền chính", ["Vai trò", "Đọc", "Điều khiển", "Quản lý thành viên", "Ghi telemetry"], [
        ["Owner", "Có", "Có", "Có", "Không"], ["Member", "Có", "Theo quyền", "Không", "Không"],
        ["Viewer", "Có", "Không", "Không", "Không"], ["Telemetry service", "Không", "Không", "Không", "Theo allowedHomeIds"],
    ])

    document.add_heading("CHƯƠNG 3. KẾT QUẢ VÀ THẢO LUẬN", level=1)
    add_mapped_content(document, content.get("CHƯƠNG 5. THỰC NGHIỆM VÀ ĐÁNH GIÁ", []), 3)
    add_table(document, "Bảng 3.1. Kết quả benchmark forecast trên tập test", ["Mô hình", "MAE (kW)", "RMSE (kW)", "MAPE (%)", "R²", "Inference (ms/mẫu)"], canonical_metric_rows(canonical))
    add_picture(document, options.forecast_figure, "Hình 3.1. So sánh MAE và RMSE của các mô hình trên nguồn canonical")
    document.add_paragraph(
        "Chưa có bảng độ trễ phần cứng vì môi trường hiện tại không có đủ raw trial được xác nhận source=plc-s7-1200 và effectiveMode=plc-real. Do đó, claim độ trễ phần cứng và tự động sa thải tải tiếp tục bị khóa."
    )

    document.add_heading("CHƯƠNG 4. KẾT LUẬN VÀ ĐỀ XUẤT", level=1)
    add_mapped_content(document, content.get("CHƯƠNG 6. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", []), 4)

    add_references(document, references)
    document.add_heading("PHỤ LỤC", level=1)
    document.add_heading("PHỤ LỤC A. LỆNH TÁI LẬP THỰC NGHIỆM", level=2)
    for command in (
        "python research/run_forecast_benchmark.py --data-source uci --folds 2 --seeds 42,3407",
        "python research/build_canonical_results.py --forecast-metrics research/results/benchmark_latest/metrics.json",
        "npm run verify",
    ):
        p = document.add_paragraph(command)
        p.paragraph_format.first_line_indent = Cm(0)
    document.add_heading("PHỤ LỤC B. TRẠNG THÁI EVIDENCE GATE", level=2)
    for key, value in canonical["claimPolicy"].items():
        if key.startswith("allow"):
            document.add_paragraph(f"{key}: {value}")

    set_update_fields(document)
    options.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(options.output)
    print(options.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
