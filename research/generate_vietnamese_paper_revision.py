from __future__ import annotations

import argparse
import json
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


TITLE_VI = "TÍCH HỢP BỘ ĐIỀU KHIỂN KHẢ TRÌNH VÀ DỰ BÁO PHỤ TẢI CHO QUẢN LÝ NĂNG LƯỢNG NHÀ THÔNG MINH"
TITLE_EN = "Integrating a programmable logic controller and load forecasting for smart home energy management"

ABSTRACT_VI = (
    "Nghiên cứu trình bày nguyên mẫu quản lý năng lượng nhà thông minh tích hợp ứng dụng di động, backend biên, "
    "bộ điều khiển khả trình Siemens S7-1200 và dự báo phụ tải 24 giờ. Mục tiêu là tạo luồng điều khiển có thể truy vết, "
    "tách tài khoản người dùng khỏi credential telemetry và phân biệt bằng chứng phần mềm, benchmark công khai với "
    "thử nghiệm phần cứng. Kiến trúc áp dụng phân quyền theo nhà, tuần tự hóa thao tác PLC, status tag độc lập và "
    "evidence gate dựa trên nguồn kết quả chuẩn. Benchmark UCI dùng rolling-origin, hai fold và hai random seed. "
    "XGBoost đạt MAE 0,486 ± 0,046 kW, thấp hơn seasonal naive 24 giờ khoảng 9,8%, nhưng chỉ thấp hơn Random Forest "
    "khoảng 1,2%; MAPE còn cao 66,2 ± 8,7%. Kiểm thử hợp đồng phần mềm xác nhận cách ly quyền, telemetry service-only, "
    "phản hồi độc lập và fail-closed, song không thay thế phép đo PLC/MFM384 thật. Bằng chứng hiện có phù hợp để định vị "
    "công trình là nguyên mẫu kiến trúc có kiểm soát claim. Độ trễ, độ chính xác địa phương và sa thải tải vẫn là giới hạn."
)

ABSTRACT_EN = (
    "This study presents a smart home energy management prototype integrating a mobile application, an edge backend, "
    "a Siemens S7-1200 programmable logic controller, and a 24-hour load forecasting module. The objective is to make "
    "control actions traceable, separate user credentials from telemetry authentication, and distinguish software "
    "evidence, public-data benchmarking, and physical experiments. The architecture applies home-scoped authorization, "
    "serialized PLC operations, an independent status tag, and an evidence gate driven by a canonical result source. "
    "The UCI benchmark uses rolling-origin evaluation with two folds and two random seeds. XGBoost obtains a test MAE "
    "of 0.486 ± 0.046 kW, approximately 9.8% lower than the 24-hour seasonal-naive baseline but only 1.2% lower than "
    "Random Forest; MAPE remains high at 66.2 ± 8.7%. Software-contract tests support authorization isolation, "
    "service-only telemetry, independent feedback, and fail-closed behavior, but they do not substitute for physical "
    "PLC/MFM384 measurements. The evidence therefore supports positioning the work as an evidence-controlled system "
    "prototype. Real-hardware latency, local forecasting accuracy, and automatic load shedding remain unverified and "
    "are explicitly retained as limitations."
)


REFERENCES = [
    "[1] Shareef H, Ahmed MS, Mohamed A, Al Hassan E. Review on home energy management system considering demand responses, smart technologies, and intelligent controllers. IEEE Access. 2018;6:24498-24509. doi:10.1109/ACCESS.2018.2831917.",
    "[2] Gomes I, Bot K, Ruano MG, Ruano A. Recent techniques used in home energy management systems: a review. Energies. 2022;15(8):2866. doi:10.3390/en15082866.",
    "[3] Motta LL, Ferreira LCBC, Cabral TW, et al. General overview and proof of concept of a smart home energy management system architecture. Electronics. 2023;12(21):4453. doi:10.3390/electronics12214453.",
    "[4] Ferreira LCBC, Borchardt ADR, Cardoso GDS, et al. Edge computing and microservices middleware for home energy management systems. IEEE Access. 2022;10:109663-109676. doi:10.1109/ACCESS.2022.3214229.",
    "[5] Xu H, König L, Cáliz D, Schmeck H. A generic user interface for energy management in smart homes. Energy Inform. 2018;1:55. doi:10.1186/s42162-018-0060-0.",
    "[6] Hlayel M, Mahdin H, Hayajneh M, Nurwarsito H. Toward Industry 5.0: a WebSocket-S7 bridge for low-latency, IEC 61588-compliant digital twins in remote industrial automation. PLoS One. 2026;21(5):e0342004. doi:10.1371/journal.pone.0342004.",
    "[7] Cao Z, Han X, Lyons W, O'Rourke F. Energy management optimisation using a combined Long Short-Term Memory recurrent neural network-Particle Swarm Optimisation model. J Clean Prod. 2021;326:129246. doi:10.1016/j.jclepro.2021.129246.",
    "[8] Semmelmann L, Henni S, Weinhardt C. Load forecasting for energy communities: a novel LSTM-XGBoost hybrid model based on smart meter data. Energy Inform. 2022;5:24. doi:10.1186/s42162-022-00212-9.",
    "[9] Siemens AG. SIMATIC S7-1200 programmable controller system manual. Version 4.6, document A5E02486680-AP. Nuremberg: Siemens AG; 2022.",
    "[10] Modbus Organization. MODBUS application protocol specification V1.1b3. Hopkinton (MA): Modbus Organization; 2012.",
    "[11] Stouffer K, Pease M, Tang C, et al. Guide to operational technology (OT) security. NIST SP 800-82 Rev. 3. Gaithersburg (MD): National Institute of Standards and Technology; 2023. doi:10.6028/NIST.SP.800-82r3.",
    "[12] Fagan M, Megas KN, Scarfone K, Smith M. IoT device cybersecurity capability core baseline. NISTIR 8259A. Gaithersburg (MD): National Institute of Standards and Technology; 2020. doi:10.6028/NIST.IR.8259A.",
    "[13] OWASP Foundation. Application Security Verification Standard [Internet]. Available from: https://owasp.org/www-project-application-security-verification-standard/ [cited 2026 Jul 16].",
    "[14] Hebrail G, Berard A. Individual household electric power consumption [dataset]. Irvine (CA): UCI Machine Learning Repository; 2006. doi:10.24432/C58K54.",
    "[15] Breiman L. Random forests. Mach Learn. 2001;45:5-32. doi:10.1023/A:1010933404324.",
    "[16] Chen T, Guestrin C. XGBoost: a scalable tree boosting system. In: Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining; 2016. p. 785-794. doi:10.1145/2939672.2939785.",
    "[17] Hyndman RJ, Koehler AB. Another look at measures of forecast accuracy. Int J Forecast. 2006;22(4):679-688. doi:10.1016/j.ijforecast.2006.03.001.",
    "[18] Bergmeir C, Benítez JM. On the use of cross-validation for time series predictor evaluation. Inf Sci. 2012;191:192-213. doi:10.1016/j.ins.2011.12.028.",
]


RELATED_WORK = [
    ["Công trình", "Bối cảnh", "PLC/phản hồi", "Quyền truy cập", "Dự báo/đánh giá"],
    ["Shareef et al. [1]", "Tổng quan HEMS", "Không", "Không trọng tâm", "Taxonomy DR/controller"],
    ["Gomes et al. [2]", "Tổng quan hệ thống", "Không", "Không trọng tâm", "Ma trận kỹ thuật HEMS"],
    ["Motta et al. [3]", "PoC end-to-end", "Thiết bị thật; middleware", "Không theo home scope", "Trace và integrity test"],
    ["Ferreira et al. [4]", "Edge microservices", "Không phải đóng góp chính", "Không trọng tâm", "Đánh giá middleware"],
    ["Xu et al. [5]", "UI smart home", "Command qua BOS", "Ba vai trò", "Usability/user study"],
    ["Hlayel et al. [6]", "PLC-cloud bridge", "S7 hai chiều", "Không trọng tâm", "RTT, p95/p99, tải đồng thời"],
    ["Cao et al. [7]", "Forecast + tối ưu", "Không", "Không", "LSTM, baseline, Wilcoxon"],
    ["Semmelmann et al. [8]", "Forecast cộng đồng", "Không", "Không", "LSTM-XGB, 12-fold, kiểm định"],
    ["Công trình này", "Prototype HEMS biên", "S7-1200; status tag độc lập", "Home scope; service token", "Software contract + UCI; hardware chờ"],
]


SOFTWARE_EVIDENCE = [
    ["Nhóm kiểm tra", "Tình huống chính", "Trạng thái", "Loại bằng chứng", "Phạm vi diễn giải"],
    ["Xác thực và phân quyền", "Thu hồi phiên, hết hạn, khóa người dùng, giới hạn đăng nhập, cách ly nhà", "Đạt", "Hợp đồng phần mềm", "Không phải penetration test"],
    ["Telemetry", "Chỉ service credential được ghi; GET chỉ đọc; mock fallback không lưu thành dữ liệu thật", "Đạt", "Hợp đồng phần mềm", "Không xác nhận cảm biến MFM384"],
    ["Lệnh PLC", "I/O tuần tự, status tag độc lập, idempotency, lỗi từng thiết bị trong scene", "Đạt", "Hợp đồng phần mềm", "Không xác nhận RTT phần cứng"],
    ["Forecast API", "Kiểm tra artifact, timestamp, model, sample source và từ chối retrain giả", "Đạt", "Hợp đồng phần mềm", "Có cảnh báo khác phiên bản artifact"],
    ["Frontend", "Auth/chat/room/theme contract", "Chưa đạt hoàn toàn", "Hợp đồng giao diện", "Không tuyên bố UI đã xác minh đầy đủ"],
]


EVIDENCE_GATE = [
    ["Claim", "Hợp đồng phần mềm", "Dữ liệu công khai", "PLC/MFM384 thật", "Trạng thái công bố"],
    ["Cách ly quyền và telemetry", "Có", "Không áp dụng", "Chưa đo", "Được mô tả là software-validated"],
    ["Tuần tự lệnh và feedback", "Có", "Không áp dụng", "Chưa đo", "Được mô tả là software contract"],
    ["Benchmark dự báo UCI", "Pipeline đã kiểm tra", "Có", "Không", "Được công bố trong phạm vi UCI"],
    ["Độ chính xác tại Cần Thơ", "Không đủ", "Không đại diện", "Chưa có", "Chỉ là hướng nghiên cứu"],
    ["Latency end-to-end", "Không thay thế phép đo", "Không áp dụng", "Chưa có", "Không đưa số liệu"],
    ["Sa thải tải tự động", "Fail-closed", "Không áp dụng", "Chưa có", "Không tuyên bố đã hoạt động"],
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the Vietnamese paper revision from the official journal template.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def set_text_preserve_first_run(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_run_texts(paragraph: Paragraph, values: list[str]) -> None:
    if len(paragraph.runs) < len(values):
        raise ValueError(f"Template paragraph has {len(paragraph.runs)} runs, expected at least {len(values)}")
    for index, run in enumerate(paragraph.runs):
        run.text = values[index] if index < len(values) else ""


def insert_before_sectpr(document: Document, element) -> None:
    body = document._body._element
    sect_pr = body.sectPr
    if sect_pr is None:
        body.append(element)
    else:
        sect_pr.addprevious(element)


def clone_paragraph(document: Document, template_paragraph: Paragraph, text: str) -> Paragraph:
    element = deepcopy(template_paragraph._p)
    text_nodes = element.xpath(".//w:t")
    if text_nodes:
        text_nodes[0].text = text
        for node in text_nodes[1:]:
            node.text = ""
    insert_before_sectpr(document, element)
    paragraph = Paragraph(element, document._body)
    if not text_nodes:
        paragraph.add_run(text)
    return paragraph


def add_figure(document: Document, image_template: Paragraph, caption_template: Paragraph, path: Path, caption: str) -> None:
    paragraph = clone_paragraph(document, image_template, "")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.75))
    clone_paragraph(document, caption_template, caption)


def set_cell_text(cell, text: str, *, bold: bool = False, font_size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        if len(values) != len(rows[0]):
            raise ValueError("Table rows must have equal column counts")
        tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for column_index, value in enumerate(values):
            set_cell_text(table.rows[row_index].cells[column_index], value, bold=row_index == 0)
            if widths:
                table.rows[row_index].cells[column_index].width = Inches(widths[column_index])


def add_captioned_table(
    document: Document,
    caption_template: Paragraph,
    note_template: Paragraph,
    caption: str,
    rows: list[list[str]],
    note: str = "",
    widths: list[float] | None = None,
) -> None:
    clone_paragraph(document, caption_template, caption)
    add_table(document, rows, widths)
    if note:
        clone_paragraph(document, note_template, note)


def model_rows(canonical: dict) -> list[list[str]]:
    names = {
        "persistence": "Persistence",
        "seasonal_naive_24h": "Seasonal naive 24 h",
        "seasonal_naive_168h": "Seasonal naive 168 h",
        "random_forest": "Random Forest",
        "xgboost": "XGBoost",
    }
    rows = [["Mô hình", "MAE (kW)", "RMSE (kW)", "MAPE (%)", "R²", "Inference (ms/mẫu)"]]
    for item in [row for row in canonical["forecast"]["models"] if row.get("split") == "test"]:
        inference = item["inference_ms_per_sample"]
        inference_text = "<0,001" if inference < 0.001 else f"{inference:.3f} ± {item['inference_ms_per_sample_std']:.3f}".replace(".", ",")
        rows.append(
            [
                names[item["model"]],
                f"{item['mae_kw']:.3f} ± {item['mae_kw_std']:.3f}".replace(".", ","),
                f"{item['rmse_kw']:.3f} ± {item['rmse_kw_std']:.3f}".replace(".", ","),
                f"{item['mape_percent']:.1f} ± {item['mape_percent_std']:.1f}".replace(".", ","),
                f"{item['r2']:.3f} ± {item['r2_std']:.3f}".replace(".", ","),
                inference_text,
            ]
        )
    return rows


def format_integer_vi(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def clear_template_body_after_abstract(document: Document) -> None:
    second_table = document.tables[1]._tbl
    current = second_table.getnext()
    body = document._body._element
    while current is not None:
        next_element = current.getnext()
        if current.tag != qn("w:sectPr"):
            body.remove(current)
        current = next_element


def populate_front_matter(document: Document) -> None:
    set_text_preserve_first_run(document.paragraphs[0], TITLE_VI)
    set_text_preserve_first_run(document.paragraphs[1], "")
    set_text_preserve_first_run(document.paragraphs[3], "[HỌ VÀ TÊN TÁC GIẢ]1, [HỌ VÀ TÊN ĐỒNG TÁC GIẢ]2")
    set_text_preserve_first_run(document.paragraphs[4], "1 [KHOA/ĐƠN VỊ, TRƯỜNG ĐẠI HỌC KỸ THUẬT - CÔNG NGHỆ CẦN THƠ]")
    set_text_preserve_first_run(document.paragraphs[5], "2 [ĐƠN VỊ CỦA ĐỒNG TÁC GIẢ, NẾU CÓ]")
    set_text_preserve_first_run(document.paragraphs[6], "Tác giả liên hệ: [EMAIL TÁC GIẢ CHỊU TRÁCH NHIỆM]")

    vi_info = document.tables[0].cell(0, 0).paragraphs[0]
    set_run_texts(
        vi_info,
        [
            "Thông tin chung\n",
            "Ngày nhận bài: dd/mm/yyyy\nNgày nhận bài sửa: dd/mm/yyyy\nNgày duyệt đăng: dd/mm/yyyy\n\n",
            "Từ khóa:\n",
            "bộ điều khiển khả trình, dự báo phụ tải, hệ thống quản lý năng lượng, nhà thông minh, phân quyền",
        ],
    )
    set_text_preserve_first_run(document.tables[0].cell(0, 1).paragraphs[1], ABSTRACT_VI)

    en_info = document.tables[1].cell(0, 0).paragraphs[0]
    set_run_texts(
        en_info,
        [
            "Title:\n",
            TITLE_EN + "\n\n",
            "Keywords:\n",
            "authorization, home energy management system, load forecasting, programmable logic controller, smart home",
        ],
    )
    set_text_preserve_first_run(document.tables[1].cell(0, 1).paragraphs[1], ABSTRACT_EN)


def build_paper(document: Document, canonical: dict, figures_dir: Path, samples: dict[str, Paragraph]) -> None:
    h1 = samples["h1"]
    h2 = samples["h2"]
    h3 = samples["h3"]
    body = samples["body"]
    table_caption = samples["table_caption"]
    table_note = samples["table_note"]
    image_placeholder = samples["image"]
    figure_caption = samples["figure_caption"]
    ack = samples["ack"]
    ref_heading = samples["ref_heading"]
    ref = samples["ref"]

    clone_paragraph(document, h1, "1. ĐẶT VẤN ĐỀ")
    clone_paragraph(
        document,
        body,
        "Hệ thống quản lý năng lượng nhà thông minh (HEMS) kết hợp giám sát, điều khiển thiết bị, dự báo và hỗ trợ quyết định trong một hạ tầng có tác động vật lý. Các tổng quan cho thấy phần lớn công trình tập trung vào demand response, scheduling hoặc thuật toán tối ưu [1], [2]. Với nguyên mẫu có PLC, tiêu chí đánh giá không thể chỉ là giao diện hoạt động hay sai số dự báo; một lệnh chỉ được xem là thành công khi đúng người dùng, đúng phạm vi nhà và nhận được trạng thái phản hồi độc lập.",
    )
    clone_paragraph(
        document,
        body,
        "Các system paper gần đề tài thường trình bày PoC end-to-end, middleware và trace phần cứng [3], [4]. Nghiên cứu giao diện HEMS đã đánh giá vai trò người dùng và usability [5], trong khi nghiên cứu PLC-cloud báo cáo RTT, p95/p99 và tải đồng thời trên testbed thật [6]. Nhánh dự báo thường dùng baseline, kiểm định thống kê và liên kết sai số với outcome quản lý năng lượng [7], [8]. Những hướng này cung cấp chuẩn đối chiếu, nhưng không thể thay thế phép đo trên chính hệ thống S7-1200/MFM384 đang nghiên cứu.",
    )
    clone_paragraph(
        document,
        body,
        "Khoảng trống của công trình nằm ở giao điểm giữa kiến trúc HEMS biên, cách ly quyền truy cập, hợp đồng lệnh-phản hồi PLC và quản trị claim bằng nguồn bằng chứng duy nhất. Bài báo không định vị XGBoost là thuật toán mới, cũng không xem kiểm thử phần mềm là bằng chứng hiệu năng phần cứng. Mô-đun dự báo được giữ như một case study định lượng để kiểm tra khả năng tái lập của pipeline.",
    )
    clone_paragraph(document, h2, "1.1. Câu hỏi nghiên cứu")
    clone_paragraph(document, body, "RQ1: Kiến trúc nào tách tài khoản ứng dụng khỏi credential telemetry và giới hạn thao tác theo phạm vi nhà?")
    clone_paragraph(document, body, "RQ2: Hợp đồng phần mềm nào giảm nguy cơ lost update, trạng thái giả thành công và lỗi từng phần khi gửi lệnh tới PLC?")
    clone_paragraph(document, body, "RQ3: Benchmark dự báo công khai hiện hỗ trợ claim nào, và claim nào phải giữ ở trạng thái chưa kiểm chứng?")
    clone_paragraph(document, h2, "1.2. Đóng góp và phạm vi")
    clone_paragraph(
        document,
        body,
        "Bài báo đóng góp bốn thành phần có thể truy vết: kiến trúc nhiều lớp với home-scoped authorization và service credential riêng; hợp đồng I/O PLC được tuần tự hóa với status tag độc lập; pipeline canonical khóa dataset, metrics và trạng thái claim bằng fingerprint; và benchmark 24 giờ so sánh XGBoost, Random Forest với ba baseline. Các đóng góp này được đánh giá bằng hai lớp bằng chứng hiện có là hợp đồng phần mềm và dataset công khai. Độ trễ PLC/MFM384 thật, độ chính xác tại Cần Thơ và hiệu quả sa thải tải tự động nằm ngoài Results.",
    )

    clone_paragraph(document, h1, "2. CƠ SỞ LÝ THUYẾT/ MÔ HÌNH/ PHƯƠNG PHÁP NGHIÊN CỨU")
    clone_paragraph(document, h2, "2.1. Công trình liên quan và khoảng trống")
    clone_paragraph(
        document,
        body,
        "Bảng 1 đối chiếu các capability trực tiếp liên quan. PoC của Motta et al. [3] có bằng chứng thiết bị và giao diện phong phú; middleware của Ferreira et al. [4] làm rõ vai trò edge; Xu et al. [5] cung cấp mô hình vai trò và UI; Hlayel et al. [6] thiết kế phép đo PLC-cloud. Hai nghiên cứu dự báo [7], [8] cho thấy baseline, multiple folds và kiểm định là cần thiết khi muốn tuyên bố ưu thế mô hình. Chưa công trình nào trong nhóm đối chiếu dùng đồng thời home scope, telemetry credential tách biệt, serialized PLC write, status feedback độc lập và evidence gate như một đóng góp thống nhất.",
    )
    add_captioned_table(
        document,
        table_caption,
        table_note,
        "Bảng 1. Đối chiếu các công trình gần đề tài",
        RELATED_WORK,
        "Ghi chú: Ô 'không' chỉ mô tả phạm vi bài được đối chiếu, không phải đánh giá chất lượng công trình. Dòng công trình này phân biệt thiết kế phần mềm với kiểm chứng phần cứng.",
        [1.15, 1.3, 1.55, 1.35, 1.65],
    )

    clone_paragraph(document, h2, "2.2. Thiết kế nghiên cứu và phân lớp bằng chứng")
    clone_paragraph(
        document,
        body,
        "Nghiên cứu áp dụng thiết kế evaluation theo artifact. Lớp thứ nhất kiểm tra hợp đồng phần mềm bằng test tự động đối với authentication, authorization, telemetry, PLC I/O, forecast API và frontend. Lớp thứ hai đánh giá mô hình trên dataset UCI có fingerprint và rolling-origin. Lớp thứ ba yêu cầu raw log từ PLC S7-1200, MFM384 và tải thật. Một kết quả ở lớp trước không được dùng thay cho lớp sau; vì vậy mọi bảng và hình đều ghi rõ loại bằng chứng.",
    )
    clone_paragraph(document, h3, "2.2.1. Quy tắc evidence gate")
    clone_paragraph(
        document,
        body,
        "Nguồn máy đọc duy nhất là canonical_results.json. Claim chỉ được mở khi cờ bằng chứng tương ứng là true và nguồn chứa dataset hash hoặc raw trial provenance. Quy tắc này ngăn việc biến kế hoạch thử nghiệm, mock fallback hay giá trị latency placeholder thành kết quả. Hình 1 trình bày ba luồng và ranh giới claim.",
    )
    add_figure(
        document,
        image_placeholder,
        figure_caption,
        figures_dir / "paper_evidence_gate_vi.png",
        "Hình 1. Evidence gate tách hợp đồng phần mềm, benchmark công khai và thử nghiệm PLC/MFM384 thật; lớp phần cứng vẫn đang chờ dữ liệu.",
    )

    clone_paragraph(document, h2, "2.3. Kiến trúc HEMS và ranh giới tin cậy")
    clone_paragraph(
        document,
        body,
        "PLC S7-1200 và Modbus cung cấp lớp điều khiển và trao đổi thanh ghi [9], [10]. Trong prototype, mọi S7 I/O đi qua một lock chung vì ghi bit PLC là thao tác read-modify-write. Cấu hình nhiều worker bị từ chối khi tiến trình còn trực tiếp sở hữu PLC. Lệnh ưu tiên onCommandTag/offCommandTag dạng xung; trạng thái hoàn thành được đọc từ statusTag độc lập thay vì suy ra từ command bit.",
    )
    clone_paragraph(
        document,
        body,
        "Kiến trúc ở Hình 2 gồm ứng dụng di động/web quản trị, API biên, forecast service, kho dữ liệu, PLC gateway và lớp vật lý. Ứng dụng không truy cập cơ sở dữ liệu hoặc PLC trực tiếp. API xác thực phiên, xác định homeId, kiểm tra role, ghi audit log và chỉ sau đó chuyển lệnh tới gateway. Telemetry ingestion dùng service token riêng; bearer token của người dùng không được phép ghi telemetry. Cách tách này bám theo nguyên tắc least privilege cho IoT/OT [11], [12], [13].",
    )
    add_figure(
        document,
        image_placeholder,
        figure_caption,
        figures_dir / "paper_architecture_vi.png",
        "Hình 2. Kiến trúc nhiều lớp phân biệt luồng điều khiển, phản hồi và telemetry; đường đến tải thật chưa được xem là bằng chứng thực nghiệm.",
    )

    clone_paragraph(document, h2, "2.4. Hợp đồng lệnh-phản hồi và trạng thái ứng dụng")
    clone_paragraph(
        document,
        body,
        "Hình 3 mô tả một thao tác điều khiển. Ứng dụng đặt trạng thái đang xử lý, backend kiểm tra home scope và quyền thiết bị, gateway đọc status hiện tại rồi chỉ phát xung nếu trạng thái đích chưa đạt. Sau khi ghi lệnh, gateway poll statusTag đến khi khớp hoặc hết timeout. Backend trả cấu trúc kết quả có nguyên nhân để ứng dụng hiển thị success, denied, connection error hoặc feedback timeout. Scene nhiều thiết bị báo lỗi từng phần thay vì che lấp thiết bị thất bại.",
    )
    clone_paragraph(
        document,
        body,
        "Chuỗi trạng thái này là bằng chứng thiết kế và hợp đồng phần mềm, không phải đo thời gian đáp ứng người dùng hay usability. Việc đưa khả năng phản hồi của app vào bài báo là hợp lý vì nó trực tiếp nối RQ2 với outcome quan sát được; screenshot trang trí, chatbot và theme không được xem là contribution.",
    )
    add_figure(
        document,
        image_placeholder,
        figure_caption,
        figures_dir / "paper_command_feedback_vi.png",
        "Hình 3. Chuỗi App-API-PLC-feedback và bốn trạng thái ứng dụng ở mức software contract; hình không biểu diễn latency đo trên phần cứng.",
    )

    clone_paragraph(document, h2, "2.5. Dữ liệu, đặc trưng và mô hình dự báo")
    dataset = canonical["forecast"]["dataset"]
    clone_paragraph(
        document,
        body,
        f"Benchmark dùng Individual Household Electric Power Consumption của UCI [14]. Canonical run chứa {format_integer_vi(dataset['raw_rows'])} dòng thô, {format_integer_vi(dataset['hourly_rows'])} dòng theo giờ và {format_integer_vi(dataset['supervised_rows'])} mẫu supervised. Công suất tác dụng, phản kháng, điện áp và dòng điện được tổng hợp theo trung bình giờ; ba kênh sub-metering được cộng theo giờ. Giá trị thiếu được nội suy theo thời gian; các biến phụ còn thiếu sau nội suy được điền 0. Việc dùng UCI chỉ nhằm kiểm tra pipeline và baseline, không đại diện hộ gia đình Cần Thơ.",
    )
    clone_paragraph(
        document,
        body,
        "Đặc trưng gồm giờ, thứ, tháng, ngày trong năm, weekend và các biến chu kỳ sin/cos; lag công suất từ 1-24 giờ, 144-168 giờ và các mốc 48, 72, 96, 120, 336 giờ; mean, standard deviation, min và max trượt ở cửa sổ 3, 6, 12, 24, 48, 72, 168 giờ; delta, ratio, exponentially weighted mean; cùng lag/rolling mean của điện áp, dòng điện, công suất phản kháng và sub-metering. Tất cả rolling feature được shift trước một bước để tránh leakage.",
    )
    clone_paragraph(
        document,
        body,
        "Bài toán dùng direct multi-horizon: mỗi mô hình huấn luyện một estimator cho từng bước từ h+1 đến h+24. Random Forest [15] dùng 50 cây, max_depth=12 và max_features=sqrt. XGBoost [16] dùng 150 cây, max_depth=5, learning_rate=0,035, subsample=0,95, colsample_bytree=0,90, reg_alpha=0,03 và reg_lambda=1,8. Đây là cấu hình cố định của pipeline hiện hành, chưa phải hyperparameter search toàn diện.",
    )

    clone_paragraph(document, h2, "2.6. Thiết kế đánh giá")
    ranges = dataset["fold_ranges"]
    clone_paragraph(
        document,
        body,
        f"Đánh giá expanding rolling-origin gồm {dataset['rolling_folds']} fold. Fold 1 huấn luyện {format_integer_vi(ranges[0]['train']['rows'])} mẫu, validation {format_integer_vi(ranges[0]['val']['rows'])} mẫu và test {format_integer_vi(ranges[0]['test']['rows'])} mẫu; fold 2 mở rộng train lên {format_integer_vi(ranges[1]['train']['rows'])} mẫu và giữ hai cửa sổ sau cùng cùng kích thước. Hai random seed là {dataset['random_seeds'][0]} và {dataset['random_seeds'][1]}; vì vậy mô hình học máy có bốn lần chạy, còn baseline có hai lần. Mô hình được chọn theo validation MAE mean.",
    )
    clone_paragraph(
        document,
        body,
        "Baseline gồm persistence, seasonal naive 24 giờ và seasonal naive 168 giờ. MAE là chỉ số chính; RMSE nhấn mạnh sai số lớn, R² mô tả phần phương sai được giải thích, còn MAPE dùng mẫu số tối thiểu 0,2 kW do tải gần 0 làm sai số phần trăm mất ổn định [17]. Mean ± sample SD được tổng hợp trên các lần chạy. Rolling-origin được dùng để giữ thứ tự thời gian và giảm optimistic bias [18].",
    )

    clone_paragraph(document, h1, "3. KẾT QUẢ NGHIÊN CỨU/ THẢO LUẬN")
    clone_paragraph(document, h2, "3.1. Kết quả hợp đồng phần mềm")
    clone_paragraph(
        document,
        body,
        "Bảng 2 trả lời RQ1 và phần software-only của RQ2. Các test xác nhận server từ chối tài khoản ngoài home scope, chỉ service credential được ghi telemetry, mock fallback không được lưu thành dữ liệu thật, lệnh lặp có cùng trạng thái không phát xung thừa và scene giữ lỗi theo từng thiết bị. Forecast API cũng từ chối model không hỗ trợ, timestamp sai và retrain giả. Frontend contract chưa đạt hoàn toàn, nên bài báo không tuyên bố toàn bộ UI đã xác minh.",
    )
    add_captioned_table(
        document,
        table_caption,
        table_note,
        "Bảng 2. Bằng chứng kiểm thử ở mức hợp đồng phần mềm",
        SOFTWARE_EVIDENCE,
        "Ghi chú: 'Đạt' nghĩa là các contract hiện hành chạy thành công; không suy ra an toàn điện, latency hoặc usability trên người dùng.",
        [1.25, 2.35, 0.85, 1.15, 1.25],
    )

    clone_paragraph(document, h2, "3.2. Kết quả benchmark dự báo công khai")
    clone_paragraph(
        document,
        body,
        "Bảng 3 trả lời phần định lượng của RQ3. XGBoost đạt MAE 0,486 ± 0,046 kW và RMSE 0,648 ± 0,097 kW, trong khi seasonal naive 24 giờ đạt MAE 0,538 ± 0,118 kW. Mức giảm MAE tương đối khoảng 9,8%. Random Forest đạt MAE 0,492 ± 0,037 kW; chênh lệch khoảng 1,2% giữa hai mô hình học máy nhỏ hơn độ phân tán của kết quả, vì vậy chưa đủ cơ sở tuyên bố XGBoost vượt trội có ý nghĩa thống kê.",
    )
    add_captioned_table(
        document,
        table_caption,
        table_note,
        "Bảng 3. Kết quả test của benchmark rolling-origin trên UCI",
        model_rows(canonical),
        "Ghi chú: Giá trị là mean ± sample SD. Inference đo trong runtime benchmark, không phải latency end-to-end từ app tới PLC.",
        [1.5, 1.05, 1.05, 1.05, 1.0, 1.35],
    )
    clone_paragraph(
        document,
        body,
        "Hình 4 cho thấy h+1 có sai số thấp nhất đối với hai mô hình học máy; từ h+6 đến h+24, MAE và RMSE tăng rồi tương đối ổn định. Seasonal naive 24 giờ ít thay đổi theo horizon nhưng kém XGBoost ở MAE tổng hợp. Mẫu hình này phù hợp với việc lag gần nhất mang nhiều thông tin cho h+1, song không chứng minh mô hình đủ chính xác cho điều khiển tải.",
    )
    add_figure(
        document,
        image_placeholder,
        figure_caption,
        figures_dir / "paper_forecast_horizon_vi.png",
        "Hình 4. Sai số theo horizon trên UCI cho thấy lợi thế rõ nhất ở h+1; error bar là sample SD qua các lần chạy canonical.",
    )

    clone_paragraph(document, h2, "3.3. Evidence gate và phạm vi claim")
    clone_paragraph(
        document,
        body,
        "Bảng 4 tổng hợp quyền công bố. Hai RQ về kiến trúc và lệnh-phản hồi được hỗ trợ ở mức software contract. Forecast được hỗ trợ ở phạm vi benchmark UCI. Ba claim bị khóa là local MFM384 accuracy, real-hardware latency và automatic load shedding. Cơ chế sa thải tải hiện fail-closed; đây là biện pháp an toàn của prototype, không phải kết quả hiệu quả demand response.",
    )
    add_captioned_table(
        document,
        table_caption,
        table_note,
        "Bảng 4. Phân loại bằng chứng và quyền công bố claim",
        EVIDENCE_GATE,
        "Ghi chú: Trạng thái lấy từ evidenceStatus và claimPolicy của canonical_results.json.",
        [1.45, 1.35, 1.25, 1.25, 1.7],
    )

    clone_paragraph(document, h2, "3.4. Thảo luận")
    clone_paragraph(
        document,
        body,
        "Kết quả cho thấy giá trị chính của prototype không nằm ở một model dự báo mới. XGBoost chỉ nhỉnh hơn Random Forest và MAPE cao cho thấy tải hộ gia đình có nhiều thời điểm thấp hoặc khó dự báo. R² khoảng 0,22 cũng giới hạn khả năng giải thích biến thiên. Do chỉ có hai fold, bài báo không thực hiện paired significance test; tăng số rolling fold là điều kiện cần trước khi nâng claim về model.",
    )
    clone_paragraph(
        document,
        body,
        "Ngược lại, việc nối app response với statusTag có ý nghĩa đối với system paper vì nó biến kết quả backend thành trạng thái người dùng quan sát được. Tuy nhiên, loading/success/error/timeout mới được mô tả bằng state transition và test contract. Bài báo không suy diễn rằng giao diện dễ dùng, phản hồi nhanh hay an toàn điện. Để có các claim đó cần user study, timestamp end-to-end, sơ đồ điện, interlock, manual override và raw trial log.",
    )
    clone_paragraph(
        document,
        body,
        "Threats to validity gồm bốn nhóm. Construct validity bị giới hạn vì test contract không đo tiếp điểm và tải thật. Internal validity bị ảnh hưởng bởi số fold ít và artifact model phát cảnh báo khác phiên bản thư viện. External validity thấp vì UCI không đại diện khí hậu, hành vi và thiết bị tại Cần Thơ. Conclusion validity bị giới hạn bởi chênh lệch XGBoost-Random Forest nhỏ và chưa có kiểm định cặp. Evidence gate làm giảm nguy cơ overclaim nhưng không loại bỏ các giới hạn này.",
    )

    clone_paragraph(document, h1, "4. KẾT LUẬN/ ĐỀ XUẤT/ GIẢI PHÁP")
    clone_paragraph(document, h2, "4.1. Kết luận")
    clone_paragraph(
        document,
        body,
        "Nghiên cứu đã định vị nguyên mẫu như một system/architecture paper được kiểm chứng bằng phần mềm và benchmark dữ liệu công khai. Kiến trúc tách user token khỏi telemetry credential, giới hạn thao tác theo nhà, tuần tự hóa PLC I/O và yêu cầu statusTag độc lập. Pipeline dự báo có baseline, rolling-origin và fingerprint; kết quả ủng hộ XGBoost như model được chọn trong canonical run, nhưng không ủng hộ tuyên bố ưu thế mạnh hoặc khả năng áp dụng trực tiếp tại Cần Thơ.",
    )
    clone_paragraph(document, h2, "4.2. Hạn chế và hướng nghiên cứu")
    clone_paragraph(
        document,
        body,
        "Bước tiếp theo là thu dữ liệu MFM384 địa phương đủ dài, tăng số rolling fold và sinh actual-versus-predicted từ prediction artifact chuẩn. Thử nghiệm điều khiển cần tối thiểu quy trình tải an toàn, interlock, emergency stop, manual override, mất mạng, feedback timeout, lệnh đồng thời và recovery. Chỉ sau khi có raw log hợp lệ mới bổ sung phân bố RTT, p95/p99 và success/timeout. Automatic load shedding tiếp tục bị khóa cho đến khi hoàn tất hazard analysis và phê duyệt vận hành.",
    )
    clone_paragraph(document, h2, "4.3. Tuyên bố dữ liệu, mã nguồn và liêm chính")
    clone_paragraph(
        document,
        body,
        "Dữ liệu benchmark công khai tại UCI theo DOI 10.24432/C58K54. Kết quả dùng trong bài nằm trong research/results/canonical/canonical_results.json; mã tái lập gồm research/run_forecast_benchmark.py, research/build_canonical_results.py và các script sinh hình/tài liệu. Raw credential, cấu hình mạng vận hành và dữ liệu hộ gia đình địa phương không được công khai nếu chưa ẩn danh và đánh giá an toàn.",
    )
    clone_paragraph(
        document,
        body,
        "Nghiên cứu không tuyển người tham gia và không thực hiện thử nghiệm trên người hoặc động vật. Tuyên bố đóng góp tác giả (CRediT), nguồn tài trợ và xung đột lợi ích cần được các tác giả điền và xác nhận trước khi nộp. Công cụ AI được dùng để hỗ trợ rà soát diễn đạt, lập trình sinh hình và kiểm tra cấu trúc; tác giả chịu trách nhiệm về số liệu, trích dẫn, quyết định học thuật và bản nộp cuối cùng.",
    )

    clone_paragraph(document, ack, "Lời cảm ơn (nếu có): [BỔ SUNG CƠ QUAN HỖ TRỢ/TÀI TRỢ SAU KHI TÁC GIẢ XÁC NHẬN].")
    clone_paragraph(document, ref_heading, "Tài liệu tham khảo")
    for reference in REFERENCES:
        reference_paragraph = clone_paragraph(document, ref, reference)
        reference_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def validate_inputs(args: argparse.Namespace) -> None:
    for path in (args.template, args.canonical):
        if not path.exists():
            raise FileNotFoundError(path)
    for name in ("paper_evidence_gate_vi.png", "paper_architecture_vi.png", "paper_command_feedback_vi.png", "paper_forecast_horizon_vi.png"):
        if not (args.figures_dir / name).exists():
            raise FileNotFoundError(args.figures_dir / name)
    if len(TITLE_VI.split()) > 20:
        raise ValueError("Vietnamese title exceeds the 20-word template limit")
    for label, abstract in (("Vietnamese", ABSTRACT_VI), ("English", ABSTRACT_EN)):
        count = len(abstract.split())
        if not 150 <= count <= 200:
            raise ValueError(f"{label} abstract has {count} words; template requires 150-200")


def main() -> int:
    args = parse_args()
    validate_inputs(args)
    canonical = json.loads(args.canonical.read_text(encoding="utf-8"))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.template, args.output)
    document = Document(args.output)

    if len(document.tables) < 3 or len(document.paragraphs) < 40:
        raise ValueError("Unexpected journal template structure")
    samples = {
        "h1": document.paragraphs[10],
        "body": document.paragraphs[11],
        "h2": document.paragraphs[16],
        "h3": document.paragraphs[21],
        "table_caption": document.paragraphs[25],
        "table_note": document.paragraphs[26],
        "image": document.paragraphs[29],
        "figure_caption": document.paragraphs[30],
        "ack": document.paragraphs[36],
        "ref_heading": document.paragraphs[37],
        "ref": document.paragraphs[39],
    }

    populate_front_matter(document)
    clear_template_body_after_abstract(document)
    build_paper(document, canonical, args.figures_dir, samples)

    if len(document.sections) != 1 or document.sections[0].start_type != WD_SECTION.NEW_PAGE:
        # A single-section template may report NEW_PAGE as its default start type.
        if len(document.sections) != 1:
            raise ValueError("The revision unexpectedly changed the template section count")
    document.core_properties.title = TITLE_VI.title()
    document.core_properties.subject = "Bản sửa đổi tiếng Việt theo mẫu Tạp chí Khoa học và Công nghệ Cần Thơ"
    document.core_properties.comments = "Experimental claims are generated from research/results/canonical/canonical_results.json."
    document.save(args.output)
    print(args.output)
    print(f"Vietnamese abstract words: {len(ABSTRACT_VI.split())}")
    print(f"English abstract words: {len(ABSTRACT_EN.split())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
