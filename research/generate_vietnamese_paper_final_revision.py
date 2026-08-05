from __future__ import annotations

import argparse
import json
import shutil
from copy import deepcopy
from pathlib import Path
from statistics import fmean

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


TITLE_VI = "THIẾT KẾ MÔ HÌNH GIÁM SÁT NĂNG LƯỢNG VÀ DỰ BÁO PHỤ TẢI CHO PHÒNG THÍ NGHIỆM ĐIỆN CÔNG NGHIỆP"
TITLE_EN = "Design of an Energy Monitoring and Load Forecasting Model for an Industrial Electrical Engineering Laboratory"

REFERENCES = [
    "[1] Shareef H, Ahmed MS, Mohamed A, Al Hassan E. Review on home energy management system considering demand responses, smart technologies, and intelligent controllers. IEEE Access. 2018;6:24498-24509. doi:10.1109/ACCESS.2018.2831917.",
    "[2] Gomes I, Bot K, Ruano MG, Ruano A. Recent techniques used in home energy management systems: a review. Energies. 2022;15(8):2866. doi:10.3390/en15082866.",
    "[3] Motta LL, Ferreira LCBC, Cabral TW, et al. General overview and proof of concept of a smart home energy management system architecture. Electronics. 2023;12(21):4453. doi:10.3390/electronics12214453.",
    "[4] Cao Z, Han X, Lyons W, O'Rourke F. Energy management optimisation using a combined Long Short-Term Memory recurrent neural network-Particle Swarm Optimisation model. Journal of Cleaner Production. 2021;326:129246. doi:10.1016/j.jclepro.2021.129246.",
    "[5] Semmelmann L, Henni S, Weinhardt C. Load forecasting for energy communities: a novel LSTM-XGBoost hybrid model based on smart meter data. Energy Informatics. 2022;5(Suppl 1):24. doi:10.1186/s42162-022-00212-9.",
    "[6] Siemens AG. SIMATIC S7-1200 programmable controller system manual. Version 4.6, document A5E02486680-AP. Nuremberg: Siemens AG; 2022.",
    "[7] Modbus Organization. MODBUS application protocol specification V1.1b3. Hopkinton (MA): Modbus Organization; 2012.",
    "[8] Modbus Organization. MODBUS over serial line specification and implementation guide V1.02. Hopkinton (MA): Modbus Organization; 2006.",
    "[9] Stouffer K, Pease M, Tang C, et al. Guide to operational technology (OT) security. NIST SP 800-82 Rev. 3. Gaithersburg (MD): National Institute of Standards and Technology; 2023. doi:10.6028/NIST.SP.800-82r3.",
    "[10] Hebrail G, Berard A. Individual household electric power consumption [dataset]. Irvine (CA): UCI Machine Learning Repository; 2006. doi:10.24432/C58K54.",
    "[11] Breiman L. Random forests. Machine Learning. 2001;45:5-32. doi:10.1023/A:1010933404324.",
    "[12] Chen T, Guestrin C. XGBoost: a scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD; 2016. p. 785-794. doi:10.1145/2939672.2939785.",
    "[13] Hyndman RJ, Koehler AB. Another look at measures of forecast accuracy. International Journal of Forecasting. 2006;22(4):679-688. doi:10.1016/j.ijforecast.2006.03.001.",
    "[14] Bergmeir C, Benítez JM. On the use of cross-validation for time series predictor evaluation. Information Sciences. 2012;191:192-213. doi:10.1016/j.ins.2011.12.028.",
    "[15] Karuna G, et al. Smart energy management: real-time prediction and optimization for IoT-enabled smart homes. Cogent Engineering. 2024;11(1):2390674. doi:10.1080/23311916.2024.2390674.",
    "[16] Muqtadir A, Li B, Ying Z, et al. Nowcasting the next hour of residential load using boosting ensemble machines. Scientific Reports. 2025;15:7157. doi:10.1038/s41598-025-91767-6.",
    "[17] Khan MF, et al. A hybrid reinforcement learning framework for adaptive multi-horizon electricity load forecasting: the DWRNet approach. Computers & Electrical Engineering. 2026;131:110926. doi:10.1016/j.compeleceng.2025.110926.",
    "[18] Ma Y, Chen X, Wang L, Yang J. Investigation of smart home energy management system for demand response application. Frontiers in Energy Research. 2021;9:772027. doi:10.3389/fenrg.2021.772027.",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the electrical-engineering-focused Vietnamese revision using the official CTUT template.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def vi_number(value: float, decimals: int = 3) -> str:
    return f"{value:.{decimals}f}".replace(".", ",")


def vi_int(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def metric_map(canonical: dict) -> dict[str, dict]:
    return {row["model"]: row for row in canonical["forecast"]["models"] if row.get("split") == "test"}


def abstracts(canonical: dict) -> tuple[str, str]:
    rows = metric_map(canonical)
    xgb = rows["xgboost"]
    rf = rows["random_forest"]
    dataset = canonical["forecast"]["dataset"]
    vi = (
        "Nghiên cứu trình bày mô hình giám sát năng lượng cho phòng thí nghiệm Điện công nghiệp. MFM384 dự kiến đo điện áp, dòng điện, công suất và điện năng; PLC S7-1200 tiếp nhận dữ liệu, điều khiển cơ cấu chấp hành và phản hồi trạng thái. "
        "Ứng dụng Web/di động và mô-đun dự báo 24 giờ hỗ trợ quan sát phụ tải, cảnh báo quota và đề xuất vận hành. "
        "Chuỗi điều khiển kiểm tra quyền, tuần tự hóa lệnh, đọc phản hồi độc lập và từ chối an toàn khi thiếu điều kiện. "
        f"Phần dự báo được đánh giá trên bộ dữ liệu UCI bằng expanding rolling-origin với {dataset['rolling_folds']} fold và {len(dataset['random_seeds'])} seed. "
        f"XGBoost đạt MAE {vi_number(xgb['mae_kw'])} ± {vi_number(xgb['mae_kw_std'])} kW; Random Forest đạt {vi_number(rf['mae_kw'])} ± {vi_number(rf['mae_kw_std'])} kW. "
        "Chênh lệch nhỏ hơn độ phân tán qua các lần chạy nên chưa chứng minh mô hình vượt trội. "
        "Kết quả hỗ trợ kiến trúc, kiểm thử phần mềm và benchmark công khai. "
        "Do chưa có log MFM384, thử tải thật và số liệu độ trễ App-PLC, nghiên cứu không tuyên bố độ chính xác địa phương, tiết kiệm hoặc sa thải tải tự động."
    )
    en = (
        "This study presents an energy monitoring model for an Industrial Electrical Engineering laboratory. "
        "The intended measurement chain uses an MFM384 meter for voltage, current, power, and energy, while a Siemens S7-1200 PLC acquires data, commands switching devices, and reports operating status. "
        "A Web/mobile application and a 24-hour load forecasting module support load observation, quota warnings, and operating recommendations. "
        "The control path checks authorization, serializes commands, reads an independent feedback status, and fails safely when required conditions are unavailable. "
        f"Forecasting is evaluated on the public UCI dataset by expanding rolling-origin validation with {dataset['rolling_folds']} folds and {len(dataset['random_seeds'])} seeds. "
        f"XGBoost obtains an MAE of {xgb['mae_kw']:.3f} ± {xgb['mae_kw_std']:.3f} kW, and Random Forest obtains {rf['mae_kw']:.3f} ± {rf['mae_kw_std']:.3f} kW. "
        "Their difference is smaller than run-to-run dispersion and does not establish model superiority. "
        "Available evidence supports the proposed architecture, software tests, and the public-data benchmark. "
        "Because accepted MFM384 logs, physical-load trials, and App-to-PLC latency measurements are unavailable, no claim is made about local accuracy, energy savings, or automatic load shedding."
    )
    return vi, en


def set_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def set_runs(paragraph: Paragraph, values: list[str]) -> None:
    if len(paragraph.runs) < len(values):
        raise ValueError("Template run structure is not compatible")
    for index, run in enumerate(paragraph.runs):
        run.text = values[index] if index < len(values) else ""


def insert_before_sectpr(document: Document, element) -> None:
    body = document._body._element
    if body.sectPr is None:
        body.append(element)
    else:
        body.sectPr.addprevious(element)


def clone_paragraph(document: Document, sample: Paragraph, text: str, *, keep_next: bool = False) -> Paragraph:
    element = deepcopy(sample._p)
    nodes = element.xpath(".//w:t")
    if nodes:
        nodes[0].text = text
        for node in nodes[1:]:
            node.text = ""
    insert_before_sectpr(document, element)
    paragraph = Paragraph(element, document._body)
    if not nodes:
        paragraph.add_run(text)
    paragraph.paragraph_format.keep_with_next = keep_next
    return paragraph


def set_alt_text(inline_shape, alt_text: str) -> None:
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", alt_text[:80])


def add_figure(document: Document, image_sample: Paragraph, caption_sample: Paragraph, path: Path, caption: str, alt_text: str, width: float = 6.65) -> None:
    paragraph = clone_paragraph(document, image_sample, "", keep_next=True)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(shape, alt_text)
    caption_paragraph = clone_paragraph(document, caption_sample, caption)
    caption_paragraph.paragraph_format.keep_together = True
    # Keep the caption with the preceding image via the image paragraph's
    # keep-with-next setting, but do not chain the caption to the next body
    # paragraph/table; that can force a large figure onto a new page.
    caption_paragraph.paragraph_format.keep_with_next = False


def set_cell(cell, text: str, *, header: bool, size: float) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or len(text) < 28 else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    run.bold = header
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_geometry(table, widths_in: list[float]) -> None:
    widths = [round(value * 1440) for value in widths_in]
    total = round(sum(widths_in) * 1440)
    widths[-1] += total - sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:type"), "dxa")
    # Match the default 120-DXA start-cell margin so the visible table border
    # aligns with the template body text in Word's fixed-layout rendering.
    indent.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))


def add_table(document: Document, rows: list[list[str]], widths: list[float], *, size: float = 8.2) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for row_index, values in enumerate(rows):
        tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for column_index, value in enumerate(values):
            set_cell(table.rows[row_index].cells[column_index], value, header=row_index == 0, size=size)


def captioned_table(document: Document, samples: dict[str, Paragraph], caption: str, rows: list[list[str]], note: str, widths: list[float], *, size: float = 8.2) -> None:
    clone_paragraph(document, samples["table_caption"], caption, keep_next=True)
    add_table(document, rows, widths, size=size)
    clone_paragraph(document, samples["table_note"], note)


def clear_after_front_matter(document: Document) -> None:
    current = document.tables[1]._tbl.getnext()
    body = document._body._element
    while current is not None:
        following = current.getnext()
        if current.tag != qn("w:sectPr"):
            body.remove(current)
        current = following


def related_work_rows() -> list[list[str]]:
    return [
        ["Công trình", "Đối tượng điện", "Đo lường/điều khiển", "Dự báo/quản lý tải", "Bằng chứng"],
        ["Shareef et al. [1]", "Phụ tải gia đình", "Bộ điều khiển thông minh", "Đáp ứng nhu cầu, dịch chuyển tải", "Tổng quan"],
        ["Motta et al. [3]", "Thiết bị điện trong HEMS", "Ổ cắm thông minh, bộ điều khiển", "Giám sát, cảnh báo, quản lý", "Proof-of-concept"],
        ["Ma et al. [18]", "Thiết bị gia dụng", "Điều phối thiết bị", "Demand response", "Mô phỏng/kịch bản"],
        ["Karuna et al. [15]", "Phụ tải nhà thông minh", "IoT gateway", "Gradient boosting", "Dữ liệu công khai"],
        ["Muqtadir et al. [16]", "Phụ tải dân dụng", "Không thuộc trọng tâm", "Boosting, dự báo 1 giờ", "Đánh giá theo thời gian"],
        ["Khan et al. [17]", "Phụ tải hệ thống điện", "EMS/SCADA", "Dự báo đa chân trời", "Rolling-origin"],
        ["Công trình này", "Mô hình tải phòng thí nghiệm", "MFM384, PLC, relay/contactor", "Dự báo 24 giờ, quota", "Test phần mềm + UCI; thử tải chờ"],
    ]


def software_rows(canonical: dict) -> list[list[str]]:
    evidence = canonical.get("softwareEvidence") or {}
    suites = {row["suite"]: row for row in evidence.get("suites") or []}
    permissions = evidence.get("claimPermissions") or {}

    def combined(names: list[str]) -> tuple[str, str, str]:
        selected = [suites[name] for name in names if name in suites]
        passed = sum(int(row.get("passed") or 0) for row in selected)
        total = sum(int(row.get("total") or 0) for row in selected)
        hashes = "/".join(str(row.get("logSha256") or "")[:8] for row in selected)
        return "+".join(names), f"{passed}/{total}", hashes

    definitions = [
        ("Kiểm tra quyền điều khiển", ["backend"], "authorizationAndHomeScope", "Không thay thế đánh giá an toàn hệ thống"),
        ("Tách kênh dữ liệu đo", ["backend"], "telemetryCredentialIsolation", "Chưa xác nhận dữ liệu MFM384 thật"),
        ("Phản hồi PLC và từ chối an toàn", ["backend"], "plcFeedbackAndFailClosed", "Chưa xác nhận tiếp điểm hoặc độ trễ phần cứng"),
        ("Dữ liệu dự báo và chống rò rỉ", ["forecast", "research"], "forecastArtifactContract", "Chỉ áp dụng cho benchmark UCI"),
        ("Trạng thái hiển thị trên ứng dụng", ["frontend_contract", "room_presentation"], "applicationStateContract", "Không phải nghiên cứu trải nghiệm người dùng"),
    ]
    rows = [["Nhóm tuyên bố", "Suite/ID", "Đạt/tổng", "Hash log", "Trạng thái", "Giới hạn diễn giải"]]
    for label, names, permission, limit in definitions:
        suite_id, ratio, hashes = combined(names)
        allowed = bool(permissions.get(permission))
        rows.append([label, suite_id, ratio, hashes, "Đạt" if allowed else "Chưa đạt", limit])
    return rows


def model_rows(canonical: dict) -> list[list[str]]:
    names = {
        "persistence": "Persistence",
        "seasonal_naive_24h": "Seasonal naive 24 h",
        "seasonal_naive_168h": "Seasonal naive 168 h",
        "random_forest": "Random Forest",
        "xgboost": "XGBoost",
    }
    rows = [["Mô hình", "MAE (kW)", "RMSE (kW)", "MAPE ngưỡng (%)", "R²", "Inference (ms/mẫu)", "n"]]
    for item in canonical["forecast"]["models"]:
        if item.get("split") != "test":
            continue
        inference = item["inference_ms_per_sample"]
        inference_text = "<0,001" if inference < 0.001 else f"{vi_number(inference)} ± {vi_number(item['inference_ms_per_sample_std'])}"
        rows.append(
            [
                names[item["model"]],
                f"{vi_number(item['mae_kw'])} ± {vi_number(item['mae_kw_std'])}",
                f"{vi_number(item['rmse_kw'])} ± {vi_number(item['rmse_kw_std'])}",
                f"{vi_number(item['mape_percent'], 1)} ± {vi_number(item['mape_percent_std'], 1)}",
                f"{vi_number(item['r2'])} ± {vi_number(item['r2_std'])}",
                inference_text,
                str(item["run_count"]),
            ]
        )
    return rows


def claim_rows(canonical: dict) -> list[list[str]]:
    policy = canonical["claimPolicy"]
    software = (canonical.get("softwareEvidence") or {}).get("claimPermissions") or {}
    return [
        ["Kịch bản vận hành", "Trạng thái", "Bằng chứng hiện có", "Kết luận được phép"],
        ["Giám sát V/I/P/E", "Thiết kế", "Chưa có raw log MFM384", "Chỉ mô tả chuỗi đo dự kiến"],
        ["Dự báo phụ tải 24 giờ", "Đã benchmark", "UCI; 3 fold × 3 seed", "Công bố trong phạm vi UCI"],
        ["Cảnh báo gần quota", "Triển khai một phần", "API/cảnh báo phần mềm", "Chưa có thử nghiệm tác động"],
        ["Khuyến nghị giảm/dịch tải", "Thiết kế", "Chưa có log người dùng", "Không gọi là tối ưu hóa"],
        ["Lệnh App/Web đến PLC", "Kiểm thử phần mềm", "Backend và trạng thái lệnh", "Chưa suy ra đóng cắt thật"],
        ["Phản hồi/từ chối/timeout", "Kiểm thử phần mềm", str(bool(software.get("plcFeedbackAndFailClosed"))).lower(), "Chưa suy ra RTT phần cứng"],
        ["Điều khiển tải công suất lớn", "Chưa thử tải", "Không có log tiếp điểm", "Không công bố kết quả"],
        ["Sa thải tải tự động", "Đang khóa", str(bool(policy["allowAutomaticLoadSheddingClaims"])).lower(), "Chỉ là thiết kế an toàn"],
        ["Manual override/khôi phục", "Thiết kế", "Chưa có thử nghiệm sự cố", "Đưa vào kế hoạch kiểm chứng"],
        ["Dự báo có thời tiết", "Hướng phát triển", "Benchmark chưa có biến thời tiết", "Không đưa vào kết quả hiện tại"],
    ]


def populate_front_matter(document: Document, canonical: dict) -> None:
    abstract_vi, abstract_en = abstracts(canonical)
    set_text(document.paragraphs[0], TITLE_VI)
    set_text(document.paragraphs[1], "")
    set_text(document.paragraphs[3], "[HỌ VÀ TÊN TÁC GIẢ]1, [HỌ VÀ TÊN ĐỒNG TÁC GIẢ]2")
    set_text(document.paragraphs[4], "1 [KHOA/ĐƠN VỊ, TRƯỜNG ĐẠI HỌC KỸ THUẬT - CÔNG NGHỆ CẦN THƠ]")
    set_text(document.paragraphs[5], "2 [ĐƠN VỊ CỦA ĐỒNG TÁC GIẢ, NẾU CÓ]")
    set_text(document.paragraphs[6], "Tác giả liên hệ: [EMAIL TÁC GIẢ CHỊU TRÁCH NHIỆM]")
    set_runs(
        document.tables[0].cell(0, 0).paragraphs[0],
        [
            "Thông tin chung\n",
            "Ngày nhận bài: dd/mm/yyyy\nNgày nhận bài sửa: dd/mm/yyyy\nNgày duyệt đăng: dd/mm/yyyy\n\n",
            "Từ khóa:\n",
            "dự báo phụ tải, giám sát năng lượng, MFM384, PLC S7-1200, quản lý phụ tải",
        ],
    )
    set_text(document.tables[0].cell(0, 1).paragraphs[1], abstract_vi)
    set_runs(
        document.tables[1].cell(0, 0).paragraphs[0],
        [
            "Title:\n",
            TITLE_EN + "\n\n",
            "Keywords:\n",
            "energy monitoring, load forecasting, MFM384, load management, Siemens S7-1200 PLC",
        ],
    )
    set_text(document.tables[1].cell(0, 1).paragraphs[1], abstract_en)


def build_body(document: Document, canonical: dict, figures: Path, samples: dict[str, Paragraph]) -> None:
    dataset = canonical["forecast"]["dataset"]
    rows = metric_map(canonical)
    xgb, rf, sn24 = rows["xgboost"], rows["random_forest"], rows["seasonal_naive_24h"]
    preprocessing = dataset["preprocessing"]
    software = canonical.get("softwareEvidence") or {}
    commit = str(software.get("commitSha") or "không xác định")[:12]

    def p(kind: str, text: str) -> Paragraph:
        return clone_paragraph(document, samples[kind], text)

    p("h1", "1. ĐẶT VẤN ĐỀ")
    p("body", "Trong phòng thí nghiệm Điện công nghiệp, một mô hình quản lý năng lượng cần thực hiện đồng thời ba nhiệm vụ: đo các đại lượng điện, điều khiển phụ tải và dự báo xu hướng tiêu thụ. HEMS và đáp ứng nhu cầu thường dùng thông tin phụ tải để cảnh báo, dịch chuyển hoặc cắt giảm tải có kiểm soát [1], [2], [18]. Tuy nhiên, quyết định điều khiển chỉ có ý nghĩa khi dữ liệu đo, trạng thái PLC và cơ cấu chấp hành được liên kết rõ ràng.")
    p("body", "Các kiến trúc HEMS đã kết hợp thiết bị đo, bộ điều khiển, lưu trữ và giao diện giám sát [3]. Nhánh dự báo dùng dữ liệu công-tơ thông minh, mô hình học máy và đặc trưng theo thời gian để hỗ trợ vận hành [4], [5], [15]-[17]. Khoảng trống của đề tài không nằm ở việc đề xuất thuật toán mới, mà ở cách tổ chức chuỗi MFM384–PLC–phụ tải–dự báo–quota trong một mô hình phòng thí nghiệm và giới hạn kết luận theo bằng chứng thực tế.")
    p("body", "Bài báo vì vậy tập trung vào thiết kế đo lường và điều khiển ở góc nhìn Điện–Điều khiển. Ứng dụng Web/di động và dịch vụ AI chỉ là lớp hỗ trợ hiển thị, dự báo và gửi yêu cầu; chúng không thay thế phép đo tại đồng hồ, phản hồi tiếp điểm hoặc thử nghiệm an toàn trên tải thật.")
    p("h2", "1.1. Câu hỏi nghiên cứu")
    p("body", "RQ1: Chuỗi thiết bị nào liên kết phép đo V/I/P/E từ MFM384, PLC S7-1200, cơ cấu đóng cắt và giao diện giám sát mà vẫn phản hồi đúng trạng thái vận hành?")
    p("body", "RQ2: Random Forest và XGBoost dự báo phụ tải h+1 đến h+24 như thế nào so với các đường cơ sở khi đánh giá theo thứ tự thời gian?")
    p("body", "RQ3: Những kịch bản quota, cảnh báo, điều khiển và sa thải phụ tải nào đã có bằng chứng, và những kịch bản nào mới dừng ở mức thiết kế an toàn?")
    p("h2", "1.2. Đóng góp và phạm vi")
    p("body", "Bài báo có ba đóng góp: thiết kế chuỗi đo–điều khiển từ MFM384 đến PLC và phụ tải; xây dựng logic phản hồi lệnh, quota và trạng thái an toàn; đánh giá tái lập mô-đun dự báo 24 giờ trên dữ liệu UCI với hai mô hình học máy và ba đường cơ sở. Phần cứng hiện chưa có raw log được chấp nhận, nên kết quả chỉ gồm kiểm thử phần mềm và benchmark công khai, không gồm sai số đo tại Cần Thơ, độ trễ App–PLC, mức tiết kiệm hoặc sa thải tải tự động.")

    p("h1", "2. CƠ SỞ LÝ THUYẾT/ MÔ HÌNH/ PHƯƠNG PHÁP NGHIÊN CỨU")
    p("h2", "2.1. Công trình liên quan và khoảng trống")
    p("body", "Bảng 1 đối chiếu các công trình theo đối tượng điện, chuỗi đo–điều khiển, chức năng quản lý tải và loại bằng chứng. Các tổng quan [1], [2] làm rõ vai trò của đáp ứng nhu cầu; Motta et al. [3] cung cấp proof-of-concept HEMS; Ma et al. [18] trình bày điều phối thiết bị; các nghiên cứu [4], [5], [15]-[17] tập trung vào dự báo và tối ưu. So với các hướng này, bài báo giới hạn ở mô hình phòng thí nghiệm có PLC và công-tơ, trong đó mọi kết luận phần cứng phải đi kèm log đo và trạng thái cơ cấu chấp hành.")
    captioned_table(document, samples, "Bảng 1. Đối chiếu các công trình đại diện gần đề tài", related_work_rows(), "Ghi chú: Bảng so sánh phạm vi và loại bằng chứng, không xếp hạng chất lượng công trình.", [1.12, 1.30, 1.42, 1.50, 1.42], size=7.8)

    p("h2", "2.2. Mô hình đo lường và chuỗi thiết bị điện")
    p("body", "Chuỗi đo dự kiến bắt đầu tại MFM384, thu điện áp, dòng điện, công suất và điện năng. Dữ liệu được truyền qua RS-485/Modbus RTU theo bảng thanh ghi và hệ số tỷ lệ của đúng phiên bản thiết bị. Đặc tả Modbus quy định mô hình thanh ghi, mã hàm và xử lý ngoại lệ [7], trong khi hướng dẫn đường truyền nối tiếp quy định khung RTU và yêu cầu triển khai RS-485 [8]. Bài chưa công bố địa chỉ thanh ghi, chu kỳ lấy mẫu hoặc cấp chính xác vì chưa có manual MFM384 đúng mã hàng và raw log thử nghiệm.")
    p("body", "PLC S7-1200 nhận giá trị đã chuẩn hóa, giám sát điều kiện cho phép và phát lệnh tới relay hoặc contactor để đóng cắt tải. Tài liệu S7-1200 [6] được dùng cho vai trò bộ điều khiển và giao tiếp, không thay thế sơ đồ đấu dây của mô hình. Trước khi thử tải phải xác nhận nguồn điều khiển, khả năng chịu dòng của cơ cấu chấp hành, bảo vệ ngắn mạch/quá tải, tiếp điểm phản hồi, nút dừng khẩn và cách ly giữa mạch lực với mạch điều khiển.")
    p("body", "Hình 1 thể hiện hai nhánh: nhánh đo từ MFM384 đến kho dữ liệu và mô-đun dự báo; nhánh điều khiển từ App/Web qua PLC đến phụ tải rồi trả trạng thái. Các kiểm soát truy cập chỉ bảo vệ đường gửi lệnh; an toàn điện vẫn phụ thuộc liên động PLC, bảo vệ phần cứng và quy trình vận hành theo nguyên tắc phòng thủ nhiều lớp cho hệ thống OT [9].")
    add_figure(document, samples["image"], samples["figure_caption"], figures / "final_architecture_vi.png", "Hình 1. Chuỗi đo–điều khiển của mô hình: MFM384 cung cấp dữ liệu điện, PLC S7-1200 điều khiển phụ tải và phản hồi trạng thái; phần cứng thật còn chờ raw log.", "Sơ đồ chuỗi điện gồm MFM384, PLC S7-1200, relay hoặc contactor, phụ tải, kho dữ liệu, mô-đun dự báo và giao diện vận hành.", width=5.45)

    p("h2", "2.3. Điều khiển PLC, phản hồi và yêu cầu an toàn")
    p("body", "Hình 2 mô tả một chu kỳ điều khiển. Ứng dụng gửi trạng thái mong muốn; lớp biên kiểm tra quyền và chuyển yêu cầu tới gateway; gateway đọc trạng thái hiện tại rồi chỉ phát xung khi mục tiêu chưa đạt. Kết quả không lấy trực tiếp từ lệnh ghi mà được xác nhận bằng statusTag độc lập. Hết thời gian chờ, mất kết nối hoặc sai quyền đều trả trạng thái lỗi thay vì hiển thị thành công.")
    p("body", "Với tải công suất lớn, lệnh phần mềm chưa đủ để cho phép đóng cắt. Logic PLC cần thêm liên động, thời gian giữ tối thiểu, chống đóng cắt liên tục, giới hạn số lần thao tác, nút dừng khẩn và quyền điều khiển tại chỗ. Khi dữ liệu đo lỗi thời, trạng thái phản hồi không nhất quán hoặc mất truyền thông, hệ thống phải giữ trạng thái an toàn và yêu cầu người vận hành kiểm tra.")
    add_figure(document, samples["image"], samples["figure_caption"], figures / "final_command_feedback_vi.png", "Hình 2. Chuỗi App–PLC xác nhận lệnh bằng statusTag và phân biệt thành công, từ chối, lỗi kết nối và timeout; hình không biểu diễn độ trễ phần cứng.", "Sơ đồ tuần tự từ ứng dụng đến PLC S7-1200, có vòng lặp đọc statusTag và các trạng thái kết quả.", width=6.0)

    p("h2", "2.4. Quota, khuyến nghị và kịch bản sa thải phụ tải")
    p("body", "Quota là giới hạn điện năng tích lũy, không phải ngưỡng công suất tức thời. Tại thời điểm t, hệ thống so sánh điện năng đã dùng E_t với quota Q và ước lượng điện năng 24 giờ từ dãy công suất dự báo. Khi E_t cộng phần điện năng dự kiến vượt Q, ứng dụng chỉ tạo cảnh báo và danh sách tải có thể giảm hoặc dịch chuyển theo mức ưu tiên. Cơ chế này là hỗ trợ quyết định; chưa phải bài toán tối ưu vì chưa có hàm mục tiêu, ràng buộc và baseline điều phối.")
    p("body", "Sa thải phụ tải chủ động được giữ ở trạng thái thiết kế. Trước khi mở chế độ tự động cần xác định tải không được cắt, thứ tự ưu tiên, ngưỡng công suất, hysteresis, thời gian vượt ngưỡng, thời gian giữ sau khi cắt, manual override, emergency stop và quy tắc khôi phục. Mỗi lần thử phải ghi timestamp, V/I/P/E trước–sau, tải mục tiêu, nguyên nhân tác động, trạng thái tiếp điểm và thao tác phục hồi. Nếu thiếu các dữ liệu này, bài không được công bố mức tiết kiệm hoặc tỷ lệ sa thải thành công.")

    p("h2", "2.5. Dữ liệu và mô hình dự báo phụ tải")
    p("body", f"Bộ dữ liệu UCI [10] gồm {vi_int(dataset['source_raw_rows'])} bản ghi một phút; {vi_int(dataset['source_power_missing_rows'])} giá trị công suất bị thiếu ({vi_number(dataset['source_power_missing_percent'], 2)}%). Pipeline giữ {vi_int(dataset['raw_rows'])} dòng trong {dataset.get('max_history_days', 730)} ngày gần nhất, tổng hợp thành {vi_int(dataset['hourly_rows'])} mốc giờ và tạo {vi_int(dataset['supervised_rows'])} mẫu. Có {vi_int(preprocessing['hourly_power_rows_not_observed'])} mốc giờ không quan sát công suất, trong đó {vi_int(preprocessing['hourly_power_rows_causally_filled'])} mốc được điền tiến từ quá khứ với giới hạn {preprocessing['causal_fill_limit_hours']} giờ. Giá trị điền chỉ làm đặc trưng lịch sử; cửa sổ có nhãn không quan sát bị loại.")
    model_config = dataset["model_config"]
    p("body", f"Đặc trưng gồm thời gian chu kỳ; độ trễ công suất 1–24 giờ, 144–168 giờ và các mốc 48, 72, 96, 120, 336 giờ; thống kê trượt 3–168 giờ; cùng độ trễ của điện áp, dòng điện, công suất phản kháng và ba kênh sub-metering. Mọi thống kê trượt đều dịch một bước để không dùng tương lai. Random Forest [11] dùng {model_config['random_forest']['n_estimators']} cây, max_depth={model_config['random_forest']['max_depth']}; XGBoost [12] dùng {model_config['xgboost']['n_estimators']} cây, max_depth={model_config['xgboost']['max_depth']} và learning_rate={str(model_config['xgboost']['learning_rate']).replace('.', ',')}. Mỗi chân trời h+1 đến h+24 có một bộ ước lượng trực tiếp. Benchmark hiện không có biến thời tiết; tích hợp thời tiết chỉ là hướng phát triển.")

    p("h2", "2.6. Thiết kế đánh giá và khả năng tái lập")
    fold_text = "; ".join(
        f"fold {index}: train {vi_int(item['train']['rows'])} ({item['train']['start']}–{item['train']['end']}), validation {vi_int(item['val']['rows'])}, test {vi_int(item['test']['rows'])}"
        for index, item in enumerate(dataset["fold_ranges"], start=1)
    )
    p("body", f"Đánh giá expanding rolling-origin gồm {dataset['rolling_folds']} fold: {fold_text}. Các seed {', '.join(str(value) for value in dataset['random_seeds'])} được dùng cho Random Forest và XGBoost; vì vậy mỗi mô hình học máy có {dataset['rolling_folds'] * len(dataset['random_seeds'])} lần chạy, mỗi baseline có {dataset['rolling_folds']} lần. Mô hình được chọn bằng validation MAE trung bình, không dùng test để chọn.")
    p("body", "Đường cơ sở gồm persistence, seasonal naive 24 giờ và seasonal naive 168 giờ. MAE là chỉ số chính; RMSE nhấn mạnh sai số lớn; R² mô tả phần phương sai được giải thích. MAPE dùng mẫu số tối thiểu 0,2 kW vì chỉ số này mất ổn định khi tải gần 0 [13]. Mean ± sample SD được tính qua các lần chạy. Ba fold chưa đủ để thực hiện kiểm định cặp đáng tin cậy, nên bài không báo cáo p-value [14].")
    p("body", f"Dữ liệu, cấu hình mô hình, timestamp, lệnh kiểm thử và SHA-256 log được hợp nhất trong tệp kết quả chuẩn hóa; commit tham chiếu là {commit}. Cách tổ chức này giúp tái lập benchmark, nhưng không biến thay đổi chưa commit hoặc log phần cứng còn thiếu thành bằng chứng.")

    p("h1", "3. KẾT QUẢ NGHIÊN CỨU/ THẢO LUẬN")
    p("h2", "3.1. Kết quả kiểm thử chuỗi giám sát–điều khiển")
    frontend_suite = next(row for row in software.get("suites", []) if row.get("suite") == "frontend_contract")
    p("body", f"Bảng 2 ghi số ca đạt và mã băm của log kiểm thử. Các ca backend hỗ trợ kiểm tra quyền điều khiển, tách kênh dữ liệu đo, phản hồi PLC độc lập và từ chối an toàn ở mức phần mềm. Bộ kiểm tra trạng thái ứng dụng chỉ đạt {frontend_suite['passed']}/{frontend_suite['total']} ca nên chưa thể kết luận toàn bộ giao diện xử lý đúng mọi lỗi. Các kết quả này không thay cho thử tiếp điểm, đo điện hoặc đánh giá an toàn tải.")
    captioned_table(document, samples, "Bảng 2. Bằng chứng kiểm thử phần mềm cho chuỗi giám sát–điều khiển", software_rows(canonical), "Ghi chú: Hash là tám ký tự đầu của SHA-256 log. “Đạt” chỉ áp dụng cho ca phần mềm, không đồng nghĩa phần cứng đã hoạt động.", [1.15, 1.05, 0.60, 0.72, 0.68, 2.45], size=7.7)

    p("h2", "3.2. Kết quả benchmark dự báo công khai")
    fold_runs = canonical["forecast"].get("runs") or {}
    xgb_gain = 100.0 * (sn24["mae_kw"] - xgb["mae_kw"]) / sn24["mae_kw"]
    rf_gap = 100.0 * abs(rf["mae_kw"] - xgb["mae_kw"]) / rf["mae_kw"]
    p("body", f"Bảng 3 cho thấy XGBoost đạt MAE {vi_number(xgb['mae_kw'])} ± {vi_number(xgb['mae_kw_std'])} kW, thấp hơn seasonal naive 24 giờ khoảng {vi_number(xgb_gain, 1)}%. Random Forest đạt {vi_number(rf['mae_kw'])} ± {vi_number(rf['mae_kw_std'])} kW; chênh lệch tương đối giữa hai mô hình học máy là {vi_number(rf_gap, 1)}%, nhỏ so với độ phân tán. Vì vậy kết quả chỉ cho phép chọn mô hình theo MAE trên tập validation, không chứng minh XGBoost vượt trội có ý nghĩa thống kê.")
    def fold_mae_text(model: str) -> str:
        grouped: dict[int, list[float]] = {}
        for run in fold_runs.get(model) or []:
            grouped.setdefault(int(run["fold"]), []).append(float(run["test"]["mae"]))
        return ", ".join(
            f"F{fold}={vi_number(fmean(values))} kW" for fold, values in sorted(grouped.items())
        )
    p("body", f"MAE test trung bình theo từng fold của XGBoost là {fold_mae_text('xgboost')}; Random Forest là {fold_mae_text('random_forest')}. Báo cáo theo từng fold cho thấy sai số thay đổi theo giai đoạn và tránh hiểu sample SD như chỉ do seed ngẫu nhiên.")
    captioned_table(document, samples, "Bảng 3. Kết quả test của benchmark expanding rolling-origin trên UCI", model_rows(canonical), "Ghi chú: Mean ± sample SD; n là số lần chạy thật. MAPE dùng mẫu số tối thiểu 0,2 kW. Inference là thời gian runtime benchmark, không phải latency App-PLC.", [1.25, 0.92, 0.92, 1.02, 0.92, 1.15, 0.57], size=7.7)
    p("body", "Hình 3 trình bày bốn chân trời đại diện h+1, h+6, h+12 và h+24, không phải toàn bộ 24 điểm. Màu, marker và kiểu đường đồng thời phân biệt năm phương pháp; error bar là sample SD và chú giải ghi n riêng cho từng mô hình/baseline. Trục bắt đầu từ 0 để tránh phóng đại chênh lệch nhỏ.")
    add_figure(document, samples["image"], samples["figure_caption"], figures / "final_forecast_horizon_vi.png", "Hình 3. Sai số tại h+1, h+6, h+12 và h+24 trên UCI; hai panel lần lượt là MAE và RMSE, error bar là sample SD, n được ghi trong chú giải.", "Hai biểu đồ đường MAE và RMSE theo bốn chân trời dự báo cho XGBoost, Random Forest và ba baseline, kèm error bar và số lần chạy.", width=6.55)

    p("h2", "3.3. Đối chiếu kịch bản vận hành và bằng chứng")
    p("body", "Bảng 4 phân biệt rõ kịch bản đã có benchmark, kịch bản mới được kiểm thử bằng phần mềm và kịch bản chỉ ở mức thiết kế. Giám sát V/I/P/E, đóng cắt tải công suất lớn, manual override và sa thải tải vẫn thiếu raw log phần cứng. Dự báo có thời tiết cũng chưa thuộc benchmark hiện tại. Vì vậy ảnh giao diện hoặc ảnh tủ điện chưa được dùng làm hình khoa học khi chưa có dữ liệu và chú giải kiểm chứng đi kèm.")
    captioned_table(document, samples, "Bảng 4. Trạng thái bằng chứng của các kịch bản vận hành", claim_rows(canonical), "Ghi chú: “Thiết kế” không được diễn giải thành kết quả thực nghiệm; mọi thử tải phải có log đo và trạng thái tiếp điểm.", [1.80, 1.25, 2.00, 1.70], size=7.4)

    p("h2", "3.4. Thảo luận và nguy cơ đối với tính hợp lệ")
    p("body", "XGBoost và Random Forest gần tương đương vì cùng khai thác tập độ trễ và thống kê trượt, trong khi phụ tải hộ gia đình biến động mạnh. Ba fold và ba seed giúp quan sát độ phân tán nhưng chưa đủ để khẳng định ưu thế mô hình. Các nghiên cứu boosting hoặc đa chân trời [16], [17] chỉ dùng để đối chiếu cách đánh giá; không so sánh trực tiếp trị số vì khác dữ liệu và chân trời dự báo.")
    p("body", "Về mặt phương pháp luận học máy, tồn tại hiện tượng lệch phân phối dữ liệu (distribution shift) giữa tập dữ liệu công khai UCI (thu thập tại hộ gia đình dân dụng ở Pháp) và đặc tính phụ tải thực tế của phòng thí nghiệm hoặc hộ gia đình tại Cần Thơ. Đồ thị phụ tải, biên độ công suất và tính chu kỳ của hai môi trường này có sự khác biệt bản chất. Do đó, việc đạt chỉ số sai số MAE/RMSE tốt trên benchmark UCI chỉ chứng minh tính đúng đắn của quy trình xử lý dữ liệu và thuật toán, không đảm bảo mô hình sẽ đạt độ chính xác tương đương khi triển khai thực tế trên dữ liệu từ cảm biến MFM384 tại địa phương khi chưa được huấn luyện lại (fine-tune).")
    p("body", "Tính hợp lệ của chuỗi đo–điều khiển còn thấp vì kiểm thử phần mềm không quan sát tiếp điểm, dòng tải và sai số đồng hồ. Benchmark được cải thiện nhờ xử lý thiếu theo chiều thời gian, nhưng vẫn phụ thuộc cấu hình cố định. Việc công bố rõ ranh giới bằng chứng giúp tránh tuyên bố quá mức nhưng không thay thế thí nghiệm vật lý.")

    p("h1", "4. KẾT LUẬN/ ĐỀ XUẤT/ GIẢI PHÁP")
    p("h2", "4.1. Kết luận")
    p("body", "Bài báo đã xây dựng khung kỹ thuật cho mô hình giám sát năng lượng phòng thí nghiệm: MFM384 cung cấp đại lượng điện, PLC S7-1200 xử lý điều kiện đóng cắt, giao diện gửi yêu cầu và nhận phản hồi, còn mô-đun AI dự báo phụ tải 24 giờ. Benchmark UCI cho thấy Random Forest và XGBoost cải thiện MAE so với các đường cơ sở nhưng gần tương đương nhau. Kết quả chưa xác nhận sai số đo, độ trễ, tiết kiệm năng lượng hoặc sa thải tải trên mô hình thật.")
    p("h2", "4.2. Hạn chế và hướng nghiên cứu")
    p("body", "Bước tiếp theo là xác nhận manual và bảng thanh ghi MFM384 đúng mã thiết bị; hoàn thiện sơ đồ mạch lực/mạch điều khiển; thu dữ liệu V/I/P/E có timestamp; so sánh với thiết bị tham chiếu; và thử relay/contactor dưới giám sát. Thử nghiệm điều khiển phải bao gồm interlock, emergency stop, manual override, mất mạng, timeout, lệnh đồng thời và khôi phục. Dự báo cần được huấn luyện lại trên dữ liệu địa phương và công bố đồ thị actual–predicted trước khi đánh giá khả năng ứng dụng.")
    p("h2", "4.3. Tuyên bố dữ liệu, đạo đức và liêm chính")
    p("body", "Dữ liệu công khai theo DOI 10.24432/C58K54 [10]. Kết quả dùng trong bài nằm trong research/results/canonical/canonical_results.json; mã tái lập gồm script benchmark, thu log kiểm thử, xây tệp kết quả, sinh hình và sinh DOCX. Nghiên cứu không tuyển người tham gia và không thử nghiệm trên người hoặc động vật. Thông tin xác thực, cấu hình mạng vận hành và dữ liệu địa phương không được công khai khi chưa ẩn danh và đánh giá an toàn.")
    p("body", "Thông tin đóng góp tác giả theo CRediT, nguồn tài trợ, xung đột lợi ích và tác giả liên hệ còn là placeholder, phải được xác nhận trước khi nộp. Công cụ AI được dùng để hỗ trợ rà soát diễn đạt, kiểm thử mã, sinh sơ đồ và kiểm tra cấu trúc; tác giả chịu trách nhiệm về số liệu, trích dẫn, quyết định học thuật và bản nộp cuối cùng.")

    clone_paragraph(document, samples["ack"], "Lời cảm ơn (nếu có): [BỔ SUNG CƠ QUAN HỖ TRỢ/TÀI TRỢ SAU KHI TÁC GIẢ XÁC NHẬN].")
    clone_paragraph(document, samples["ref_heading"], "Tài liệu tham khảo")
    for reference in REFERENCES:
        paragraph = clone_paragraph(document, samples["ref"], reference)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def validate(args: argparse.Namespace, canonical: dict) -> None:
    for path in (args.template, args.canonical):
        if not path.exists():
            raise FileNotFoundError(path)
    for name in ("final_architecture_vi.png", "final_command_feedback_vi.png", "final_forecast_horizon_vi.png"):
        if not (args.figures_dir / name).exists():
            raise FileNotFoundError(args.figures_dir / name)
    if len(TITLE_VI.split()) > 20:
        raise ValueError("Vietnamese title exceeds 20 words")
    for label, abstract in zip(("Vietnamese", "English"), abstracts(canonical)):
        count = len(abstract.split())
        if not 150 <= count <= 200:
            raise ValueError(f"{label} abstract has {count} words")
    if not canonical.get("softwareEvidence"):
        raise ValueError("softwareEvidence is required for Table 2")


def main() -> int:
    args = parse_args()
    canonical = json.loads(args.canonical.read_text(encoding="utf-8"))
    validate(args, canonical)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.template, args.output)
    document = Document(args.output)
    if len(document.tables) < 3 or len(document.paragraphs) < 40:
        raise ValueError("Unexpected journal template structure")
    samples = {
        "h1": document.paragraphs[10],
        "body": document.paragraphs[11],
        "h2": document.paragraphs[16],
        "h3": document.paragraphs[21],
        "table_caption": document.paragraphs[25],
        "table_note": document.paragraphs[26],
        "image": document.paragraphs[29],
        "figure_caption": document.paragraphs[30],
        "ack": document.paragraphs[36],
        "ref_heading": document.paragraphs[37],
        "ref": document.paragraphs[39],
    }
    populate_front_matter(document, canonical)
    clear_after_front_matter(document)
    build_body(document, canonical, args.figures_dir, samples)
    document.core_properties.title = TITLE_VI.title()
    document.core_properties.subject = "Bản sửa đổi theo trọng tâm Điện - Điều khiển và mẫu Tạp chí Khoa học và Công nghệ Cần Thơ"
    document.core_properties.comments = "Forecast and software results are generated from canonical_results.json; physical hardware claims remain pending."
    document.save(args.output)
    vi_abstract, en_abstract = abstracts(canonical)
    print(args.output)
    print(f"Title words: {len(TITLE_VI.split())}")
    print(f"Vietnamese abstract words: {len(vi_abstract.split())}")
    print(f"English abstract words: {len(en_abstract.split())}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
