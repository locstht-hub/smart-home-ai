from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from statistics import fmean

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph

import generate_vietnamese_paper_final_revision as base


TITLE_VI = "THIẾT KẾ HỆ THỐNG GIÁM SÁT NĂNG LƯỢNG IOT HỖ TRỢ DỰ BÁO PHỤ TẢI BẰNG TRÍ TUỆ NHÂN TẠO"
TITLE_EN = "Design of an IoT Energy Monitoring System Supporting Artificial-Intelligence-Based Load Forecasting"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the objective-aligned Vietnamese CTUT journal paper.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def abstracts(canonical: dict) -> tuple[str, str]:
    rows = base.metric_map(canonical)
    rf, xgb = rows["random_forest"], rows["xgboost"]
    dataset = canonical["forecast"]["dataset"]
    vi = (
        "Nghiên cứu trình bày thiết kế hệ thống giám sát năng lượng IoT cho mô hình phụ tải tại phòng thí nghiệm Điện công nghiệp. "
        "Dự báo phụ tải 24 giờ là lõi AI; tài khoản, phân quyền theo hộ, Web Dashboard, ứng dụng di động và cổng PLC tạo lớp triển khai. "
        "Chuỗi xử lý đi từ đo lường, lưu trữ, dự báo đến quota, khuyến nghị và điều khiển có phản hồi. "
        "Kiến trúc dùng credential riêng cho telemetry, I/O tuần tự, statusTag và từ chối an toàn. "
        f"Benchmark UCI được đánh giá bằng expanding rolling-origin với {dataset['rolling_folds']} fold và {len(dataset['random_seeds'])} seed. "
        f"Random Forest đạt MAE {base.vi_number(rf['mae_kw'])} ± {base.vi_number(rf['mae_kw_std'])} kW; XGBoost đạt {base.vi_number(xgb['mae_kw'])} ± {base.vi_number(xgb['mae_kw_std'])} kW. "
        "Chênh lệch nhỏ so với độ phân tán nên không chứng minh mô hình vượt trội. "
        "Bằng chứng xác nhận hợp đồng phần mềm và benchmark; kiểm thử giao diện chưa đạt đầy đủ. "
        "Do thiếu raw log MFM384, PLC và tải thật, bài không tuyên bố độ chính xác, độ trễ phần cứng, tiết kiệm hoặc sa thải tự động. "
        "Kết quả xác lập quy trình triển khai, ranh giới bằng chứng và kế hoạch kiểm chứng mô hình."
    )
    en = (
        "This study presents the design of an IoT energy monitoring system for a small-scale load model in an Industrial Electrical Engineering laboratory. "
        "The 24-hour load forecast is the AI core, while personal accounts, household-scoped authorization, a Web dashboard, a mobile application, and a PLC gateway form its deployment layer. "
        "The workflow proceeds from measurement and storage to forecasting, quota assessment, recommendations, and feedback-aware control. "
        "The software separates telemetry credentials, serializes I/O, uses an independent statusTag, and fails closed. "
        f"The public UCI benchmark is evaluated by expanding rolling-origin validation with {dataset['rolling_folds']} folds and {len(dataset['random_seeds'])} seeds. "
        f"Random Forest obtains an MAE of {rf['mae_kw']:.3f} ± {rf['mae_kw_std']:.3f} kW, while XGBoost obtains {xgb['mae_kw']:.3f} ± {xgb['mae_kw_std']:.3f} kW. "
        "Their difference is small relative to run-to-run dispersion and does not establish model superiority. "
        "Available evidence supports selected software behavior contracts and a public-data benchmark, although the interface test suite is not yet fully passing. "
        "Because no accepted raw logs from the MFM384 meter, PLC, and physical loads are available, the paper makes no claim about accuracy in Can Tho, hardware latency, energy savings, or automatic load shedding. "
        "The contribution is therefore an evidence-bounded implementation workflow and a verification plan for completing the laboratory prototype."
    )
    return vi, en


def populate_front_matter(document: Document, canonical: dict) -> None:
    abstract_vi, abstract_en = abstracts(canonical)
    base.set_text(document.paragraphs[0], TITLE_VI)
    base.set_text(document.paragraphs[1], "")
    base.set_text(document.paragraphs[3], "[HỌ VÀ TÊN TÁC GIẢ]1, [HỌ VÀ TÊN ĐỒNG TÁC GIẢ]2")
    base.set_text(document.paragraphs[4], "1 [KHOA/ĐƠN VỊ, TRƯỜNG ĐẠI HỌC KỸ THUẬT - CÔNG NGHỆ CẦN THƠ]")
    base.set_text(document.paragraphs[5], "2 [ĐƠN VỊ CỦA ĐỒNG TÁC GIẢ, NẾU CÓ]")
    base.set_text(document.paragraphs[6], "Tác giả liên hệ: [EMAIL TÁC GIẢ CHỊU TRÁCH NHIỆM]")
    base.set_runs(
        document.tables[0].cell(0, 0).paragraphs[0],
        [
            "Thông tin chung\n",
            "Ngày nhận bài: dd/mm/yyyy\nNgày nhận bài sửa: dd/mm/yyyy\nNgày duyệt đăng: dd/mm/yyyy\n\n",
            "Từ khóa:\n",
            "dự báo phụ tải, giám sát năng lượng, Internet vạn vật, quota năng lượng, phân quyền người dùng",
        ],
    )
    base.set_text(document.tables[0].cell(0, 1).paragraphs[1], abstract_vi)
    base.set_runs(
        document.tables[1].cell(0, 0).paragraphs[0],
        [
            "Title:\n",
            TITLE_EN + "\n\n",
            "Keywords:\n",
            "energy monitoring, Internet of Things, load forecasting, energy quota, user authorization",
        ],
    )
    base.set_text(document.tables[1].cell(0, 1).paragraphs[1], abstract_en)


def system_evidence_rows(canonical: dict) -> list[list[str]]:
    suites = {row["suite"]: row for row in canonical["softwareEvidence"]["suites"]}
    return [
        ["Khối chức năng", "Đầu vào → đầu ra", "Trạng thái và bằng chứng", "Phạm vi kết luận"],
        ["Đo lường và PLC", "V, I, P, E từ MFM384 → tag PLC/telemetry", "Có gateway/cấu hình; chưa có raw trial", "Kiến trúc mục tiêu, chưa phải kết quả đo"],
        ["Tài khoản và RBAC", "Token, homeId, role → phạm vi thao tác/audit", f"Backend {suites['backend']['passed']}/{suites['backend']['total']}; hash {suites['backend']['logSha256'][:8]}", "Hợp đồng phần mềm đã kiểm thử"],
        ["Web Dashboard và ứng dụng", "Telemetry, dự báo, quota → hiển thị/lệnh/trạng thái", f"Frontend {suites['frontend_contract']['passed']}/{suites['frontend_contract']['total']}; room {suites['room_presentation']['passed']}/{suites['room_presentation']['total']}", "Đã triển khai một phần; chưa xác nhận toàn bộ UI"],
        ["Dự báo phụ tải", "Lịch sử theo giờ → vectơ P̂(h+1…h+24)", f"Forecast {suites['forecast']['passed']}/{suites['forecast']['total']}; research {suites['research']['passed']}/{suites['research']['total']}", "Được công bố trong phạm vi UCI"],
        ["Quota và khuyến nghị", "E tích lũy, quota, P̂, ưu tiên → cảnh báo/gợi ý", "Có CRUD, guard và cảnh báo; chưa có log tác động", "Chức năng hỗ trợ quyết định"],
        ["Sa thải tự động", "Nguy cơ vượt ngưỡng → lệnh cắt tải an toàn", "Cờ an toàn khóa; chưa có thuật toán kW/thử tải", "Không được công bố là đã hoạt động"],
    ]


def objective_rows(canonical: dict) -> list[list[str]]:
    evidence = canonical["evidenceStatus"]
    policy = canonical["claimPolicy"]
    software = canonical["softwareEvidence"]["claimPermissions"]
    return [
        ["Mục tiêu dự án", "Phương pháp trong bài", "Bằng chứng hiện có", "Kết luận được phép"],
        ["Đo lường và IoT từ xa", "MFM384–PLC–gateway–API", f"Raw phần cứng: {str(evidence['realHardwareLatency']).lower()}", "Kiến trúc mục tiêu; cần thử nghiệm PTN"],
        ["Quản lý tài khoản và phân quyền", "Token, home scope, RBAC, audit", f"Contract: {str(software['authorizationAndHomeScope']).lower()}", "Được hỗ trợ ở mức kiểm thử phần mềm"],
        ["Điều khiển trên Web và di động", "Lệnh có phản hồi và trạng thái nguyên nhân", f"Application-state contract: {str(software['applicationStateContract']).lower()}", "Có triển khai; chưa xác nhận đầy đủ"],
        ["Dự báo phụ tải bằng AI", "RF/XGBoost và ba baseline", f"UCI benchmark: {str(evidence['publicDatasetBenchmark']).lower()}", "Chỉ khái quát cho benchmark công khai"],
        ["Khuyến nghị theo quota/ưu tiên", "Dự báo + quota + chính sách ưu tiên", "CRUD/guard/cảnh báo; chưa có thử nghiệm tác động", "Cơ chế hỗ trợ quyết định, không gọi là tối ưu hóa"],
        ["Sa thải tải chủ động", "Interlock và thuật toán công suất tức thời dự kiến", f"Permission: {str(policy['allowAutomaticLoadSheddingClaims']).lower()}", "Chưa được tuyên bố; là công việc tương lai"],
    ]


def build_body(document: Document, canonical: dict, figures: Path, samples: dict[str, Paragraph]) -> None:
    dataset = canonical["forecast"]["dataset"]
    preprocessing = dataset["preprocessing"]
    rows = base.metric_map(canonical)
    rf, xgb, sn24 = rows["random_forest"], rows["xgboost"], rows["seasonal_naive_24h"]
    software = canonical["softwareEvidence"]
    suites = {row["suite"]: row for row in software["suites"]}

    def p(style: str, text: str, *, keep_next: bool = False) -> Paragraph:
        return base.clone_paragraph(document, samples[style], text, keep_next=keep_next)

    p("h1", "1. ĐẶT VẤN ĐỀ")
    p("body", "Các phòng thí nghiệm Điện công nghiệp cần quan sát công suất, điện năng và trạng thái phụ tải để hỗ trợ giảng dạy về đo lường, điều khiển và quản lý nhu cầu. Hệ thống quản lý năng lượng gia đình thường kết hợp thiết bị đo, truyền thông, giao diện người dùng và bộ điều khiển thông minh [1], [2]. Tuy nhiên, một mô hình phòng thí nghiệm còn phải phân biệt rõ phần đã triển khai trong mã, phần đã được kiểm thử bằng phần mềm và phần chỉ có thể xác nhận sau khi đấu nối thiết bị thật.")
    p("body", "Đề tài hướng tới xây dựng mô hình giám sát năng lượng ứng dụng AI để dự báo phụ tải trên quy mô nhỏ: đo lường qua MFM384, điều khiển bằng PLC S7-1200, kết nối IoT, quản lý tài khoản và phân quyền, vận hành qua Web Dashboard/ứng dụng di động, đồng thời theo dõi quota để hỗ trợ khuyến nghị và sa thải tải chủ động. Trong bài báo, dự báo phụ tải là lõi AI; IoT, tài khoản, giao diện, quota và phản hồi PLC là lớp triển khai để biến kết quả dự báo thành thông tin có thể quan sát và hành động. Khi chưa có raw log MFM384–PLC–phụ tải, đóng góp được định vị là thiết kế hệ thống và đánh giá phần mềm/dữ liệu công khai, không phải đánh giá phần cứng tại Cần Thơ.")
    p("body", "Khoảng trống nghiên cứu được đặt ở sự liên kết có thể truy vết giữa bốn lớp: nhận dạng người dùng, điều khiển có phản hồi, dự báo phụ tải và chính sách quota. Nhiều kiến trúc HEMS/edge tập trung vào middleware [3], [4], giao diện [5] hoặc cầu nối PLC [6], trong khi các nghiên cứu dự báo và điều phối năng lượng thường giả định sẵn dữ liệu hiện trường [7], [8]. Với một đồ án phòng thí nghiệm, cần thêm cổng kiểm soát tuyên bố để tránh dùng kết quả benchmark công khai làm bằng chứng cho phần cứng địa phương.")
    p("body", "Bốn câu hỏi nghiên cứu được đặt ra: RQ1—kiến trúc nào liên kết đo lường, PLC, IoT, Web/mobile và AI mà vẫn chỉ rõ ranh giới bằng chứng? RQ2—cơ chế tài khoản, phân quyền và phản hồi lệnh được hỗ trợ đến mức nào bởi kiểm thử phần mềm? RQ3—Random Forest và XGBoost dự báo 24 giờ như thế nào so với baseline trên UCI khi đánh giá theo thời gian? RQ4—những điều kiện nào còn thiếu trước khi có thể tuyên bố khuyến nghị theo quota và sa thải tải tự động an toàn?")
    p("body", "Quy trình triển khai được tổ chức theo chuỗi: MFM384 đo V/I/P/E → PLC/gateway thu và chuẩn hóa telemetry → API lưu lịch sử → mô hình sinh vectơ dự báo h+1…h+24 → dịch vụ quota đối chiếu điện năng tích lũy và ưu tiên tải → App/Web hiển thị cảnh báo hoặc khuyến nghị → lệnh đã cấp quyền được gửi tới PLC và xác nhận bằng statusTag. Đóng góp của bài gồm kiến trúc mục tiêu, hợp đồng lệnh–phản hồi, benchmark dự báo tái lập và ma trận đối chiếu mục tiêu–bằng chứng.")

    p("h1", "2. PHƯƠNG PHÁP NGHIÊN CỨU")
    p("h2", "2.1. Công trình liên quan và định vị nghiên cứu")
    p("body", "Các tổng quan gần đây cho thấy HEMS bao gồm giám sát, phản hồi nhu cầu, giao diện và điều khiển thông minh [1], [2]. Kiến trúc proof-of-concept và microservices edge cung cấp nền tảng tích hợp [3], [4]; giao diện dùng chung hỗ trợ người dùng quan sát và thao tác năng lượng [5]. Ở lớp công nghiệp, PLC S7-1200 và Modbus yêu cầu tuân thủ đúng mô hình tag/thanh ghi [9], [10], trong khi NIST và OWASP nhấn mạnh phân đoạn OT, định danh thiết bị, phân quyền tối thiểu và kiểm tra ứng dụng [11]–[13]. Các nghiên cứu mới về IoT, kiểm soát truy cập, nowcasting và dự báo đa chân trời [19]–[24] được dùng để đối chiếu cấu trúc, không dùng để so sánh trực tiếp trị số do khác dữ liệu và thiết kế thí nghiệm.")
    base.captioned_table(document, samples, "Bảng 1. Định vị công trình so với các hướng nghiên cứu liên quan", base.related_work_rows(), "Ghi chú: Bảng tập trung vào phạm vi hệ thống và loại bằng chứng; không xếp hạng chất lượng công trình.", [1.15, 1.05, 1.15, 1.55, 1.75], size=7.7)

    p("h2", "2.2. Kiến trúc hệ thống và mô hình phòng thí nghiệm")
    p("body", "Kiến trúc mục tiêu và luồng chuyển giao giữa các khối được trình bày ở Hình 1.")
    base.add_figure(document, samples["image"], samples["figure_caption"], figures / "final_architecture_vi.png", "Hình 1. Kiến trúc mục tiêu App/Web–API–dự báo–PLC/MFM384 và ranh giới bằng chứng; phần cứng thật chưa có raw trial.", "Sơ đồ kiến trúc mục tiêu và ranh giới bằng chứng từ App/Web qua API, dự báo, gateway và telemetry tới PLC S7-1200, MFM384 cùng phụ tải phòng thí nghiệm.", width=6.0)
    p("body", "Ở nhánh dữ liệu, MFM384 cung cấp V/I/P/E, PLC/gateway chuẩn hóa telemetry, kho dữ liệu tạo lịch sử và dịch vụ dự báo sinh chuỗi h+1 đến h+24. Ở nhánh quyết định, API kết hợp dự báo với quota và mức ưu tiên để trả cảnh báo/khuyến nghị cho App/Web. Khi người dùng thao tác, API kiểm tra phạm vi hộ và RBAC; gateway tuần tự hóa I/O và xác nhận bằng statusTag độc lập. Telemetry dùng service credential riêng. Các đường MFM384–PLC và PLC–phụ tải vẫn là mục tiêu tích hợp vì chưa có raw trial được chấp nhận.")
    p("body", "Bảng 2 cụ thể hóa đầu vào, đầu ra và trạng thái của từng khối, tương tự vai trò của bảng tham số trong một bài điều khiển: người đọc có thể lần theo dữ liệu từ phép đo đến quyết định. Trạng thái triển khai được lấy từ nguồn canonical và log kiểm thử, không suy ra từ hình minh họa.")
    base.captioned_table(document, samples, "Bảng 2. Chuỗi đầu vào–đầu ra và trạng thái triển khai của hệ thống", system_evidence_rows(canonical), "Ghi chú: P̂ là công suất dự báo; E là điện năng. Hash là tám ký tự đầu của SHA-256 log. Mock không phải phép đo MFM384/PLC thật.", [1.35, 2.30, 1.70, 1.30], size=7.7)

    p("h2", "2.3. Đo lường, PLC và kết nối IoT")
    p("body", "Lớp đo lường dự kiến đọc điện áp, dòng điện, công suất và điện năng từ MFM384 qua Modbus RTU, sau đó ánh xạ vào PLC S7-1200 [9], [10]. Gateway chỉ phát xung lệnh khi trạng thái mục tiêu chưa đạt, giữ command tag tách khỏi statusTag và thực hiện polling có timeout. Kiến trúc này phù hợp để thu log gồm timestamp, định danh tải, lệnh, trạng thái phản hồi và giá trị đo; tuy nhiên bài chưa báo cáo tần số lấy mẫu, độ trễ, tỷ lệ thành công hoặc sai số cảm biến vì chưa có raw log phần cứng.")
    p("body", "Để xác nhận RQ1 ở mức thực nghiệm, thử nghiệm tương lai phải ghi đồng bộ thời gian ở App, API, gateway, PLC và đồng hồ; kiểm tra mất mạng, timeout, lệnh đồng thời, manual override và emergency stop. Chỉ khi tệp thử nghiệm chứa cấu hình thiết bị, mốc thời gian, đơn vị đo và checksum được chấp nhận mới có thể báo cáo median/p95/p99 hoặc tỷ lệ thành công.")

    p("h2", "2.4. Quản lý người dùng, Web Dashboard và ứng dụng di động")
    p("body", "Tài khoản cá nhân được gắn với phạm vi hộ để ngăn người dùng đọc hoặc điều khiển thiết bị ngoài phạm vi. API kiểm tra token, role, homeId và quyền thiết bị trước khi gửi lệnh; telemetry dùng credential dịch vụ riêng. Audit log lưu tác nhân, mục tiêu và kết quả. Cách phân tách này bám theo nguyên tắc đặc quyền tối thiểu, định danh rõ tài sản IoT và kiểm soát lớp OT [11]–[13].")
    p("body", "Web Dashboard và ứng dụng di động có vai trò tương đương ở lớp nghiệp vụ: hiển thị telemetry, dự báo, mức quota và trạng thái thao tác. Phản hồi của ứng dụng cần phân biệt đang xử lý, thành công khi feedback khớp, bị từ chối, lỗi kết nối và timeout; không được đổi trạng thái chỉ dựa trên việc API đã nhận yêu cầu. Bằng chứng hiện tại hỗ trợ backend 16/16 ca, nhưng frontend contract chỉ đạt 16/20 ca. Vì vậy bài mô tả khả năng xử lý phản hồi như một hợp đồng thiết kế đã được kiểm thử một phần, không tuyên bố usability hay độ ổn định đầy đủ.")
    base.add_figure(document, samples["image"], samples["figure_caption"], figures / "final_command_feedback_vi.png", "Hình 2. Chuỗi App/Web–API–PLC biểu diễn kiểm tra tài khoản, phạm vi, vai trò, quota, polling statusTag và trả trạng thái có nguyên nhân; hình không biểu diễn độ trễ phần cứng.", "Sơ đồ tuần tự từ App hoặc Web qua API có RBAC và quota đến PLC gateway và S7-1200, kết thúc bằng trạng thái verified, denied, error hoặc timeout.")

    p("h2", "2.5. Dữ liệu và mô hình dự báo phụ tải")
    p("body", f"Benchmark dùng bộ Individual Household Electric Power Consumption của UCI [14] với {base.vi_int(dataset['source_raw_rows'])} dòng nguồn; {base.vi_int(dataset['source_power_missing_rows'])} giá trị công suất thiếu, tương đương {base.vi_number(dataset['source_power_missing_percent'], 2)}%. Pipeline giữ tối đa {dataset.get('max_history_days', 730)} ngày gần nhất, tạo {base.vi_int(dataset['hourly_rows'])} mốc giờ và {base.vi_int(dataset['supervised_rows'])} mẫu supervised. Có {base.vi_int(preprocessing['hourly_power_rows_not_observed'])} mốc giờ không quan sát; chỉ {base.vi_int(preprocessing['hourly_power_rows_causally_filled'])} mốc được điền tiến từ quá khứ trong giới hạn {preprocessing['causal_fill_limit_hours']} giờ. Cửa sổ có target không quan sát bị loại để tránh biến dữ liệu thiếu thành nhãn giả.")
    config = dataset["model_config"]
    p("body", f"Đặc trưng gồm thời gian chu kỳ, các lag ngắn/dài và thống kê trượt đã shift một bước. Random Forest [15] dùng {config['random_forest']['n_estimators']} cây với max_depth={config['random_forest']['max_depth']}; XGBoost [16] dùng {config['xgboost']['n_estimators']} cây, max_depth={config['xgboost']['max_depth']} và learning_rate={str(config['xgboost']['learning_rate']).replace('.', ',')}. Mỗi chân trời h+1 đến h+24 có estimator trực tiếp; đầu ra là vectơ 24 giá trị công suất P̂ để lớp quota sử dụng. Persistence, seasonal naive 24 giờ và seasonal naive 168 giờ là baseline bắt buộc.")

    p("h2", "2.6. Cơ chế quota và khuyến nghị quản lý năng lượng")
    p("body", "Bài sử dụng đúng thuật ngữ “cơ chế khuyến nghị quản lý năng lượng dựa trên dự báo, quota và mức ưu tiên phụ tải”. Đây chưa phải bài toán tối ưu hóa vì chưa xác lập đầy đủ hàm mục tiêu, biến quyết định, ràng buộc, solver, baseline điều phối và kết quả thực nghiệm. Quota hiện được quản lý bằng API CRUD, guard và cảnh báo Dashboard.")
    p("body", "Logic thiết kế gồm sáu bước: (1) lấy điện năng tích lũy E_t và quota Q; (2) kiểm tra chất lượng artifact dự báo; (3) tính điện năng 24 giờ Ê_24 = ΣP̂_hΔt; (4) tính mức thiếu hụt dự kiến R = E_t + Ê_24 − Q; (5) nếu R > 0, xếp các tải được phép điều chỉnh theo mức ưu tiên để tạo danh sách trì hoãn/giảm sử dụng; (6) hiển thị nguyên nhân và chờ người dùng phê duyệt trước khi gửi lệnh. Vì cửa sổ dự báo chỉ dài 24 giờ, R là chỉ báo rủi ro gần hạn chứ không phải dự báo đầy đủ đến cuối tháng.")

    p("h2", "2.7. Thiết kế sa thải phụ tải chủ động")
    p("body", "Sa thải phụ tải chủ động được mô tả như một thiết kế an toàn cho giai đoạn thử nghiệm tiếp theo, không phải chức năng đã được xác nhận. Cờ an toàn trong mã đang khóa khả năng bật tự động vì routine theo điện năng tháng không thay thế thuật toán ra quyết định theo công suất tức thời. Điều kiện kích hoạt dự kiến phải kết hợp ngưỡng công suất, thời gian vượt ngưỡng và chất lượng telemetry; dự báo hoặc quota không được tự mình phát lệnh cắt tải.")
    p("body", "Trước khi mở cờ tự động cần xác lập danh sách tải không được cắt, thứ tự ưu tiên, giới hạn công suất, hysteresis, thời gian giữ tối thiểu, interlock, emergency stop, manual override và rollback. Hệ thống phải từ chối lệnh khi telemetry lỗi thời, statusTag không nhất quán, mất kết nối hoặc thiếu quyền; mỗi thử nghiệm phải ghi timestamp, tải mục tiêu, lý do cắt, trạng thái trước–sau và thao tác phục hồi. Chỉ raw log trên tải thật có giám sát mới cho phép báo cáo tỷ lệ thành công hoặc năng lượng dịch chuyển. Nguyên tắc fail-closed hiện có là bằng chứng phần mềm, không phải kết quả sa thải tải.")

    p("h2", "2.8. Thiết kế đánh giá và khả năng tái lập")
    folds = "; ".join(f"fold {i}: train {base.vi_int(item['train']['rows'])}, validation {base.vi_int(item['val']['rows'])}, test {base.vi_int(item['test']['rows'])}" for i, item in enumerate(dataset["fold_ranges"], 1))
    p("body", f"Đánh giá expanding rolling-origin gồm {dataset['rolling_folds']} fold ({folds}) và các seed {', '.join(str(v) for v in dataset['random_seeds'])}. Mô hình được chọn theo validation MAE, không dùng test để chọn. MAE là chỉ số chính; RMSE nhấn mạnh sai số lớn; R² mô tả phương sai giải thích. MAPE chỉ tính với mẫu số tối thiểu 0,2 kW do tính bất ổn gần 0 [17]. Mean ± sample SD được báo cáo qua các lần chạy; ba fold chưa đủ để diễn giải p-value ghép cặp đáng tin cậy [18].")
    p("body", "Kết quả dự báo, trạng thái bằng chứng, log kiểm thử và SHA-256 được hợp nhất trong canonical_results.json. Cách tổ chức này hỗ trợ provenance và tự động hóa tái lập theo thực hành MLOps [22]. Tệp canonical đồng thời khóa các quyền tuyên bố: cho phép kết quả hợp đồng phần mềm và benchmark công khai, nhưng chặn độ chính xác địa phương, độ trễ phần cứng và sa thải tự động.")

    p("h1", "3. KẾT QUẢ NGHIÊN CỨU VÀ THẢO LUẬN")
    p("h2", "3.1. Kết quả xây dựng lớp phần mềm và quản lý người dùng")
    p("body", f"Backend đạt {suites['backend']['passed']}/{suites['backend']['total']} ca, forecast {suites['forecast']['passed']}/{suites['forecast']['total']} và research {suites['research']['passed']}/{suites['research']['total']}. Các kết quả này hỗ trợ RQ2 ở mức kiểm thử phần mềm đối với phạm vi hộ, RBAC, credential telemetry, phản hồi PLC và từ chối an toàn. Tuy nhiên, bộ verify toàn kho chỉ đạt {suites['verify']['passed']}/{suites['verify']['total']} và frontend contract đạt {suites['frontend_contract']['passed']}/{suites['frontend_contract']['total']}; do đó không thể kết luận toàn bộ ứng dụng đã vượt qua xác minh. Kết quả cũng không thay thế penetration test hoặc thử nghiệm người dùng.")

    p("h2", "3.2. Kết quả benchmark dự báo công khai")
    xgb_gain = 100.0 * (sn24["mae_kw"] - xgb["mae_kw"]) / sn24["mae_kw"]
    model_gap = 100.0 * abs(rf["mae_kw"] - xgb["mae_kw"]) / rf["mae_kw"]
    p("body", f"Bảng 3 cho thấy Random Forest đạt MAE {base.vi_number(rf['mae_kw'])} ± {base.vi_number(rf['mae_kw_std'])} kW và XGBoost đạt {base.vi_number(xgb['mae_kw'])} ± {base.vi_number(xgb['mae_kw_std'])} kW. XGBoost thấp hơn seasonal naive 24 giờ khoảng {base.vi_number(xgb_gain, 1)}%, nhưng chênh lệch giữa hai mô hình học máy chỉ khoảng {base.vi_number(model_gap, 2)}% và nhỏ hơn độ phân tán qua các lần chạy. Vì vậy RQ3 được trả lời theo hướng cả hai mô hình cải thiện baseline theo MAE tổng hợp, nhưng chưa có bằng chứng một mô hình vượt trội có ý nghĩa thống kê.")
    base.captioned_table(document, samples, "Bảng 3. Kết quả test của benchmark expanding rolling-origin trên UCI", base.model_rows(canonical), "Ghi chú: Mean ± sample SD; n là số lần chạy. MAPE dùng mẫu số tối thiểu 0,2 kW. Inference là thời gian runtime benchmark, không phải latency App–PLC.", [1.25, 0.92, 0.92, 1.02, 0.92, 1.15, 0.57], size=7.7)
    runs = canonical["forecast"].get("runs") or {}

    def fold_mae(model: str) -> str:
        grouped: dict[int, list[float]] = {}
        for run in runs.get(model) or []:
            grouped.setdefault(int(run["fold"]), []).append(float(run["test"]["mae"]))
        return ", ".join(f"F{fold}={base.vi_number(fmean(values))} kW" for fold, values in sorted(grouped.items()))

    p("body", f"MAE test theo fold của Random Forest là {fold_mae('random_forest')}; của XGBoost là {fold_mae('xgboost')}. Hình 3 chỉ trình bày bốn chân trời đại diện h+1, h+6, h+12 và h+24, với trục bắt đầu từ 0 và error bar là sample SD. Không có đường actual–predicted vì nguồn canonical hiện chưa công bố prediction artifact theo timestamp; việc không dựng dữ liệu giả bảo toàn khả năng kiểm chứng.")
    base.add_figure(document, samples["image"], samples["figure_caption"], figures / "final_forecast_horizon_vi.png", "Hình 3. MAE và RMSE tại h+1, h+6, h+12 và h+24 trên UCI cho hai mô hình học máy và ba baseline; error bar là sample SD.", "Hai biểu đồ MAE và RMSE theo bốn chân trời dự báo cho Random Forest, XGBoost và ba baseline, có error bar và số lần chạy.", width=6.55)

    p("h2", "3.3. Trạng thái bằng chứng đo lường và điều khiển phần cứng")
    p("body", "RQ1 hiện chỉ được trả lời ở mức kiến trúc. Không có raw log MFM384, PLC S7-1200 và phụ tải thật trong nguồn bằng chứng được chấp nhận; tệp latency chỉ có header. Vì vậy không có số liệu để tính độ trễ end-to-end, tỷ lệ thành công, độ chính xác tại phòng thí nghiệm hoặc mức năng lượng tiết kiệm. Chế độ mock và unit test có giá trị phát triển phần mềm nhưng không được chuyển thành kết quả phần cứng.")

    p("h2", "3.4. Đánh giá cơ chế quota, khuyến nghị và sa thải tải")
    p("body", "RQ4 cho thấy quota đã có vòng đời quản lý, guard và cảnh báo ở lớp phần mềm; vectơ dự báo 24 giờ tạo được đầu vào định lượng Ê_24 cho chỉ báo R. Tuy nhiên, chuỗi sáu bước mới là logic thiết kế: chưa có log quota thực, phản ứng người dùng, số lần tránh vượt ngưỡng hoặc năng lượng dịch chuyển để đánh giá tác động. Sa thải tự động vẫn bị khóa và chưa có thuật toán công suất tức thời, thử nghiệm interlock hoặc tải an toàn. Do đó bài chỉ công bố cơ chế hỗ trợ quyết định và thiết kế fail-closed, không công bố tối ưu hóa, tiết kiệm hay sa thải thành công.")

    p("h2", "3.5. Đánh giá tổng hợp theo mục tiêu nghiên cứu")
    p("body", "Bảng 4 đánh giá từng mục tiêu theo cùng một chuỗi: phương pháp, chỉ báo bằng chứng và kết luận được phép. Hai mục tiêu phần mềm và benchmark đã có chỉ báo tương ứng; đo lường vật lý, tác động của khuyến nghị và sa thải tải vẫn cần thực nghiệm. Cách trình bày này giữ ưu điểm của bài kỹ thuật theo chuỗi mô hình–triển khai–đánh giá, đồng thời ngăn kết quả UCI bị suy rộng sang phòng thí nghiệm Cần Thơ.")
    table4_caption = base.clone_paragraph(document, samples["table_caption"], "Bảng 4. Đánh giá tổng hợp các mục tiêu nghiên cứu", keep_next=True)
    table4_caption.paragraph_format.page_break_before = True
    base.add_table(document, objective_rows(canonical), [1.45, 1.80, 1.45, 1.95], size=7.7)
    base.clone_paragraph(document, samples["table_note"], "Ghi chú: “false” là cổng tuyên bố đang đóng, không đồng nghĩa thành phần không tồn tại trong thiết kế hoặc mã nguồn.")

    p("h2", "3.6. Thảo luận và nguy cơ đối với tính hợp lệ")
    p("body", "Construct validity bị giới hạn vì contract phần mềm không đo tiếp điểm và tải thật. Internal validity của benchmark được cải thiện nhờ tiền xử lý nhân quả, baseline và rolling-origin, nhưng cấu hình mô hình cố định chưa phải tìm kiếm siêu tham số toàn diện. External validity thấp vì UCI được thu tại một hộ ở Pháp, không đại diện khí hậu, hành vi hoặc cấu hình phòng thí nghiệm tại Cần Thơ. Conclusion validity bị giới hạn bởi ba fold, ba seed và chênh lệch Random Forest–XGBoost rất nhỏ. Việc công bố ranh giới bằng chứng làm giảm overclaim nhưng không loại bỏ các hạn chế này.")
    p("body", "Về tính thuyết phục của hình và bảng, ba hình hiện phục vụ ba luận điểm khác nhau: ranh giới kiến trúc, lệnh–phản hồi và kết quả dự báo. Bốn bảng lần lượt định vị nghiên cứu, kiểm kê bằng chứng, báo cáo benchmark và ánh xạ mục tiêu. Screenshot chatbot hoặc giao diện trang trí không được đưa vào vì không trả lời trực tiếp câu hỏi nghiên cứu; khi có nghiên cứu usability hoặc log quota, ảnh giao diện mới có thể trở thành bằng chứng bổ sung.")

    p("h1", "4. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    p("h2", "4.1. Kết luận")
    p("body", "Bài báo xác lập chuỗi triển khai MFM384 → PLC/gateway → telemetry → dự báo 24 giờ → quota/ưu tiên → khuyến nghị → lệnh có phản hồi. RQ1 được giải quyết ở mức kiến trúc nhưng chưa được xác nhận vật lý. RQ2 được hỗ trợ ở mức backend và hợp đồng phần mềm, trong khi toàn bộ frontend chưa đạt đầy đủ. RQ3 được trả lời bằng benchmark UCI: Random Forest và XGBoost gần tương đương, không có cơ sở tuyên bố mô hình vượt trội. RQ4 xác nhận quota, chỉ báo rủi ro và cảnh báo ở mức thiết kế/phần mềm nhưng sa thải tải tự động chưa được mở. IoT và quản lý người dùng đóng vai trò lớp triển khai; dự báo phụ tải vẫn là lõi AI của đề tài.")
    p("h2", "4.2. Hạn chế và kế hoạch kiểm chứng tiếp theo")
    p("body", "Công việc tiếp theo gồm: hoàn tất bốn ca frontend còn lỗi; thu raw log MFM384/PLC/phụ tải có timestamp và checksum; đánh giá sai số đo, RTT, p95/p99, timeout và recovery; huấn luyện/kiểm tra lại trên dữ liệu phòng thí nghiệm theo chia thời gian; công bố prediction artifact để vẽ actual–predicted; xác lập chính sách quota, mức ưu tiên và khảo sát phản hồi người dùng. Sa thải tải chỉ được kích hoạt sau khi có thuật toán theo kW, hysteresis, interlock, emergency stop, manual override và thử nghiệm an toàn có giám sát.")
    p("h2", "4.3. Dữ liệu, đạo đức và liêm chính nghiên cứu")
    p("body", "Dữ liệu công khai theo DOI 10.24432/C58K54 [14]. Kết quả sử dụng trong bài nằm trong research/results/canonical/canonical_results.json cùng script benchmark, thu bằng chứng, sinh hình và sinh DOCX. Nghiên cứu không tuyển người tham gia và không thử nghiệm trên người hoặc động vật. Credential, cấu hình mạng vận hành và dữ liệu địa phương phải được ẩn danh, kiểm soát truy cập và đánh giá an toàn trước khi công bố.")
    p("body", "Đóng góp tác giả theo CRediT, nguồn tài trợ, xung đột lợi ích và tác giả liên hệ vẫn là placeholder, phải được tác giả xác nhận trước khi nộp. Công cụ AI được dùng để hỗ trợ rà soát diễn đạt, kiểm thử mã, sinh hình và kiểm tra cấu trúc; tác giả chịu trách nhiệm về số liệu, trích dẫn, quyết định học thuật và bản nộp cuối cùng.")

    base.clone_paragraph(document, samples["ack"], "Lời cảm ơn (nếu có): [BỔ SUNG CƠ QUAN HỖ TRỢ/TÀI TRỢ SAU KHI TÁC GIẢ XÁC NHẬN].")
    base.clone_paragraph(document, samples["ref_heading"], "Tài liệu tham khảo")
    for reference in base.REFERENCES:
        paragraph = base.clone_paragraph(document, samples["ref"], reference)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def validate(args: argparse.Namespace, canonical: dict) -> None:
    for path in (args.template, args.canonical):
        if not path.exists():
            raise FileNotFoundError(path)
    for name in ("final_architecture_vi.png", "final_command_feedback_vi.png", "final_forecast_horizon_vi.png"):
        if not (args.figures_dir / name).exists():
            raise FileNotFoundError(args.figures_dir / name)
    if len(TITLE_VI.split()) != 20:
        raise ValueError(f"Vietnamese title must contain exactly 20 words, found {len(TITLE_VI.split())}")
    for label, abstract in zip(("Vietnamese", "English"), abstracts(canonical)):
        count = len(abstract.split())
        if not 150 <= count <= 200:
            raise ValueError(f"{label} abstract has {count} words")
    if not canonical.get("softwareEvidence"):
        raise ValueError("softwareEvidence is required")


def main() -> int:
    args = parse_args()
    canonical = json.loads(args.canonical.read_text(encoding="utf-8"))
    validate(args, canonical)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.template, args.output)
    document = Document(args.output)
    if len(document.tables) < 3 or len(document.paragraphs) < 40:
        raise ValueError("Unexpected journal template structure")
    samples = {
        "h1": document.paragraphs[10],
        "body": document.paragraphs[11],
        "h2": document.paragraphs[16],
        "h3": document.paragraphs[21],
        "table_caption": document.paragraphs[25],
        "table_note": document.paragraphs[26],
        "image": document.paragraphs[29],
        "figure_caption": document.paragraphs[30],
        "ack": document.paragraphs[36],
        "ref_heading": document.paragraphs[37],
        "ref": document.paragraphs[39],
    }
    populate_front_matter(document, canonical)
    base.clear_after_front_matter(document)
    build_body(document, canonical, args.figures_dir, samples)
    document.core_properties.title = TITLE_VI
    document.core_properties.subject = "Hệ thống giám sát năng lượng IoT hỗ trợ dự báo phụ tải bằng AI"
    document.core_properties.keywords = "IoT; giám sát năng lượng; dự báo phụ tải; quota; RBAC"
    document.save(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
