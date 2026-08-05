from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


REFERENCES = [
    "[1] G. Hebrail and A. Berard, Individual Household Electric Power Consumption [Dataset], UCI Machine Learning Repository, 2006, doi:10.24432/C58K54.",
    "[2] L. Breiman, Random Forests, Machine Learning, vol. 45, pp. 5-32, 2001, doi:10.1023/A:1010933404324.",
    "[3] T. Chen and C. Guestrin, XGBoost: A Scalable Tree Boosting System, Proc. KDD, pp. 785-794, 2016, doi:10.1145/2939672.2939785.",
    "[4] C. Li et al., Power Load Forecasting Based on the Combined Model of LSTM and XGBoost, Proc. PRAI, pp. 46-51, 2019, doi:10.1145/3357777.3357792.",
    "[5] I. Gomes et al., Recent Techniques Used in Home Energy Management Systems: A Review, Energies, vol. 15, no. 8, 2866, 2022, doi:10.3390/en15082866.",
    "[6] N. Koltsaklis et al., Smart home energy management processes support through machine learning algorithms, Energy Reports, vol. 8, Suppl. 3, pp. 1-6, 2022, doi:10.1016/j.egyr.2022.01.033.",
    "[7] M. Kim et al., Stochastic optimization of home energy management system using clustered quantile scenario reduction, Applied Energy, vol. 349, 121555, 2023, doi:10.1016/j.apenergy.2023.121555.",
    "[8] H. Mortaji et al., Load Shedding and Smart-Direct Load Control Using Internet of Things in Smart Grid Demand Response Management, IEEE Trans. Ind. Appl., vol. 53, no. 6, pp. 5155-5163, 2017, doi:10.1109/TIA.2017.2740832.",
    "[9] Siemens AG, SIMATIC S7-1200 Programmable Controller System Manual, V4.6, A5E02486680-AP, 2022.",
    "[10] Modbus Organization, MODBUS Application Protocol Specification V1.1b3, 2012.",
    "[11] M. Hlayel et al., Toward Industry 5.0: A WebSocket-S7 Bridge for Low-Latency, IEC 61588-Compliant Digital Twins in Remote Industrial Automation, PLOS ONE, vol. 21, no. 5, e0342004, 2026, doi:10.1371/journal.pone.0342004.",
    "[12] NIST, IoT Device Cybersecurity Capability Core Baseline, NISTIR 8259A, 2020, doi:10.6028/NIST.IR.8259A.",
    "[13] OWASP Foundation, Application Security Verification Standard, official project documentation.",
    "[14] NIST, Guide to Operational Technology (OT) Security, NIST SP 800-82 Rev. 3, 2023, doi:10.6028/NIST.SP.800-82r3.",
]

RELATED_WORK_ROWS = [
    ["Công trình", "PLC thật", "Forecast", "RBAC/service auth", "Feedback vật lý", "Vai trò đối chiếu"],
    ["Gomes et al. (2022)", "Không", "Tổng quan", "Không", "Không", "Phân loại kỹ thuật HEMS"],
    ["Koltsaklis et al. (2022)", "Không", "Có", "Không", "Không", "Forecast hỗ trợ tối ưu HEMS"],
    ["Kim et al. (2023)", "Không", "Xác suất", "Không", "Không", "Tối ưu dưới bất định"],
    ["Mortaji et al. (2017)", "IoT/điều khiển", "ARIMA", "Không nêu", "Không tương đương", "Direct load control"],
    ["Hlayel et al. (2026)", "S7-1500", "Không", "Không trọng tâm", "PLC-DT", "Latency PLC-cloud"],
    ["Công trình này", "S7-1200", "24 giờ", "Có", "statusTag độc lập", "Evidence-safe prototype"],
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate paper and thesis from canonical results.")
    parser.add_argument("--canonical", type=Path, required=True)
    parser.add_argument("--paper-template", type=Path, required=True)
    parser.add_argument("--thesis-template", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--references", type=Path, default=Path("Danh_Mục_Tài_Liệu_Tham_Khảo.md"))
    return parser.parse_args()


def load_references(path: Path) -> list[str]:
    references = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if re.match(r"^\[\d+\]\s", line):
            references.append(line)
    if len(references) != 22:
        raise ValueError(f"Expected 22 verified references, found {len(references)} in {path}")
    return references


def replace_section_references(document: Document, references: list[str]) -> None:
    heading_index = next(
        index for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip().upper() in {"TÀI LIỆU THAM KHẢO", "TÀI LIỆU THAM KHẢO"}
    )
    paragraphs = document.paragraphs
    stop_index = next(
        (index for index in range(heading_index + 1, len(paragraphs)) if paragraphs[index].style.name.startswith("Heading 1")),
        len(paragraphs),
    )
    existing = paragraphs[heading_index + 1:stop_index]
    for index, reference in enumerate(references):
        if index < len(existing):
            existing[index].text = reference
        else:
            paragraph = document.add_paragraph(reference)
            if stop_index < len(paragraphs):
                paragraphs[stop_index]._p.addprevious(paragraph._p)
    for paragraph in existing[len(references):]:
        paragraph._element.getparent().remove(paragraph._element)


def update_in_text_citations(document: Document) -> None:
    for paragraph in document.paragraphs:
        original = paragraph.text.strip()
        text = re.sub(r"\s*\[(?:\d+(?:\s*,\s*\d+)*)\]", "", original).strip()
        if text != original:
            paragraph.text = text
        if text.startswith("HEMS cần đồng thời"):
            paragraph.text = re.sub(r"(?:\s*\[\d+\][,;]?)+\s*$", "", text) + " [1],[2],[3]."
        elif text.startswith("S7-1200 cung cấp lớp điều khiển"):
            paragraph.text = (
                "S7-1200 cung cấp lớp điều khiển công nghiệp, trong khi Modbus định nghĩa function code, "
                "register và cơ chế trao đổi dữ liệu; triển khai RTU/RS485 cần thêm quy tắc serial, CRC, baud và parity [6],[7],[8]."
            )
        elif text.startswith("Các nghiên cứu HEMS nhấn mạnh"):
            paragraph.text = (
                "Các nghiên cứu HEMS và demand response nhấn mạnh việc phối hợp giám sát, dự báo và điều phối tải, "
                "nhưng kết quả của kiến trúc khác không thể thay thế phép đo trên hệ thống đang nghiên cứu [1],[2],[4],[5]."
            )
        elif text.startswith("Random Forest và XGBoost được chọn"):
            paragraph.text = (
                "Dự báo năng lượng đã được nghiên cứu bằng mô hình thống kê, cây quyết định và mạng LSTM [9],[10],[11]. "
                "Trong công trình này, XGBoost và Random Forest được đánh giá bằng cùng pipeline và cùng split [12],[13]. "
                "Dữ liệu UCI chỉ dùng cho benchmark công khai, không đại diện cho phụ tải Cần Thơ [14]."
            )
        elif text.startswith("MAPE sử dụng mẫu số tối thiểu"):
            paragraph.text = (
                "MAE là chỉ số chính; RMSE, R² và MAPE là chỉ số bổ sung. MAPE dùng floor 0,2 kW vì sai số phần trăm "
                "không ổn định khi tải gần 0 [15]. Dữ liệu được đánh giá theo chronological split và rolling-origin "
                "để mô phỏng dự báo tương lai [16]. ISO 50001 cung cấp khung quản lý và cải tiến hiệu quả năng lượng [17]."
            )
        elif text.startswith("Các quyết định xác thực") or "least privilege và logging" in text:
            paragraph.text = (
                "Threat model và biện pháp kiểm soát tham chiếu hướng dẫn smart-grid, application security, IoT baseline "
                "và OT security [19],[20],[21],[22]. Các tài liệu này định hướng risk control nhưng không thay thế "
                "hazard analysis hoặc interlock vật lý."
            )
        elif text.startswith("Ứng dụng mobile và admin site không giao tiếp trực tiếp"):
            paragraph.text = text + " Backend cục bộ đóng vai trò edge gateway giữa ứng dụng và PLC [18]."

    anchor = next((p for p in document.paragraphs if p.text.strip().startswith("HEMS cần đồng thời")), None)
    if anchor is not None:
        coverage = document.add_paragraph(
            "Nhóm tài liệu nền tảng được tổ chức theo thứ tự xuất hiện: HEMS và demand response [1],[2],[3],[4],[5]; "
            "Modbus và PLC [6],[7],[8]; forecasting, mô hình, dataset và đánh giá [9],[10],[11],[12],[13],[14],[15],[16]; "
            "quản lý năng lượng và edge computing [17],[18]; bảo mật smart-grid, ứng dụng, IoT và OT [19],[20],[21],[22]."
        )
        anchor._p.addnext(coverage._p)


def insert_related_work_table(document: Document, anchor_text: str) -> None:
    if any("Công trình này" in cell.text for table in document.tables for row in table.rows for cell in row.cells):
        return
    anchor = next(paragraph for paragraph in document.paragraphs if anchor_text in paragraph.text)
    lead = document.add_paragraph(
        "Bảng dưới đây làm rõ khoảng trống: các nghiên cứu HEMS tập trung dự báo hoặc tối ưu, "
        "còn nghiên cứu PLC-cloud tập trung truyền thông; ít công trình đồng thời đánh giá home-scoped authorization, "
        "service credential, serialized PLC write, feedback độc lập và provenance của kết quả."
    )
    anchor._p.addnext(lead._p)
    table = document.add_table(rows=1, cols=len(RELATED_WORK_ROWS[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for column, value in enumerate(RELATED_WORK_ROWS[0]):
        table.rows[0].cells[column].text = value
    for values in RELATED_WORK_ROWS[1:]:
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = value
    lead._p.addnext(table._tbl)


def set_forecast_table(document: Document, models: list[dict]) -> None:
    candidate = next((table for table in document.tables if "MAE" in " ".join(cell.text for cell in table.rows[0].cells)), None)
    if candidate is None:
        return
    while len(candidate.rows) > 1:
        candidate._tbl.remove(candidate.rows[-1]._tr)
    headers = [cell.text.strip().lower() for cell in candidate.rows[0].cells]
    for row in [item for item in models if item.get("split") == "test"]:
        cells = candidate.add_row().cells
        values = {
            "model": row.get("model"),
            "mô hình": row.get("model"),
            "mae": row.get("mae_kw"),
            "mae (kw)": row.get("mae_kw"),
            "rmse": row.get("rmse_kw"),
            "rmse (kw)": row.get("rmse_kw"),
            "mape": row.get("mape_percent"),
            "mape (%)": row.get("mape_percent"),
            "r²": row.get("r2"),
            "inference": row.get("inference_ms_per_sample"),
        }
        for index, header in enumerate(headers):
            value = values.get(header)
            if isinstance(value, float):
                std_key = {"mae": "mae_kw_std", "mae (kw)": "mae_kw_std", "rmse": "rmse_kw_std", "rmse (kw)": "rmse_kw_std", "mape": "mape_percent_std", "mape (%)": "mape_percent_std", "r²": "r2_std", "inference": "inference_ms_per_sample_std"}.get(header)
                std = row.get(std_key) if std_key else None
                cells[index].text = f"{value:.3f} ± {std:.3f}" if isinstance(std, float) else f"{value:.3f}"
            else:
                cells[index].text = str(value or "Chờ bằng chứng")


def add_generation_note(document: Document, canonical: dict) -> None:
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "Nguồn số liệu: canonical_results.json; generatedAtUtc="
        + str(canonical.get("generatedAtUtc"))
        + ". Không chỉnh số liệu thủ công."
    ).italic = True


def update_result_narrative(document: Document, canonical: dict) -> None:
    test_rows = {
        row["model"]: row for row in canonical["forecast"]["models"] if row.get("split") == "test"
    }
    xgb = test_rows["xgboost"]
    seasonal = test_rows["seasonal_naive_24h"]
    improvement = (seasonal["mae_kw"] - xgb["mae_kw"]) / seasonal["mae_kw"] * 100
    result_sentence = (
        f"Benchmark UCI mới dùng hai rolling-origin folds và hai random seed. XGBoost đạt MAE "
        f"{xgb['mae_kw']:.3f} ± {xgb['mae_kw_std']:.3f} kW và RMSE "
        f"{xgb['rmse_kw']:.3f} ± {xgb['rmse_kw_std']:.3f} kW trên test, thấp hơn seasonal-naive "
        f"24 giờ ({seasonal['mae_kw']:.3f} ± {seasonal['mae_kw_std']:.3f} kW) khoảng {improvement:.1f}% theo MAE. "
        f"MAPE vẫn cao ({xgb['mape_percent']:.1f} ± {xgb['mape_percent_std']:.1f}%), nên kết quả chỉ chứng minh "
        "lợi ích tương đối trên UCI và không đại diện cho độ chính xác tại Cần Thơ."
    )
    for paragraph in document.paragraphs:
        text = paragraph.text
        if "Artifact UCI cũ cho thấy" in text or "Kết quả artifact UCI cũ" in text:
            paragraph.text = result_sentence
        elif "A legacy UCI artifact reports" in text:
            paragraph.text = (
                "The repeated UCI benchmark uses two rolling-origin folds and two random seeds. "
                f"XGBoost obtains a test MAE of {xgb['mae_kw']:.3f} ± {xgb['mae_kw_std']:.3f} kW, "
                f"compared with {seasonal['mae_kw']:.3f} ± {seasonal['mae_kw_std']:.3f} kW for the 24-hour seasonal-naive baseline. "
                "The high MAPE and domain mismatch prevent these results from being generalized to Can Tho households."
            )
        elif "XGBoost có sai số thấp hơn nhẹ Random Forest trong artifact cũ" in text:
            paragraph.text = result_sentence
        elif "Không có persistence nên chưa biết" in text:
            paragraph.text = result_sentence
        elif "Kết quả AI hiện có chỉ là artifact sơ bộ trên dữ liệu UCI" in text:
            paragraph.text = (
                "Benchmark forecast công khai trên UCI đã được chạy lại bằng hai rolling-origin folds, hai random seed, "
                "persistence baseline và seasonal baseline. Kết quả này đủ cho claim so sánh trên dataset UCI, nhưng "
                "không đại diện cho độ chính xác trên dữ liệu MFM384 tại Cần Thơ; benchmark địa phương vẫn chờ dữ liệu thật."
            )
        elif text.startswith("Evidence gate hiện đánh dấu bốn claim là chưa được phép"):
            paragraph.text = (
                "Evidence gate hiện cho phép công bố benchmark UCI công khai. Ba claim vẫn chưa được phép là "
                "độ chính xác trên dữ liệu MFM384 địa phương, độ trễ phần cứng thật và hiệu quả sa thải tải tự động. "
                "Mỗi claim chỉ được mở khi canonical source ghi đủ raw data, số mẫu, nguồn PLC thật và fingerprint."
            )

    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers and headers[0] == "ID":
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == "AI-01":
                    row.cells[-1].text = "Đạt benchmark canonical (2 folds, 2 seeds)"
        if headers and headers[0] == "Claim":
            policy = canonical["claimPolicy"]
            for row in table.rows[1:]:
                key = row.cells[0].text.strip()
                if key in policy:
                    row.cells[1].text = str(bool(policy[key]))
                    row.cells[2].text = "Canonical evidence gate" if policy[key] else "Chờ bằng chứng tương ứng"


def generate(template: Path, output: Path, canonical: dict, thesis: bool, references: list[str]) -> None:
    shutil.copy2(template, output)
    document = Document(output)
    anchor = "Khoảng trống của prototype" if thesis else "Dữ liệu UCI"
    insert_related_work_table(document, anchor)
    set_forecast_table(document, canonical["forecast"]["models"])
    update_result_narrative(document, canonical)
    update_in_text_citations(document)
    replace_section_references(document, references)
    add_generation_note(document, canonical)
    document.save(output)


def main() -> int:
    args = parse_args()
    canonical = json.loads(args.canonical.read_text(encoding="utf-8"))
    references = load_references(args.references)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    generate(args.paper_template, args.output_dir / "HEMS_Paper_Canonical.docx", canonical, thesis=False, references=references)
    generate(args.thesis_template, args.output_dir / "Smart_Home_Thesis_Canonical.docx", canonical, thesis=True, references=references)
    print(args.output_dir / "HEMS_Paper_Canonical.docx")
    print(args.output_dir / "Smart_Home_Thesis_Canonical.docx")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
